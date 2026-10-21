"""
Camada de serviços do BolsaTrader:

- integração com a API pública brapi.dev (cotações da B3);
- cálculo de posições e percentual de lucro/perda;
- análise simples de tendência (alta/baixa) a partir do histórico de cotações;
- geração de alertas/lembretes.

Manter a lógica de negócio aqui (fora das views) facilita testes e reuso
tanto pelas views quanto pelos comandos de management (cron).
"""

from __future__ import annotations

import logging
import re
import statistics
import threading
import time
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone as dt_timezone
from decimal import Decimal, ROUND_HALF_UP
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from django.conf import settings
from django.contrib.auth import get_user_model
from django.utils import timezone
from django.utils.dateparse import parse_datetime

logger = logging.getLogger(__name__)

from django.db import transaction
from django.db.models import F, Q, Sum

from .models import (
    ConsumoApiBrapi, MensagemWhatsapp,
    Ativo, Cotacao, Operacao, Alerta, FonteNoticia, Noticia, AcaoB3, CotacaoIndice,
    RegistroAtualizacaoCarteira, RegistroCotacao, ContaCorrente, LancamentoContaCorrente, PostIt,
    CriptoAtivo, OperacaoCripto, CotacaoCripto, RegistroAtualizacaoCarteiraCripto, PrevisaoAnalista,
    FundamentosAtivo, VendaLote,
)
from tributario.models import (
    AvisoLimiteIsencao, Darf, EventoCorporativo, MapeamentoTitulo, NotaCorretagem, ParametroTributario, Provento,
    SaldoInicialFiscal, TaxaSelicMensal,
)


# --------------------------------------------------------------------------
# Atividade recente de todos os usuários (mural da comunidade)
# --------------------------------------------------------------------------
def atividade_recente(limite: int = 15) -> list[dict]:
    """
    Retorna as operações mais recentes de TODOS os usuários (nome + tipo de
    operação + quando), para um mural de atividade da comunidade. Não expõe
    ticker, quantidade nem preço - esses detalhes continuam privados de
    cada usuário.
    """
    operacoes = (
        Operacao.objects.select_related("usuario").order_by("-criado_em")[:limite]
    )
    return [
        {
            "nome": op.usuario.first_name or op.usuario.username,
            "tipo": op.tipo,
            "tipo_label": op.get_tipo_display(),
            "quando": op.criado_em,
        }
        for op in operacoes
    ]


# --------------------------------------------------------------------------
# Horário de negociação da B3 (configurável em B3_HORARIO_ABERTURA/
# B3_HORARIO_FECHAMENTO no .env)
# --------------------------------------------------------------------------
def mercado_b3_aberto() -> bool:
    """
    True quando agora é dia útil (segunda a sexta) e está dentro do horário
    de negociação configurado - usado tanto para mostrar o badge de mercado
    aberto/fechado quanto para bloquear "Atualizar cotações agora" fora do
    horário (fora do pregão a cotação não muda, então não há nada de novo
    pra buscar).
    """
    from datetime import datetime

    abertura = datetime.strptime(settings.B3_HORARIO_ABERTURA, "%H:%M").time()
    fechamento = datetime.strptime(settings.B3_HORARIO_FECHAMENTO, "%H:%M").time()

    agora = timezone.localtime()
    dia_util = agora.weekday() < 5  # segunda(0) a sexta(4)
    return dia_util and abertura <= agora.time() <= fechamento


# --------------------------------------------------------------------------
# Integração com a API de cotações (brapi.dev)
# --------------------------------------------------------------------------
class BrapiError(Exception):
    """Erro ao consultar a API de cotações."""


def _detalhe_limite_brapi(resposta) -> str | None:
    """
    A brapi.dev responde 429 tanto pra um limite passageiro (por minuto -
    resolve tentando de novo em instantes) quanto pro limite MENSAL do
    plano gratuito esgotado (só volta a funcionar na próxima virada do
    ciclo, às vezes semanas depois - "tente novamente em instantes" seria
    enganoso nesse caso). O corpo da resposta traz um "code" que distingue
    os dois; quando é o limite mensal, monta uma mensagem com a data real
    de renovação em vez da genérica.
    """
    try:
        corpo = resposta.json()
    except (ValueError, AttributeError):
        return None

    if corpo.get("code") != "MONTHLY_LIMIT_EXCEEDED":
        return None

    resets_at = (corpo.get("details") or {}).get("usage", {}).get("resetsAt")
    instante = parse_datetime(resets_at) if resets_at else None
    quando = (
        f"em {timezone.localtime(instante).strftime('%d/%m/%Y')}" if instante else "no início do próximo ciclo"
    )
    return f"Limite mensal de consultas da API de cotações foi atingido - só volta a funcionar {quando}"


def _mensagem_falha_api(exc: requests.RequestException, contexto: str) -> str:
    """
    Mensagem amigável para falhas de rede numa API externa (brapi.dev ou
    Banco Central), sem embutir a exceção crua na mensagem - o texto de uma
    requests.RequestException inclui a URL completa da requisição, e para a
    brapi.dev isso significa expor o BRAPI_TOKEN em texto puro. Essas
    mensagens às vezes chegam até a tela do usuário (ver
    core.views.consultar_cotacao_avulsa, acoes_b3_atualizar,
    atualizar_benchmarks_agora), então vazar o token ali seria uma falha de
    segurança desnecessária.

    Trata especificamente o limite de requisições (HTTP 429) com uma
    mensagem própria, distinguindo o limite mensal esgotado (ver
    _detalhe_limite_brapi) do limite passageiro por minuto - são situações
    bem diferentes na prática.
    """
    resposta = getattr(exc, "response", None)
    status = getattr(resposta, "status_code", None)
    if status == 429:
        detalhe_mensal = _detalhe_limite_brapi(resposta) if resposta is not None else None
        if detalhe_mensal:
            return f"{detalhe_mensal} ({contexto})."
        return f"Limite de consultas da API atingido para {contexto}. Tente novamente em alguns minutos."
    return f"Falha ao consultar a API para {contexto}. Tente novamente em instantes."


# --------------------------------------------------------------------------
# Orçamento de requisições à brapi.dev
#
# O plano contratado tem um limite por ciclo de 30 dias (BRAPI_LIMITE_MENSAL,
# Startup = 150.000) e estourar deixa o sistema sem cotações até a renovação.
# Toda requisição à brapi.dev passa por _get_brapi, que conta num contador
# diário (ConsumoApiBrapi, compartilhado entre os processos que usam o mesmo
# banco) e trava antes de estourar.
# --------------------------------------------------------------------------
BRAPI_TETO_ABSOLUTO_PCT = 95  # acima disso NENHUMA requisição é feita (nem manual)
DIAS_UTEIS_POR_CICLO = 22  # dias úteis médios em 30 dias corridos


def consumo_api_ultimos_30_dias() -> int:
    """Requisições feitas à brapi.dev nos últimos 30 dias (janela móvel, conservadora frente ao ciclo do plano)."""
    inicio = timezone.localdate() - timedelta(days=29)
    return ConsumoApiBrapi.objects.filter(data__gte=inicio).aggregate(total=Sum("requisicoes"))["total"] or 0


def limite_automatico_api() -> int:
    """Consumo a partir do qual o ciclo AUTOMÁTICO pausa - deixa a margem de segurança para uso manual/imprevistos."""
    return int(settings.BRAPI_LIMITE_MENSAL * (100 - settings.BRAPI_MARGEM_SEGURANCA_PCT) / 100)


def limite_absoluto_api() -> int:
    """Consumo a partir do qual QUALQUER requisição é bloqueada."""
    return int(settings.BRAPI_LIMITE_MENSAL * BRAPI_TETO_ABSOLUTO_PCT / 100)


def orcamento_automatico_disponivel() -> bool:
    return consumo_api_ultimos_30_dias() < limite_automatico_api()


def projetar_consumo_mensal(qtd_ativos: int, intervalo_minutos: int | None = None) -> dict:
    """
    Estima quantas requisições/mês o ciclo automático consome: ciclos por dia
    (só dentro do pregão) x requisições por ciclo (ativos / tickers por lote)
    x dias úteis. Serve para conferir se COTACOES_INTERVALO_MINUTOS cabe no
    orçamento antes de o consumo real mostrar o problema.
    """
    from datetime import datetime as _dt

    intervalo = intervalo_minutos or settings.COTACOES_INTERVALO_MINUTOS
    abertura = _dt.strptime(settings.B3_HORARIO_ABERTURA, "%H:%M")
    fechamento = _dt.strptime(settings.B3_HORARIO_FECHAMENTO, "%H:%M")
    minutos_pregao = max(int((fechamento - abertura).total_seconds() // 60), 0)
    ciclos_dia = minutos_pregao // max(intervalo, 1) + 1
    lote = max(getattr(settings, "BRAPI_TICKERS_POR_LOTE", BRAPI_TICKERS_POR_LOTE_PADRAO), 1)
    requisicoes_por_ciclo = -(-qtd_ativos // lote) if qtd_ativos else 0  # divisão com teto
    mensal = ciclos_dia * requisicoes_por_ciclo * DIAS_UTEIS_POR_CICLO
    return {
        "ciclos_dia": ciclos_dia,
        "requisicoes_por_ciclo": requisicoes_por_ciclo,
        "requisicoes_mes": mensal,
        "limite_automatico": limite_automatico_api(),
        "cabe_no_orcamento": mensal <= limite_automatico_api(),
        "percentual_do_limite": round(mensal / settings.BRAPI_LIMITE_MENSAL * 100, 1),
    }


def _get_brapi(url: str, params: dict, timeout: int):
    """requests.get para a brapi.dev com controle de orçamento: bloqueia acima do teto absoluto e conta a requisição."""
    if consumo_api_ultimos_30_dias() >= limite_absoluto_api():
        raise BrapiError(
            "Orçamento mensal de requisições à API de cotações quase esgotado - consultas bloqueadas por "
            "segurança até o consumo dos últimos 30 dias baixar."
        )
    dia = ConsumoApiBrapi.objects.get_or_create(data=timezone.localdate())[0]
    ConsumoApiBrapi.objects.filter(pk=dia.pk).update(requisicoes=F("requisicoes") + 1)
    return requests.get(url, params=params, timeout=timeout)


def buscar_cotacao_atual(ticker: str) -> dict:
    """
    Consulta a cotação atual de um ticker na API brapi.dev.

    Retorna um dicionário com pelo menos: regularMarketPrice, shortName,
    regularMarketChangePercent.

    Documentação: https://brapi.dev/docs
    """
    url = f"{settings.BRAPI_BASE_URL}/quote/{ticker.upper()}"
    params = {}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN

    try:
        resposta = _get_brapi(url, params=params, timeout=10)
        resposta.raise_for_status()
        dados = resposta.json()
    except requests.RequestException as exc:
        raise BrapiError(_mensagem_falha_api(exc, f"o ticker {ticker}")) from exc

    resultados = dados.get("results") or []
    if not resultados:
        raise BrapiError(f"Ticker {ticker} não encontrado na API de cotações.")

    return resultados[0]


def buscar_cotacao_atual_com_historico(ticker: str, dias_historico: int = 30) -> tuple[dict, list[float]]:
    """
    Busca a cotação atual de um ticker JUNTO com um histórico recente de
    fechamentos, numa única chamada à API - a brapi.dev retorna os mesmos
    campos de "cotação agora" (regularMarketPrice etc.) e também
    "historicalDataPrice" na mesma resposta quando os parâmetros
    range/interval são informados, então dá pra ter os dois sem uma segunda
    requisição.

    Usado pela consulta avulsa (ver core.views.consultar_cotacao_avulsa)
    para calcular o IFR/RSI (ver calcular_rsi) de qualquer ticker pesquisado,
    mesmo um que o usuário não tenha em carteira, sem gastar chamada extra
    de API.

    Retorna (dados_atuais, precos_cronologicos) - o segundo item já vem do
    fechamento mais antigo para o mais recente, pronto para calcular_rsi.
    """
    url = f"{settings.BRAPI_BASE_URL}/quote/{ticker.upper()}"
    params = {"range": _range_brapi(dias_historico), "interval": "1d"}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN

    try:
        resposta = _get_brapi(url, params=params, timeout=15)
        resposta.raise_for_status()
        dados = resposta.json()
    except requests.RequestException as exc:
        raise BrapiError(_mensagem_falha_api(exc, f"o ticker {ticker}")) from exc

    resultados = dados.get("results") or []
    if not resultados:
        raise BrapiError(f"Ticker {ticker} não encontrado na API de cotações.")

    dados_atuais = resultados[0]

    pontos_ordenados = []
    vistos = set()
    for ponto in dados_atuais.get("historicalDataPrice") or []:
        fechamento = ponto.get("close")
        data_convertida = _converter_data_historico(ponto.get("date"))
        if fechamento is None or data_convertida is None or data_convertida in vistos:
            continue
        vistos.add(data_convertida)
        pontos_ordenados.append((data_convertida, float(fechamento)))
    pontos_ordenados.sort(key=lambda item: item[0])

    precos_cronologicos = [preco for _, preco in pontos_ordenados]
    return dados_atuais, precos_cronologicos


# brapi.dev aceita vários tickers separados por vírgula no mesmo endpoint de
# cotação avulsa (ex: /quote/PETR4,VALE3,ITUB4) - agrupar em lotes assim, em
# vez de uma requisição HTTP por ativo, reduz drasticamente o número de
# chamadas à API e ajuda a evitar o erro 429 ("Too Many Requests"). O número
# máximo de tickers por requisição depende do PLANO contratado na brapi.dev
# (ex: gratuito e Startup = 10, Pro = 20 - pedir mais que o limite do plano
# não dá 429, dá 400 "QUOTES_PER_REQUEST_EXCEEDED", e o lote inteiro falha em
# silêncio) - configurável via BRAPI_TICKERS_POR_LOTE no .env pra acompanhar
# o plano sem precisar mexer no código.
BRAPI_TICKERS_POR_LOTE_PADRAO = 10


def buscar_cotacoes_em_lote(tickers: list[str]) -> dict[str, dict]:
    """
    Busca a cotação atual de vários tickers de uma vez (ver
    BRAPI_TICKERS_POR_LOTE no .env), em lotes de tamanho fixo. Retorna um
    dict {TICKER: dados_da_api} só com os tickers efetivamente encontrados -
    tickers não encontrados ou cujo lote falhou (erro de rede, 429, ou 400
    quando o lote excede o limite de ativos por requisição do plano
    contratado) simplesmente não aparecem no retorno; cabe ao chamador (ver
    atualizar_cotacoes_ativos) tratar a ausência como falha por ativo, sem
    que a falha de um lote derrube os demais lotes que deram certo.
    """
    tamanho_lote = getattr(settings, "BRAPI_TICKERS_POR_LOTE", BRAPI_TICKERS_POR_LOTE_PADRAO) or 1
    resultado = {}
    for inicio in range(0, len(tickers), tamanho_lote):
        lote = tickers[inicio:inicio + tamanho_lote]
        url = f"{settings.BRAPI_BASE_URL}/quote/{','.join(t.upper() for t in lote)}"
        params = {}
        if settings.BRAPI_TOKEN:
            params["token"] = settings.BRAPI_TOKEN

        try:
            resposta = _get_brapi(url, params=params, timeout=15)
            resposta.raise_for_status()
            dados = resposta.json()
        except requests.RequestException:
            continue  # esse lote falhou - os tickers dele ficam de fora do resultado

        for item in dados.get("results") or []:
            simbolo = (item.get("symbol") or "").upper()
            if simbolo:
                resultado[simbolo] = item

    return resultado


def atualizar_cotacoes_ativos(ativos) -> tuple[int, int]:
    """
    Atualiza a cotação de uma lista/queryset de ativos com uma (ou poucas,
    ver BRAPI_TICKERS_POR_LOTE no .env) chamada(s) em lote à API, em vez de
    uma requisição HTTP por ativo - usado tanto pelo botão "Atualizar
    cotações agora" quanto pelo ciclo automático (ver
    executar_ciclo_atualizacao_cotacoes). Retorna (atualizados, falhas).
    """
    ativos = list(ativos)
    try:
        lote = buscar_cotacoes_em_lote([a.ticker for a in ativos])
    except BrapiError:
        lote = {}  # orçamento de requisições esgotado (ver _get_brapi) - todos contam como falha

    atualizados, falhas = 0, 0
    for ativo in ativos:
        dados_api = lote.get(ativo.ticker.upper())
        try:
            if dados_api is None:
                raise BrapiError(f"Ticker {ativo.ticker} não encontrado na consulta em lote.")
            atualizar_cotacao_diaria(ativo, dados_api=dados_api)
            atualizados += 1
        except BrapiError:
            falhas += 1

    return atualizados, falhas


def buscar_maiores_variacoes(limite: int = 5) -> tuple[list[dict], list[dict]]:
    """
    Consulta a lista geral de ações da B3 na API brapi.dev e retorna as
    maiores altas e as maiores baixas do dia (ordenadas pela variação
    percentual atual), cada uma limitada a `limite` itens.
    """
    url = f"{settings.BRAPI_BASE_URL}/quote/list"
    params_base = {"type": "stock", "limit": limite}
    if settings.BRAPI_TOKEN:
        params_base["token"] = settings.BRAPI_TOKEN

    def _consultar(sort_order: str) -> list[dict]:
        params = {**params_base, "sortBy": "change", "sortOrder": sort_order}
        try:
            resposta = _get_brapi(url, params=params, timeout=10)
            resposta.raise_for_status()
            dados = resposta.json()
        except requests.RequestException as exc:
            raise BrapiError(_mensagem_falha_api(exc, "as maiores variações do dia")) from exc
        return dados.get("stocks") or []

    maiores_altas = _consultar("desc")
    maiores_baixas = _consultar("asc")
    return maiores_altas, maiores_baixas


def _para_decimal(valor) -> Decimal | None:
    """Converte um valor numérico da API para Decimal, tolerando None/ausente."""
    if valor is None:
        return None
    try:
        return Decimal(str(valor)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (ArithmeticError, ValueError, TypeError):
        return None


def sincronizar_acoes_b3() -> int:
    """
    Sincroniza o catálogo de ações da B3 (AcaoB3) com a lista atual da API
    brapi.dev: apaga o catálogo inteiro e cadastra de novo a partir do que a
    API retornar agora - botão "Atualizar lista" na tela Ações da B3.

    Se a API não retornar nenhuma ação, trata como falha (BrapiError) em vez
    de apagar o catálogo existente - evita esvaziar a lista por causa de uma
    leitura ruim. Retorna quantas ações foram cadastradas.
    """
    url = f"{settings.BRAPI_BASE_URL}/quote/list"
    params = {"type": "stock", "limit": 1000}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN

    try:
        resposta = _get_brapi(url, params=params, timeout=30)
        resposta.raise_for_status()
        dados = resposta.json()
    except requests.RequestException as exc:
        raise BrapiError(_mensagem_falha_api(exc, "a lista de ações da B3")) from exc

    acoes = dados.get("stocks") or []
    if not acoes:
        raise BrapiError("A API não retornou nenhuma ação da B3 agora.")

    vistos = set()
    novas = []
    for item in acoes:
        ticker = (item.get("stock") or "").strip().upper()
        if not ticker or ticker in vistos:
            continue  # ignora entradas sem ticker ou duplicadas na resposta da API
        vistos.add(ticker)
        novas.append(AcaoB3(
            ticker=ticker,
            nome=item.get("name") or "",
            setor=item.get("sector") or "",
            logo_url=item.get("logo") or "",
            preco_atual=_para_decimal(item.get("close")),
            variacao_dia_pct=_para_decimal(item.get("change")),
        ))

    AcaoB3.objects.all().delete()
    AcaoB3.objects.bulk_create(novas)
    return len(novas)


def atualizar_cotacao_diaria(ativo: Ativo, dados_api: dict | None = None) -> Cotacao:
    """Busca (ou recebe) a cotação atual de um ativo e grava/atualiza o registro do dia."""
    if dados_api is None:
        dados_api = buscar_cotacao_atual(ativo.ticker)

    # a brapi.dev às vezes aninha os campos em "data" (ex: consultas em lote);
    # trata os dois formatos pra não quebrar se a API mudar o formato de novo.
    dados = dados_api.get("data") if isinstance(dados_api.get("data"), dict) else dados_api

    preco = dados.get("regularMarketPrice")
    variacao = dados.get("regularMarketChangePercent")

    if preco is None:
        raise BrapiError(f"API não retornou preço para {ativo.ticker}.")

    ativo.nome = dados.get("shortName") or ativo.nome
    ativo.nome_longo = dados.get("longName") or ativo.nome_longo
    ativo.moeda = dados.get("currency") or ativo.moeda
    ativo.logo_url = dados.get("logourl") or ativo.logo_url
    ativo.maxima_dia = _para_decimal(dados.get("regularMarketDayHigh"))
    ativo.minima_dia = _para_decimal(dados.get("regularMarketDayLow"))
    ativo.abertura = _para_decimal(dados.get("regularMarketOpen"))
    ativo.fechamento_anterior = _para_decimal(dados.get("regularMarketPreviousClose"))
    ativo.variacao_dia_valor = _para_decimal(dados.get("regularMarketChange"))
    ativo.volume = dados.get("regularMarketVolume")
    ativo.valor_mercado = _para_decimal(dados.get("marketCap"))
    ativo.minima_52_semanas = _para_decimal(dados.get("fiftyTwoWeekLow"))
    ativo.maxima_52_semanas = _para_decimal(dados.get("fiftyTwoWeekHigh"))

    hora_cotacao = dados.get("regularMarketTime")
    if hora_cotacao:
        ativo.hora_cotacao = parse_datetime(hora_cotacao)

    if not ativo.setor:
        # nem toda resposta da brapi.dev traz "sector" na cotação avulsa (é
        # mais comum vir na lista usada por sincronizar_acoes_b3) - quando
        # falta, tenta completar pelo catálogo Ações da B3 já sincronizado
        # (ver core.services.sincronizar_acoes_b3), sem custo de mais uma
        # chamada de API. Usado pela análise de concentração por setor da
        # carteira (ver calcular_concentracao_setor).
        setor = dados.get("sector") or dados.get("sectorName")
        if not setor:
            catalogo = AcaoB3.objects.filter(ticker=ativo.ticker).values_list("setor", flat=True).first()
            setor = catalogo or ""
        ativo.setor = setor

    ativo.save()  # também atualiza 'atualizado_em' (auto_now) com o horário desta atualização

    cotacao, _ = Cotacao.objects.update_or_create(
        ativo=ativo,
        data=timezone.localdate(),
        defaults={
            "preco_fechamento": Decimal(str(preco)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
            "variacao_dia_pct": Decimal(str(variacao)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if variacao is not None
            else None,
            "volume": dados.get("regularMarketVolume"),
            "maxima": ativo.maxima_dia,
            "minima": ativo.minima_dia,
        },
    )
    return cotacao


# --------------------------------------------------------------------------
# Ciclo completo de atualização automática (cotações + alertas + robô +
# histórico da carteira) - lógica compartilhada entre o comando de
# management "atualizar_cotacoes" (--loop) e o agendador embutido no
# próprio processo do servidor (ver iniciar_agendador_cotacoes_embutido,
# ligado em bolsatrader/wsgi.py) - um único lugar pra não duplicar a regra
# entre os dois "chamadores".
# --------------------------------------------------------------------------
def executar_ciclo_atualizacao_cotacoes() -> dict:
    """
    Atualiza a cotação de todos os ativos com alguma operação registrada (e
    dos que têm previsão de analistas ainda dentro do prazo - ver
    ativos_com_previsao_em_andamento), depois gera os alertas de
    lucro/perda/tendência, roda o robô consultor e grava o retrato da
    carteira (ver registrar_atualizacao_carteira) de cada usuário com
    operações. Retorna um resumo numérico do que foi feito.
    """
    ativos = Ativo.objects.filter(
        Q(operacoes__isnull=False) | Q(id__in=ativos_com_previsao_em_andamento())
    ).distinct()
    atualizados, falhas = atualizar_cotacoes_ativos(ativos)

    total_alertas, total_sinais_robo = 0, 0
    cache_ia_diaria = {}
    for usuario in get_user_model().objects.filter(operacoes__isnull=False).distinct():
        total_alertas += len(gerar_alertas_para_usuario(usuario))
        total_sinais_robo += len(gerar_sinais_robo_para_usuario(usuario))
        notificar_compras_scanner_telegram(usuario)
        notificar_reservas_atingidas_telegram(usuario)
        enviar_ia_diaria_telegram(usuario, cache_ia_diaria)
        _lembrar_darfs(usuario)
        registrar_atualizacao_carteira(usuario)
        registrar_cotacoes_compradas(usuario, RegistroCotacao.AUTOMATICA)

    return {
        "ativos_total": ativos.count(),
        "ativos_atualizados": atualizados,
        "ativos_falha": falhas,
        "alertas_gerados": total_alertas,
        "sinais_robo_gerados": total_sinais_robo,
    }


def _lembrar_darfs(usuario) -> None:
    """
    Lembrete de DARF a vencer/vencida pelo Telegram (ver tributario.darf).
    Importado aqui dentro porque tributario.darf importa este módulo; e uma
    falha na apuração de um usuário não pode interromper o ciclo dos outros.
    """
    from tributario.darf import notificar_darfs_telegram
    try:
        notificar_darfs_telegram(usuario)
    except Exception:
        logger.exception("Falha ao verificar as DARFs de %s.", usuario.get_username())


_agendador_cotacoes_lock = threading.Lock()
_agendador_cotacoes_iniciado = False


def iniciar_agendador_cotacoes_embutido() -> bool:
    """
    Inicia, numa thread em segundo plano dentro do próprio processo do
    servidor web, o mesmo ciclo do comando "atualizar_cotacoes --loop" -
    repete a cada COTACOES_INTERVALO_MINUTOS minutos (.env), sem precisar
    manter um processo/terminal separado rodando esse comando à parte.

    Ligado em bolsatrader/wsgi.py (só quando um servidor WSGI de verdade
    carrega o app - Waitress/gunicorn -, nunca durante "manage.py migrate",
    "test" etc., já que esses comandos não importam o módulo wsgi).
    Idempotente (chamar mais de uma vez não inicia uma segunda thread) e
    desligado por padrão quando COTACOES_INTERVALO_MINUTOS é zero/negativo,
    ou explicitamente via AGENDADOR_COTACOES_EMBUTIDO=False no .env (usado
    em deploys com múltiplos processos worker - Docker/gunicorn com
    --workers > 1 -, onde cada worker chamaria isto e duplicaria o
    trabalho; nesse caso, use o comando "atualizar_cotacoes --loop" à parte,
    em um único processo, em vez do agendador embutido).

    Retorna True se a thread foi (ou já estava) iniciada, False se está
    desligado por configuração.
    """
    global _agendador_cotacoes_iniciado

    if not getattr(settings, "AGENDADOR_COTACOES_EMBUTIDO", True):
        return False

    intervalo = getattr(settings, "COTACOES_INTERVALO_MINUTOS", 0)
    if not intervalo or intervalo <= 0:
        return False

    with _agendador_cotacoes_lock:
        if _agendador_cotacoes_iniciado:
            return True
        _agendador_cotacoes_iniciado = True

    def _loop():
        # dorme antes do primeiro ciclo (em vez de rodar na hora) de propósito:
        # a thread é iniciada na importação de bolsatrader/wsgi.py, junto com
        # a subida do processo do servidor - rodar um ciclo completo (rede +
        # banco, para todos os usuários) nesse instante competiria com o
        # próprio servidor terminando de subir. Quem precisa de uma
        # atualização imediata usa o botão "Atualizar cotações agora".
        while True:
            time.sleep(intervalo * 60)
            _ciclo_agendador_embutido()

    threading.Thread(target=_loop, name="bolsatrader-cotacoes-agendador", daemon=True).start()
    # aviso no console do servidor: o intervalo é lido do .env só na subida do
    # processo - mudar COTACOES_INTERVALO_MINUTOS com o servidor rodando não
    # tem efeito até reiniciar, e sem esta linha não dá pra saber qual valor
    # está valendo de verdade.
    print(
        f"[BolsaTrader] Agendador de cotações ligado: a cada {intervalo} minuto(s), "
        "só durante o pregão da B3.", flush=True,
    )
    return True


def _ciclo_agendador_embutido() -> None:
    """
    Um "tick" do agendador embutido (ver iniciar_agendador_cotacoes_embutido):
    só executa o ciclo de atualização dentro do horário de negociação da B3
    (B3_HORARIO_ABERTURA/FECHAMENTO no .env) - fora do pregão o preço não
    muda, então não há nada de novo pra buscar. Extraído à parte pra poder
    ser testado sem precisar controlar a thread/sleep de verdade.
    """
    if not mercado_b3_aberto():
        return
    if not orcamento_automatico_disponivel():
        logger.warning(
            "Atualização automática de cotações pausada: %s requisições nos últimos 30 dias, acima do limite "
            "automático de %s (margem de segurança do plano).",
            consumo_api_ultimos_30_dias(), limite_automatico_api(),
        )
        return
    try:
        resultado = executar_ciclo_atualizacao_cotacoes()
        logger.info(
            "Atualização automática de cotações: %s/%s ativo(s), %s alerta(s), %s sinal(is) do robô.",
            resultado["ativos_atualizados"], resultado["ativos_total"],
            resultado["alertas_gerados"], resultado["sinais_robo_gerados"],
        )
    except Exception:
        logger.exception("Falha na atualização automática de cotações (agendador embutido).")


# --------------------------------------------------------------------------
# Histórico de preços (brapi.dev) e backfill de Cotacao
#
# Usado tanto para preencher de uma vez o histórico de um ativo recém-
# cadastrado (RSI/MACD exigem ~35 pregões - sem isso o robô consultor fica
# "Aguardando histórico" por semanas, um dia de cada vez) quanto para o
# histórico do Ibovespa usado no comparativo com benchmarks (ver
# atualizar_benchmarks, em core.services).
# --------------------------------------------------------------------------
_FAIXAS_RANGE_BRAPI = [
    (30, "1mo"), (90, "3mo"), (180, "6mo"), (365, "1y"), (730, "2y"), (1825, "5y"),
]


def _range_brapi(dias: int) -> str:
    """
    Converte uma quantidade de dias no parâmetro `range` aceito pela
    brapi.dev - o menor range que cobre os dias pedidos, limitado ao
    BRAPI_HISTORICO_MAXIMO do .env: o plano gratuito só aceita até "3mo"
    (mais que isso dá 400 Bad Request), os pagos aceitam "1y", "2y"... Um
    pedido acima do máximo é cortado no máximo em vez de estourar a requisição.
    """
    maximo = settings.BRAPI_HISTORICO_MAXIMO
    limite_maximo = next((limite for limite, faixa in _FAIXAS_RANGE_BRAPI if faixa == maximo), 90)
    for limite, range_str in _FAIXAS_RANGE_BRAPI:
        if dias <= limite or limite >= limite_maximo:
            return range_str
    return maximo


def _converter_data_historico(valor) -> date | None:
    """
    Converte a data de um ponto do histórico da brapi.dev para date - a API
    retorna timestamp Unix (segundos, UTC) no campo "date" de
    historicalDataPrice, mas trata string ISO também, por segurança caso o
    formato mude (mesma cautela já aplicada em atualizar_cotacao_diaria).
    """
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        try:
            return datetime.fromtimestamp(valor, tz=dt_timezone.utc).date()
        except (ValueError, OSError, OverflowError):
            return None
    if isinstance(valor, str):
        convertido = parse_datetime(valor)
        if convertido:
            return convertido.date()
        try:
            return date.fromisoformat(valor[:10])
        except ValueError:
            return None
    return None


def buscar_historico_precos(ticker: str, dias: int = 180) -> list[dict]:
    """
    Busca o histórico diário de fechamento (e máxima/mínima/volume) de um
    ticker (ação da B3, ou um índice como "^BVSP" para o Ibovespa) na API
    brapi.dev.

    Retorna uma lista de {"data": date, "fechamento": Decimal, "volume": int
    | None, "maxima": Decimal | None, "minima": Decimal | None}, em ordem
    cronológica (mais antigo primeiro) e sem datas duplicadas.
    """
    url = f"{settings.BRAPI_BASE_URL}/quote/{ticker.upper()}"
    params = {"range": _range_brapi(dias), "interval": "1d"}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN

    try:
        resposta = _get_brapi(url, params=params, timeout=30)
        resposta.raise_for_status()
        dados = resposta.json()
    except requests.RequestException as exc:
        raise BrapiError(_mensagem_falha_api(exc, f"o histórico de {ticker}")) from exc

    resultados = dados.get("results") or []
    if not resultados:
        raise BrapiError(f"Ticker {ticker} não encontrado na API de cotações (histórico).")

    pontos = resultados[0].get("historicalDataPrice") or []

    historico = []
    vistos = set()
    for ponto in pontos:
        fechamento = ponto.get("close")
        data_convertida = _converter_data_historico(ponto.get("date"))
        if fechamento is None or data_convertida is None or data_convertida in vistos:
            continue
        vistos.add(data_convertida)
        preco = _para_decimal(fechamento)
        if preco is not None:
            # volume, máxima e mínima já vêm de graça no mesmo ponto do
            # histórico (sem custo extra de API) - usados pelo Scanner
            # Técnico (ver analisar_volume_precos, calcular_atr,
            # calcular_suporte_resistencia)
            historico.append({
                "data": data_convertida, "fechamento": preco, "volume": ponto.get("volume"),
                "maxima": _para_decimal(ponto.get("high")), "minima": _para_decimal(ponto.get("low")),
            })

    historico.sort(key=lambda item: item["data"])
    return historico


def backfill_historico_cotacoes(ativo: Ativo, dias: int = 180) -> int:
    """
    Preenche de uma vez o histórico de cotações diárias (Cotacao) de um
    ativo, a partir do endpoint histórico da brapi.dev - chamado
    automaticamente ao registrar a primeira operação de um ativo novo (ver
    core.views.operacao_nova) e também disponível para ativos já existentes
    via o comando "python manage.py backfill_cotacoes".

    Não sobrescreve preço/variação de cotações que já existem (ex: a de hoje,
    gravada por atualizar_cotacao_diaria) - só grava as datas que ainda
    faltam. Retorna quantas cotações novas foram gravadas.

    Efeito colateral "auto-cura": os campos volume, máxima e mínima (ver
    Scanner Técnico, core.services.analisar_volume_precos, calcular_atr e
    calcular_suporte_resistencia) foram adicionados depois de várias
    cotações já existirem sem eles - como o histórico da brapi.dev já traz
    tudo isso de graça (mesma chamada, sem custo extra), aproveita pra
    completar o que estiver faltando em cotações JÁ existentes dentro do
    período buscado, sem tocar no preço/variação delas.
    """
    historico = buscar_historico_precos(ativo.ticker, dias=dias)
    if not historico:
        return 0

    datas_buscadas = [item["data"] for item in historico]
    existentes = {
        c.data: c for c in ativo.cotacoes.filter(data__in=datas_buscadas)
    }

    novas = [
        Cotacao(
            ativo=ativo,
            data=item["data"],
            preco_fechamento=item["fechamento"],
            volume=item.get("volume"),
            maxima=item.get("maxima"),
            minima=item.get("minima"),
        )
        for item in historico
        if item["data"] not in existentes
    ]
    if novas:
        Cotacao.objects.bulk_create(novas, ignore_conflicts=True)

    por_data = {item["data"]: item for item in historico}
    a_completar = []
    for data_, cotacao in existentes.items():
        item = por_data.get(data_)
        if not item:
            continue
        campos = []
        if cotacao.volume is None and item.get("volume") is not None:
            cotacao.volume = item["volume"]
            campos.append("volume")
        if cotacao.maxima is None and item.get("maxima") is not None:
            cotacao.maxima = item["maxima"]
            campos.append("maxima")
        if cotacao.minima is None and item.get("minima") is not None:
            cotacao.minima = item["minima"]
            campos.append("minima")
        if campos:
            a_completar.append((cotacao, campos))
    if a_completar:
        todos_os_campos = {campo for _, campos in a_completar for campo in campos}
        Cotacao.objects.bulk_update([c for c, _ in a_completar], list(todos_os_campos))

    return len(novas)


# --------------------------------------------------------------------------
# Benchmark de mercado (Ibovespa / CDI)
#
# Responde a pergunta mais importante para decidir se vale a pena continuar
# operando ativamente: "minha carteira está rendendo mais do que simplesmente
# deixar esse dinheiro no Ibovespa (renda variável passiva) ou no CDI (renda
# fixa)?" - sem isso, saber que a carteira está "com +8% de lucro" não diz
# muita coisa sozinho.
# --------------------------------------------------------------------------
IBOVESPA_TICKER = "^BVSP"
_BCB_SGS_CDI = 12  # código da série do CDI (taxa diária, % ao dia) no SGS do Banco Central


def buscar_historico_cdi(dias: int = 180) -> list[dict]:
    """
    Busca o histórico diário da taxa do CDI (série 12 do SGS/Banco Central) -
    cada registro é a taxa do próprio dia (% ao dia), não um índice
    acumulado, então o retorno de um período precisa compor (juros
    compostos) as taxas, não só comparar o primeiro e o último valor (ver
    _retorno_cdi_periodo). Retorna uma lista de {"data": date, "taxa_pct":
    Decimal} em ordem cronológica.

    Reaproveita BrapiError para o erro, ainda que a fonte aqui seja o Banco
    Central (não a brapi.dev) - mantém um único tipo de exceção pros
    chamadores (atualizar_benchmarks, comandos de management) tratarem.

    Usa o endpoint por intervalo de datas (`dados?dataInicial=...`), não o
    `dados/ultimos/N` - esse último retorna 400 Bad Request ("quantidade
    máxima de valores deve ser 20") pra série do CDI acima de 20 registros.
    """
    data_final = date.today()
    data_inicial = data_final - timedelta(days=max(dias, 1))
    url = f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{_BCB_SGS_CDI}/dados"
    params = {
        "formato": "json",
        "dataInicial": data_inicial.strftime("%d/%m/%Y"),
        "dataFinal": data_final.strftime("%d/%m/%Y"),
    }
    try:
        resposta = requests.get(url, params=params, timeout=30)
        resposta.raise_for_status()
        dados = resposta.json()
    except requests.RequestException as exc:
        raise BrapiError(_mensagem_falha_api(exc, "o histórico do CDI no Banco Central")) from exc

    historico = []
    for item in dados or []:
        data_str = item.get("data")
        valor_str = item.get("valor")
        if not data_str or valor_str is None:
            continue
        try:
            data_convertida = datetime.strptime(data_str, "%d/%m/%Y").date()
            taxa = Decimal(str(valor_str))
        except (ValueError, ArithmeticError):
            continue
        historico.append({"data": data_convertida, "taxa_pct": taxa})

    historico.sort(key=lambda item: item["data"])
    return historico


def _gravar_historico_indice(indice: str, pontos: list[tuple]) -> int:
    """Grava em CotacaoIndice os pontos (data, valor) cuja data ainda não existe para aquele índice."""
    if not pontos:
        return 0
    datas = [data for data, _ in pontos]
    existentes = set(
        CotacaoIndice.objects.filter(indice=indice, data__in=datas).values_list("data", flat=True)
    )
    novos = [
        CotacaoIndice(indice=indice, data=data, valor=valor)
        for data, valor in pontos
        if data not in existentes
    ]
    if novos:
        CotacaoIndice.objects.bulk_create(novos, ignore_conflicts=True)
    return len(novos)


def atualizar_benchmarks(dias: int = 180) -> dict:
    """
    Atualiza o histórico de benchmarks (Ibovespa e CDI) usado no comparativo
    de desempenho da carteira (ver calcular_comparativo_benchmark) - grava em
    CotacaoIndice as datas que ainda não existiam. Chamado pelo comando
    "atualizar_benchmarks" (agendável via cron/--loop, separado de
    "atualizar_cotacoes" de propósito - os benchmarks não dependem de nenhum
    usuário ter ativos cadastrados e não precisam da mesma frequência).

    Ibovespa e CDI são buscados de fontes independentes - uma fonte fora do
    ar não impede a atualização da outra. Retorna {"ibovespa": <quantas
    cotações novas>, "cdi": <quantas taxas novas>, "erros": [...]}.
    """
    resultado = {"ibovespa": 0, "cdi": 0, "erros": []}

    try:
        historico_ibov = buscar_historico_precos(IBOVESPA_TICKER, dias=dias)
        resultado["ibovespa"] = _gravar_historico_indice(
            CotacaoIndice.IBOVESPA, [(item["data"], item["fechamento"]) for item in historico_ibov]
        )
    except BrapiError as exc:
        resultado["erros"].append(str(exc))

    try:
        historico_cdi = buscar_historico_cdi(dias=dias)
        resultado["cdi"] = _gravar_historico_indice(
            CotacaoIndice.CDI, [(item["data"], item["taxa_pct"]) for item in historico_cdi]
        )
    except BrapiError as exc:
        resultado["erros"].append(str(exc))

    return resultado


def ultimo_valor_ibovespa() -> dict | None:
    """
    Valor do Ibovespa (pontos de fechamento) DE HOJE e a variação percentual
    em relação ao pregão anterior já salvo - vem do histórico gravado por
    atualizar_benchmarks (botão "Atualizar benchmarks (Ibovespa / CDI)" ou o
    comando "atualizar_benchmarks"), não do ciclo de cotações dos ativos, que
    não mexe em índices.

    None quando ainda não há nenhum valor salvo, OU quando o valor mais
    recente salvo é de um dia anterior a hoje - de propósito, pra não mostrar
    um número desatualizado como se fosse o de agora; nesse caso a tela deve
    sugerir rodar aquela atualização.
    """
    ultimos = list(
        CotacaoIndice.objects.filter(indice=CotacaoIndice.IBOVESPA).order_by("-data")[:2]
    )
    if not ultimos or ultimos[0].data != timezone.localdate():
        return None

    atual = ultimos[0]
    variacao_pct = None
    if len(ultimos) == 2 and ultimos[1].valor:
        variacao_pct = ((atual.valor - ultimos[1].valor) / ultimos[1].valor * 100).quantize(Decimal("0.01"))

    return {"data": atual.data, "valor": atual.valor, "variacao_pct": variacao_pct}


def _retorno_indice_periodo(indice: str, data_inicio: date, data_fim: date) -> Decimal | None:
    """Retorno % de um índice tipo Ibovespa (comparando o primeiro e o último valor do período)."""
    pontos = list(
        CotacaoIndice.objects.filter(indice=indice, data__gte=data_inicio, data__lte=data_fim)
        .order_by("data").values_list("valor", flat=True)
    )
    if len(pontos) < 2 or not pontos[0]:
        return None
    return ((pontos[-1] - pontos[0]) / pontos[0] * 100).quantize(Decimal("0.01"))


def _retorno_cdi_periodo(data_inicio: date, data_fim: date) -> Decimal | None:
    """
    Retorno acumulado (%) do CDI no período, compondo (juros compostos) as
    taxas diárias - cada registro de CotacaoIndice(indice=CDI) é a taxa do
    próprio dia, não um índice, então soma-los diretamente subestimaria o
    efeito de juros sobre juros.
    """
    taxas = list(
        CotacaoIndice.objects.filter(indice=CotacaoIndice.CDI, data__gte=data_inicio, data__lte=data_fim)
        .order_by("data").values_list("valor", flat=True)
    )
    if not taxas:
        return None
    fator = Decimal("1")
    for taxa in taxas:
        fator *= (1 + taxa / 100)
    return ((fator - 1) * 100).quantize(Decimal("0.01"))


def calcular_comparativo_benchmark(usuario, posicoes: list | None = None, dias: int = 90) -> dict | None:
    """
    Compara o retorno da carteira comprada (não considera reservas) com o
    retorno do Ibovespa e o retorno acumulado do CDI no mesmo período - ajuda
    a decidir se vale a pena continuar operando ativamente em vez de deixar o
    dinheiro num índice ou na renda fixa.

    O período vai da abertura da posição comprada mais antiga até hoje,
    limitado a `dias` corridos (carteiras antigas comparam só a janela mais
    recente, pra manter a comparação relevante ao momento atual). O retorno
    da carteira é uma aproximação: usa a cotação mais próxima do início do
    período e a mais recente de cada ativo, ponderada pela quantidade em
    carteira hoje - não reconstitui compras/vendas feitas dentro da janela.

    Retorna None quando não há posição comprada com histórico suficiente
    para montar a comparação, ou quando os benchmarks ainda não têm dados no
    período (ver atualizar_benchmarks).
    """
    if posicoes is None:
        posicoes = calcular_posicoes(usuario)

    compradas = [p for p in posicoes if not p.apenas_reservado and p.quantidade > 0]
    datas_abertura = [p.data_abertura for p in compradas if p.data_abertura is not None]
    if not datas_abertura:
        return None

    hoje = timezone.localdate()
    limite_periodo = hoje - timedelta(days=dias)
    data_inicio = max(min(datas_abertura), limite_periodo)
    if data_inicio >= hoje:
        return None

    valor_inicial_total = Decimal("0")
    valor_final_total = Decimal("0")
    for p in compradas:
        cotacao_inicial = p.ativo.cotacoes.filter(data__gte=data_inicio).order_by("data").first()
        cotacao_final = p.ativo.ultima_cotacao()
        if not cotacao_inicial or not cotacao_final:
            continue
        valor_inicial_total += cotacao_inicial.preco_fechamento * p.quantidade
        valor_final_total += cotacao_final.preco_fechamento * p.quantidade

    if not valor_inicial_total:
        return None

    retorno_carteira_pct = ((valor_final_total - valor_inicial_total) / valor_inicial_total * 100).quantize(
        Decimal("0.01")
    )
    retorno_ibovespa_pct = _retorno_indice_periodo(CotacaoIndice.IBOVESPA, data_inicio, hoje)
    retorno_cdi_pct = _retorno_cdi_periodo(data_inicio, hoje)

    return {
        "data_inicio": data_inicio,
        "data_fim": hoje,
        "dias_periodo": (hoje - data_inicio).days,
        "retorno_carteira_pct": retorno_carteira_pct,
        "retorno_ibovespa_pct": retorno_ibovespa_pct,
        "retorno_cdi_pct": retorno_cdi_pct,
        "carteira_bateu_ibovespa": (
            retorno_ibovespa_pct is not None and retorno_carteira_pct > retorno_ibovespa_pct
        ),
        "carteira_bateu_cdi": retorno_cdi_pct is not None and retorno_carteira_pct > retorno_cdi_pct,
    }


# --------------------------------------------------------------------------
# Cálculo de posições (compra/venda) e lucro/perda
# --------------------------------------------------------------------------
@dataclass
class Posicao:
    ativo: Ativo
    quantidade: int
    preco_medio: Decimal
    valor_investido: Decimal
    preco_atual: Decimal | None = None
    variacao_dia_pct: Decimal | None = None
    valor_atual: Decimal | None = None
    lucro_perda_valor: Decimal | None = None
    lucro_perda_pct: Decimal | None = None
    meta_lucro_pct: Decimal | None = None
    meta_perda_pct: Decimal | None = None
    tendencia: str | None = None
    data_abertura: date | None = None
    quantidade_reservada: int = 0
    reserva_do_robo: bool = False
    reserva_operacao_ids: list[int] = field(default_factory=list)
    lotes_compra: list = field(default_factory=list)  # lotes (Operacao COMPRA) com saldo > 0 que compõem a posição

    @property
    def situacao(self) -> str:
        if self.lucro_perda_pct is None:
            return "SEM_COTACAO"
        if self.lucro_perda_pct > 0:
            return "LUCRO"
        if self.lucro_perda_pct < 0:
            return "PERDA"
        return "NEUTRO"

    @property
    def apenas_reservado(self) -> bool:
        """True quando o ativo não está de fato em carteira - só há reserva(s) de intenção de compra."""
        return self.quantidade == 0 and self.quantidade_reservada > 0

    @property
    def dias_desde_compra(self) -> int | None:
        """Dias corridos desde que a posição foi aberta (primeira compra após ela estar zerada)."""
        if self.data_abertura is None:
            return None
        return (timezone.localdate() - self.data_abertura).days

    @property
    def lucro_perda_por_dia_valor(self) -> Decimal | None:
        """
        Lucro/perda atual (R$, ainda não realizado) dividido pelos dias que a
        posição está em carteira - mede a "velocidade" do resultado (R$/dia),
        igual à mesma ideia já usada em Operacao.lucro_perda_por_dia para
        vendas, mas aqui sobre a posição aberta. None quando comprado hoje (0
        dias, não dá pra ratear) ou sem lucro/perda apurado (sem cotação).
        """
        dias = self.dias_desde_compra
        if self.lucro_perda_valor is None or not dias:
            return None
        return (self.lucro_perda_valor / dias).quantize(Decimal("0.01"))

    @property
    def lucro_perda_por_dia_pct(self) -> Decimal | None:
        """Lucro/perda atual (%) dividido pelos dias em carteira - mesma ideia de lucro_perda_por_dia_valor, em percentual."""
        dias = self.dias_desde_compra
        if self.lucro_perda_pct is None or not dias:
            return None
        return (self.lucro_perda_pct / dias).quantize(Decimal("0.01"))

    @property
    def meta_lucro_atingida(self) -> bool:
        """
        True quando o lucro/perda atual da posição já atingiu a meta de lucro
        (a própria, ou a padrão do sistema quando não há uma definida) - usado
        para destacar na grid que é hora de considerar a venda.
        """
        if self.lucro_perda_pct is None or self.apenas_reservado:
            return False
        # a meta de lucro é sempre um ganho (valor positivo) - normaliza caso
        # tenha sido salva como negativa (dado antigo, de antes da validação)
        meta = abs(
            self.meta_lucro_pct
            if self.meta_lucro_pct is not None
            else Decimal(str(settings.META_LUCRO_PADRAO))
        )
        return self.lucro_perda_pct >= meta

    @property
    def preco_alvo(self) -> Decimal | None:
        """
        Preço (R$) em que a posição atinge a meta de lucro: preço médio de
        compra acrescido do percentual de meta cadastrado na compra (ou do
        padrão do sistema quando não há um definido - mesma regra de
        meta_lucro_atingida). None para posições só de reserva.
        """
        if self.apenas_reservado or not self.preco_medio:
            return None
        meta = abs(
            self.meta_lucro_pct
            if self.meta_lucro_pct is not None
            else Decimal(str(settings.META_LUCRO_PADRAO))
        )
        return (self.preco_medio * (1 + meta / 100)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    @property
    def variacao_pct_reserva(self) -> Decimal | None:
        """
        Para posições só de reserva (sem compra real): variação percentual do
        preço atual em relação ao preço pretendido na reserva - usada para
        saber se a baixa desejada para comprar já foi atingida.
        """
        if not self.apenas_reservado or self.preco_atual is None or not self.preco_medio:
            return None
        return ((self.preco_atual - self.preco_medio) / self.preco_medio * 100).quantize(Decimal("0.01"))

    @property
    def meta_compra_atingida(self) -> bool:
        """
        True quando uma posição só de reserva já caiu até (ou além) a meta de
        baixa definida (a própria, ou a padrão do sistema quando não há uma
        definida) - usado para destacar na grid que é hora de considerar a compra.
        """
        variacao = self.variacao_pct_reserva
        if variacao is None:
            return False
        meta_bruta = (
            self.meta_perda_pct
            if self.meta_perda_pct is not None
            else Decimal(str(settings.META_PERDA_PADRAO))
        )
        # a meta de perda é sempre uma queda (valor negativo) - normaliza caso
        # tenha sido salva como positiva (dado antigo, de antes da validação)
        meta = -abs(meta_bruta)
        return variacao <= meta


def calcular_posicoes(usuario) -> list[Posicao]:
    """
    Consolida os lotes de compra do usuário em posições por ativo (preço
    médio de compra do saldo não vendido, quantidade líquida em carteira) e
    calcula o lucro/perda percentual com base na última cotação disponível.

    Cada Operacao tipo=COMPRA é um lote com seu próprio saldo (quantidade -
    quantidade_vendida); só o saldo ainda não vendido entra na posição.
    """
    operacoes = (
        Operacao.objects.filter(usuario=usuario)
        .select_related("ativo")
        .order_by("data_operacao", "criado_em")
    )

    agregados = defaultdict(lambda: {
        "quantidade": 0,
        "custo_total": Decimal("0"),
        "meta_lucro_pct": None,
        "meta_perda_pct": None,
        "data_abertura": None,
        "quantidade_reservada": 0,
        "custo_reservado": Decimal("0"),
        "reserva_do_robo": False,
        "reserva_operacao_ids": [],
        "lotes_compra": [],
    })

    for op in operacoes:
        item = agregados[op.ativo_id]
        if op.tipo == Operacao.COMPRA:
            saldo_lote = op.saldo
            if saldo_lote > 0:
                # a posição é composta pelo(s) lote(s) mais antigo(s) que ainda
                # têm saldo - "há quanto tempo está na carteira" olha pra eles
                if item["data_abertura"] is None or op.data_operacao < item["data_abertura"]:
                    item["data_abertura"] = op.data_operacao
                item["quantidade"] += saldo_lote
                item["custo_total"] += saldo_lote * op.preco_unitario
                item["lotes_compra"].append(op)
        else:  # RESERVAR: intenção de compra futura (no máximo 1 unidade), não altera a posição real
            item["quantidade_reservada"] += op.quantidade
            item["custo_reservado"] += op.valor_total
            item["reserva_operacao_ids"].append(op.id)
            if "robô consultor" in (op.observacao or "").lower():
                item["reserva_do_robo"] = True

        # a meta mais recente informada pelo usuário prevalece
        if op.meta_lucro_pct is not None:
            item["meta_lucro_pct"] = op.meta_lucro_pct
        if op.meta_perda_pct is not None:
            item["meta_perda_pct"] = op.meta_perda_pct
        item["ativo"] = op.ativo

    posicoes = []
    for ativo_id, item in agregados.items():
        if item["quantidade"] <= 0 and item["quantidade_reservada"] <= 0:
            continue  # posição zerada (tudo vendido) e sem nenhuma reserva - nada a mostrar

        ativo = item["ativo"]

        if item["quantidade"] > 0:
            preco_medio = (item["custo_total"] / item["quantidade"]).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP
            )
            valor_investido = (preco_medio * item["quantidade"]).quantize(Decimal("0.01"))
        else:
            # só há reserva(s): usa o preço pretendido na reserva como referência,
            # sem valor investido de verdade (nenhum capital foi comprometido ainda)
            preco_medio = (item["custo_reservado"] / item["quantidade_reservada"]).quantize(
                Decimal("0.01"), rounding=ROUND_HALF_UP
            )
            valor_investido = Decimal("0.00")

        posicao = Posicao(
            ativo=ativo,
            quantidade=item["quantidade"],
            preco_medio=preco_medio,
            valor_investido=valor_investido,
            meta_lucro_pct=item["meta_lucro_pct"],
            meta_perda_pct=item["meta_perda_pct"],
            data_abertura=item["data_abertura"],
            quantidade_reservada=item["quantidade_reservada"],
            reserva_do_robo=item["reserva_do_robo"],
            reserva_operacao_ids=item["reserva_operacao_ids"],
            lotes_compra=item["lotes_compra"],
        )

        ultima_cotacao = ativo.ultima_cotacao()
        if ultima_cotacao:
            posicao.preco_atual = ultima_cotacao.preco_fechamento
            posicao.variacao_dia_pct = ultima_cotacao.variacao_dia_pct
            posicao.tendencia = analisar_tendencia(ativo)

        if ultima_cotacao and not posicao.apenas_reservado:
            posicao.valor_atual = (posicao.preco_atual * posicao.quantidade).quantize(Decimal("0.01"))
            posicao.lucro_perda_valor = (posicao.valor_atual - posicao.valor_investido).quantize(Decimal("0.01"))
            if posicao.valor_investido:
                posicao.lucro_perda_pct = (
                    (posicao.lucro_perda_valor / posicao.valor_investido) * 100
                ).quantize(Decimal("0.01"))

        posicoes.append(posicao)

    # menos dias em carteira primeiro; posições só com reserva (sem compra
    # real, sem "dias em carteira" pra comparar) ficam por último
    posicoes.sort(key=lambda p: (
        p.dias_desde_compra is None,
        p.dias_desde_compra if p.dias_desde_compra is not None else 0,
        p.ativo.ticker,
    ))
    return posicoes


# --------------------------------------------------------------------------
# Histórico de atualizações da carteira (snapshot a cada atualização de
# cotações) - ver core.models.RegistroAtualizacaoCarteira e a tela Histórico
# de Atualizações.
# --------------------------------------------------------------------------
def registrar_atualizacao_carteira(usuario, posicoes: list[Posicao] | None = None) -> RegistroAtualizacaoCarteira:
    """
    Grava um retrato (snapshot) dos totais da carteira comprada do usuário
    (total de ativos, valor investido, valor atual, lucro/perda) - chamado
    toda vez que as cotações são atualizadas, seja pelo botão manual
    "Atualizar cotações agora" (core.views.atualizar_cotacoes_agora) seja
    pelo comando "atualizar_cotacoes" (inclusive em --loop). Mesma conta
    usada nos cartões de resumo do Painel e de Posições em Carteira.
    """
    if posicoes is None:
        posicoes = calcular_posicoes(usuario)

    compradas = [p for p in posicoes if not p.apenas_reservado]
    valor_investido = sum((p.valor_investido for p in compradas), Decimal("0"))
    valor_atual = sum((p.valor_atual for p in compradas if p.valor_atual is not None), Decimal("0"))
    lucro_perda = valor_atual - valor_investido
    lucro_perda_pct = (
        (lucro_perda / valor_investido * 100).quantize(Decimal("0.01")) if valor_investido else None
    )

    return RegistroAtualizacaoCarteira.objects.create(
        usuario=usuario,
        total_ativos=len(compradas),
        valor_investido=valor_investido,
        valor_atual=valor_atual,
        lucro_perda=lucro_perda,
        lucro_perda_pct=lucro_perda_pct,
    )


# --------------------------------------------------------------------------
# Registros de cotações dos ativos comprados (log a cada atualização) - ver
# core.models.RegistroCotacao, a grid em Detalhes de Cotações e o gráfico
# em popup (compra x alvo) das grids de ativos comprados.
# --------------------------------------------------------------------------
def registrar_cotacao(usuario, ativo: Ativo, preco, origem: str = RegistroCotacao.AUTOMATICA) -> RegistroCotacao:
    """Grava um registro de cotação (ativo, data, hora, valor) com a data/hora local de agora."""
    agora = timezone.localtime()
    return RegistroCotacao.objects.create(
        usuario=usuario,
        ativo=ativo,
        data=agora.date(),
        hora=agora.time().replace(microsecond=0),
        preco=Decimal(str(preco)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        origem=origem,
    )


def registrar_cotacoes_compradas(usuario, origem: str = RegistroCotacao.AUTOMATICA) -> int:
    """
    Grava um registro de cotação para cada ativo que o usuário tem em
    carteira (saldo comprado > 0) - chamado logo depois de atualizar as
    cotações, pelo ciclo automático (executar_ciclo_atualizacao_cotacoes) e
    pelo botão "Atualizar cotações agora". Usa só a cotação de HOJE: se a
    atualização do ativo falhou e a última cotação é de outro dia, não
    registra um valor velho como se fosse de agora. Retorna quantos gravou.
    """
    hoje = timezone.localdate()
    total = 0
    for ativo in Ativo.objects.filter(id__in=ativos_distintos_comprados(usuario)):
        cotacao = ativo.cotacoes.filter(data=hoje).first()
        if cotacao is None:
            continue
        registrar_cotacao(usuario, ativo, cotacao.preco_fechamento, origem)
        total += 1
    return total


def excluir_registros_cotacao_ativo(usuario, ativo: Ativo) -> int:
    """Exclui todos os registros de cotação do usuário para o ativo (pergunta feita após a venda)."""
    excluidos, _ = RegistroCotacao.objects.filter(usuario=usuario, ativo=ativo).delete()
    return excluidos


def construir_grafico_registros_cotacao(
    registros_cronologico: list[RegistroCotacao], preco_compra: Decimal | None, preco_alvo: Decimal | None,
    largura: int = 640, altura: int = 260, padding: int = 30,
) -> dict | None:
    """
    Monta os dados (SVG) do gráfico em popup dos registros de cotação de um
    ativo, com duas linhas horizontais de referência: o preço de compra e o
    preço alvo (meta de lucro). A escala vertical inclui as duas
    referências, pra elas aparecerem sempre dentro do gráfico mesmo quando
    a cotação ainda está longe do alvo. Recebe os registros em ordem
    cronológica (mais antigo primeiro). None quando não há registros.
    """
    if not registros_cronologico:
        return None

    precos = [float(r.preco) for r in registros_cronologico]
    referencias = [float(v) for v in (preco_compra, preco_alvo) if v is not None]
    valor_min = min(precos + referencias)
    valor_max = max(precos + referencias)
    margem = (valor_max - valor_min) * 0.08 or max(valor_max * 0.01, 0.01)
    valor_min -= margem
    valor_max += margem
    faixa = valor_max - valor_min

    plot_largura = largura - (padding * 2)
    plot_altura = altura - (padding * 2)

    def y_de(valor: float) -> float:
        return round(padding + (1 - (valor - valor_min) / faixa) * plot_altura, 2)

    def real(valor: float) -> str:
        return f"R$ {valor:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")

    total = len(registros_cronologico)
    pontos = []
    comandos_path = []
    for i, (registro, preco) in enumerate(zip(registros_cronologico, precos)):
        x = round(padding + (i * plot_largura / (total - 1) if total > 1 else plot_largura / 2), 2)
        y = y_de(preco)
        comandos_path.append(f"{'M' if i == 0 else 'L'}{x} {y}")
        pontos.append({
            "x": x,
            "y": y,
            "data_label": f"{registro.data:%d/%m/%Y} {registro.hora:%H:%M}",
            "preco_label": real(preco),
            "variacao_label": (
                f"{(preco - float(preco_compra)) / float(preco_compra) * 100:+.2f}% s/ compra".replace(".", ",")
                if preco_compra else None
            ),
            "variacao_positiva": bool(preco_compra) and preco >= float(preco_compra),
        })

    def pct_de(y: float) -> float:
        # Posição vertical (% da altura do SVG) dos rótulos HTML à esquerda do gráfico.
        return round(y / altura * 100, 2)

    rotulos_esquerda = []
    if preco_alvo is not None:
        rotulos_esquerda.append({"classe": "referencia-alvo", "titulo": "Alvo",
                                 "valor": real(float(preco_alvo)), "topo_pct": pct_de(y_de(float(preco_alvo)))})
    if preco_compra is not None:
        rotulos_esquerda.append({"classe": "referencia-compra", "titulo": "Compra",
                                 "valor": real(float(preco_compra)), "topo_pct": pct_de(y_de(float(preco_compra)))})
    rotulos_esquerda.append({
        "classe": "linha-alta" if preco_compra is None or precos[-1] >= float(preco_compra) else "linha-baixa",
        "titulo": "Última", "valor": real(precos[-1]), "topo_pct": pct_de(pontos[-1]["y"]),
    })
    # Afasta rótulos de linhas muito próximas pra não sobrepor o texto (~32px num gráfico de 300px).
    distancia_min = 11.0
    rotulos_esquerda.sort(key=lambda r: r["topo_pct"])
    for anterior, atual in zip(rotulos_esquerda, rotulos_esquerda[1:]):
        if atual["topo_pct"] - anterior["topo_pct"] < distancia_min:
            atual["topo_pct"] = round(anterior["topo_pct"] + distancia_min, 2)
    excesso = rotulos_esquerda[-1]["topo_pct"] - 96
    if excesso > 0:
        for r in rotulos_esquerda:
            r["topo_pct"] = round(r["topo_pct"] - excesso, 2)

    return {
        "id_svg": "pontos-registros-cotacao",
        "largura": largura,
        "altura": altura,
        "rotulos_esquerda": rotulos_esquerda,
        "x_inicio": padding,
        "x_fim": largura - padding,
        "y_topo": padding,
        "y_base": altura - padding,
        "path_d": " ".join(comandos_path),
        "pontos": pontos,
        "y_compra": y_de(float(preco_compra)) if preco_compra is not None else None,
        "y_alvo": y_de(float(preco_alvo)) if preco_alvo is not None else None,
        "compra_label": real(float(preco_compra)) if preco_compra is not None else None,
        "alvo_label": real(float(preco_alvo)) if preco_alvo is not None else None,
        "faixa_min_label": real(valor_min),
        "faixa_max_label": real(valor_max),
        "tendencia_alta": preco_compra is None or precos[-1] >= float(preco_compra),
    }


def excluir_registros_atualizacao_antigos(usuario, dias: int) -> int:
    """
    Exclui os registros de atualização da carteira do usuário com mais de
    `dias` dias - usado no formulário "Excluir por dias" da tela Histórico
    de Atualizações, pra não deixar a tabela crescer indefinidamente quando
    o comando "atualizar_cotacoes --loop" fica rodando o dia inteiro.
    """
    limite = timezone.now() - timedelta(days=dias)
    excluidos, _ = RegistroAtualizacaoCarteira.objects.filter(
        usuario=usuario, criado_em__lt=limite,
    ).delete()
    return excluidos


def excluir_registros_atualizacao_por_periodo(usuario, data_inicio: date, data_fim: date) -> int:
    """
    Exclui os registros de atualização da carteira do usuário com data/hora
    entre `data_inicio` e `data_fim` (ambas incluídas) - usado no formulário
    "Excluir por período" da tela Histórico de Atualizações, alternativa ao
    "Excluir por dias" quando o usuário quer limpar uma janela específica em
    vez de "tudo mais antigo que N dias".
    """
    inicio = timezone.make_aware(datetime.combine(data_inicio, datetime.min.time()))
    fim = timezone.make_aware(datetime.combine(data_fim, datetime.max.time()))
    excluidos, _ = RegistroAtualizacaoCarteira.objects.filter(
        usuario=usuario, criado_em__gte=inicio, criado_em__lte=fim,
    ).delete()
    return excluidos


def calcular_variacoes_historico(registros_desc: list[RegistroAtualizacaoCarteira]) -> list[dict]:
    """
    Para cada registro do histórico de atualizações (ordenados do mais
    recente para o mais antigo, mesma ordem da tela), calcula a variação do
    valor atual da carteira (R$ e %) em relação ao registro imediatamente
    anterior no tempo - "quanto mudou desde a atualização passada". O
    registro mais antigo da lista fica sem variação (não há um anterior pra
    comparar), assim como quando o valor atual anterior era zero (não dá
    pra calcular percentual sobre zero).

    Retorna uma lista de dicts na mesma ordem de entrada, cada um com
    "registro", "variacao_valor" e "variacao_pct".
    """
    resultado = []
    for i, registro in enumerate(registros_desc):
        anterior = registros_desc[i + 1] if i + 1 < len(registros_desc) else None
        variacao_valor = None
        variacao_pct = None
        if anterior is not None:
            variacao_valor = (registro.valor_atual - anterior.valor_atual).quantize(Decimal("0.01"))
            if anterior.valor_atual:
                variacao_pct = (variacao_valor / anterior.valor_atual * 100).quantize(Decimal("0.01"))
        resultado.append({
            "registro": registro,
            "variacao_valor": variacao_valor,
            "variacao_pct": variacao_pct,
        })
    return resultado


def construir_grafico_atualizacoes_dia(
    registros_dia: list[RegistroAtualizacaoCarteira], largura: int = 640, altura: int = 200, padding: int = 40,
    id_svg: str = "pontos-atualizacoes-dia",
) -> dict | None:
    """
    Monta os dados (SVG) do gráfico "Variações do dia" da tela Histórico de
    Atualizações: evolução do lucro/perda (%) da carteira ao longo dos
    registros de hoje, na ordem em que foram gravados. Recebe os registros
    já filtrados pelo dia atual, em ordem cronológica (mais antigo primeiro).

    Funciona com qualquer objeto que tenha os campos criado_em/valor_atual/
    lucro_perda/lucro_perda_pct - RegistroAtualizacaoCarteira (ações) ou
    RegistroAtualizacaoCarteiraCripto, sem precisar duplicar esta função.
    `id_svg` deve ser único na página quando os dois gráficos (ações e
    cripto) aparecem juntos (ver core.views.dashboard), senão o crosshair/
    tooltip (grafico.js) quebra por causa do id de SVG duplicado.

    Retorna None quando não há pelo menos 2 registros hoje (não dá pra
    traçar uma linha com um ponto só).
    """
    pontos_validos = [r for r in registros_dia if r.lucro_perda_pct is not None]
    if len(pontos_validos) < 2:
        return None

    valores = [float(r.lucro_perda_pct) for r in pontos_validos]
    valor_min, valor_max = min(valores), max(valores)
    faixa = (valor_max - valor_min) or 1.0

    plot_largura = largura - (padding * 2)
    plot_altura = altura - (padding * 2)
    passo_x = plot_largura / (len(pontos_validos) - 1)

    pontos = []
    comandos_path = []
    for i, (registro, valor) in enumerate(zip(pontos_validos, valores)):
        x = round(padding + i * passo_x, 2)
        y = round(padding + (1 - (valor - valor_min) / faixa) * plot_altura, 2)
        comandos_path.append(f"{'M' if i == 0 else 'L'}{x} {y}")

        pontos.append({
            "x": x,
            "y": y,
            "hora_label": timezone.localtime(registro.criado_em).strftime("%H:%M"),
            "valor_atual_label": f"R$ {registro.valor_atual:,.2f}".replace(",", "#").replace(".", ",").replace("#", "."),
            "lucro_perda_label": f"R$ {registro.lucro_perda:,.2f}".replace(",", "#").replace(".", ",").replace("#", "."),
            "lucro_perda_pct_label": f"{valor:.2f}%".replace(".", ","),
            "positivo": valor >= 0,
        })

    return {
        "id_svg": id_svg,
        "largura": largura,
        "altura": altura,
        "x_inicio": padding,
        "x_fim": largura - padding,
        "y_topo": padding,
        "y_base": altura - padding,
        "path_d": " ".join(comandos_path),
        "pontos": pontos,
        "faixa_min_label": f"{valor_min:.2f}%".replace(".", ","),
        "faixa_max_label": f"{valor_max:.2f}%".replace(".", ","),
        "tendencia_alta": valores[-1] >= valores[0],
    }


COLUNAS_HISTORICO_ATUALIZACOES = [
    "Data/hora", "Total de ativos", "Valor investido (R$)", "Valor atual (R$)",
    "Lucro/Perda (R$)", "Lucro/Perda (%)",
]


def _linha_historico_atualizacao(registro: RegistroAtualizacaoCarteira) -> list:
    return [
        timezone.localtime(registro.criado_em),
        registro.total_ativos,
        float(registro.valor_investido),
        float(registro.valor_atual),
        float(registro.lucro_perda),
        float(registro.lucro_perda_pct) if registro.lucro_perda_pct is not None else None,
    ]


def gerar_excel_historico_atualizacoes(registros: list[RegistroAtualizacaoCarteira]) -> bytes:
    """Gera uma planilha .xlsx com o histórico de atualizações da carteira do usuário, pronta para download."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Histórico de atualizações"

    ws.append(["Histórico de Atualizações"])
    ws.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])
    ws.append(COLUNAS_HISTORICO_ATUALIZACOES)
    for celula in ws[4]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D1526")
        celula.alignment = Alignment(horizontal="center")

    for registro in registros:
        linha = _linha_historico_atualizacao(registro)
        linha[0] = linha[0].strftime("%d/%m/%Y %H:%M")
        ws.append(linha)

    for indice in range(1, len(COLUNAS_HISTORICO_ATUALIZACOES) + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 19
    ws.freeze_panes = "A5"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_historico_atualizacoes(registros: list[RegistroAtualizacaoCarteira], nome_usuario: str) -> bytes:
    """Gera um PDF com o histórico de atualizações da carteira do usuário, pronto para download."""
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()

    elementos = [
        Paragraph("BolsaTrader - Histórico de Atualizações", estilos["Title"]),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    def fmt(valor, sufixo="", quando_none="—"):
        return f"{valor:.2f}{sufixo}" if valor is not None else quando_none

    if not registros:
        elementos.append(Paragraph("Nenhum registro de atualização até o momento.", estilos["Normal"]))
    else:
        dados = [COLUNAS_HISTORICO_ATUALIZACOES]
        for registro in registros:
            linha = _linha_historico_atualizacao(registro)
            dados.append([
                linha[0].strftime("%d/%m/%Y %H:%M"), str(linha[1]), fmt(linha[2]), fmt(linha[3]),
                fmt(linha[4]), fmt(linha[5], "%"),
            ])

        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela)

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Concentração por setor
# --------------------------------------------------------------------------
def calcular_concentracao_setor(posicoes: list) -> list[dict]:
    """
    Agrupa o valor atual das posições compradas (reservas não entram) por
    setor do ativo (ver Ativo.setor, preenchido em atualizar_cotacao_diaria)
    e calcula o percentual de cada setor sobre o valor total da carteira -
    ativos sem setor conhecido entram em "Sem setor". Ordenado do maior para
    o menor percentual.

    O limite de ativos diferentes em carteira (MAX_ATIVOS_EM_CARTEIRA) não
    enxerga isso: dá pra estar "diversificado" em 10 tickers e todos do mesmo
    setor - esta função é o que permite avisar sobre esse tipo de risco.
    """
    valores_por_setor = defaultdict(lambda: Decimal("0"))
    valor_total = Decimal("0")
    for p in posicoes:
        if p.apenas_reservado or p.valor_atual is None:
            continue
        setor = p.ativo.setor or "Sem setor"
        valores_por_setor[setor] += p.valor_atual
        valor_total += p.valor_atual

    if not valor_total:
        return []

    concentracao = [
        {"setor": setor, "valor": valor, "percentual": (valor / valor_total * 100).quantize(Decimal("0.01"))}
        for setor, valor in valores_por_setor.items()
    ]
    concentracao.sort(key=lambda item: item["percentual"], reverse=True)
    return concentracao


# --------------------------------------------------------------------------
# Métricas de risco: volatilidade, drawdown e concentração por ativo
#
# Complementam o lucro/perda: duas carteiras com o mesmo retorno podem ter
# risco bem diferente - uma diversificada em vários ativos estáveis, outra
# concentrada numa única posição muito volátil. O percentual de lucro/perda
# sozinho não distingue os dois casos.
# --------------------------------------------------------------------------
DIAS_PREGAO_POR_ANO = 252
LIMIAR_CONCENTRACAO_ALERTA_PCT = Decimal("40")


def _volatilidade_e_drawdown_precos(precos_cronologico: list[float]) -> dict:
    """
    Núcleo puro (sem acesso ao banco) do cálculo de volatilidade anualizada e
    drawdown máximo a partir de uma lista de preços em ordem cronológica
    (mais antigo primeiro). Extraído de calcular_metricas_risco para ser
    reaproveitado pelo Scanner Técnico (ver analisar_volatilidade_precos),
    sem duplicar a fórmula. Ambos ficam None sem histórico suficiente
    (< 3 preços).

    Volatilidade: desvio padrão dos retornos diários no período, anualizado
    (× √252, dias de pregão por ano) - maior valor = preço oscila mais no dia
    a dia. Drawdown máximo: maior queda percentual do topo ao fundo dentro da
    mesma janela.
    """
    if len(precos_cronologico) < 3:
        return {"volatilidade_pct": None, "drawdown_pct": None}

    retornos = [
        (atual - anterior) / anterior
        for anterior, atual in zip(precos_cronologico, precos_cronologico[1:])
        if anterior
    ]
    volatilidade_pct = (
        round(statistics.pstdev(retornos) * (DIAS_PREGAO_POR_ANO ** 0.5) * 100, 2)
        if len(retornos) >= 2
        else None
    )

    pico = precos_cronologico[0]
    maior_queda = 0.0
    for preco in precos_cronologico:
        pico = max(pico, preco)
        if pico:
            maior_queda = min(maior_queda, (preco - pico) / pico)
    drawdown_pct = round(maior_queda * 100, 2)

    return {"volatilidade_pct": volatilidade_pct, "drawdown_pct": drawdown_pct}


def calcular_metricas_risco(posicoes: list, dias_janela: int = 30) -> dict:
    """
    Para cada posição comprada, calcula a volatilidade anualizada e o
    drawdown máximo dos últimos `dias_janela` pregões, mais o quanto aquela
    posição representa do valor atual total da carteira - e, a partir disso,
    a maior concentração individual e um índice de Herfindahl simples (soma
    dos quadrados das participações: quanto mais perto de 1, mais concentrada
    a carteira; quanto mais perto de 1/nº de ativos, mais diversificada).

    Volatilidade: desvio padrão dos retornos diários no período, anualizado
    (× √252, dias de pregão por ano) - maior valor = preço oscila mais no dia
    a dia. Drawdown máximo: maior queda percentual do topo ao fundo dentro da
    mesma janela. Ambos ficam None sem histórico suficiente (< 3 pregões).
    """
    compradas = [p for p in posicoes if not p.apenas_reservado and p.quantidade > 0]
    valor_total = sum((p.valor_atual for p in compradas if p.valor_atual is not None), Decimal("0"))

    por_ativo = []
    for p in compradas:
        historico_desc = list(
            p.ativo.cotacoes.order_by("-data").values_list("preco_fechamento", flat=True)[:dias_janela]
        )
        precos = [float(v) for v in reversed(historico_desc)]  # cronológico: mais antigo -> mais recente

        metricas = _volatilidade_e_drawdown_precos(precos)
        volatilidade_pct = metricas["volatilidade_pct"]
        drawdown_pct = metricas["drawdown_pct"]

        percentual_carteira = (
            (p.valor_atual / valor_total * 100).quantize(Decimal("0.01"))
            if p.valor_atual is not None and valor_total
            else None
        )
        por_ativo.append({
            "ativo": p.ativo,
            "volatilidade_pct": volatilidade_pct,
            "drawdown_pct": drawdown_pct,
            "percentual_carteira": percentual_carteira,
        })

    por_ativo.sort(key=lambda item: item["percentual_carteira"] or Decimal("0"), reverse=True)
    maior_posicao = por_ativo[0] if por_ativo else None
    indice_herfindahl = sum((float(item["percentual_carteira"] or 0) / 100) ** 2 for item in por_ativo)

    return {
        "por_ativo": por_ativo,
        "maior_posicao_ativo": maior_posicao["ativo"] if maior_posicao else None,
        "maior_posicao_pct": maior_posicao["percentual_carteira"] if maior_posicao else None,
        "concentracao_alta": bool(
            maior_posicao and maior_posicao["percentual_carteira"] is not None
            and maior_posicao["percentual_carteira"] >= LIMIAR_CONCENTRACAO_ALERTA_PCT
        ),
        "indice_herfindahl": round(indice_herfindahl, 3),
    }


# --------------------------------------------------------------------------
# Correlação e teste de estresse da carteira (aba Indicadores de risco)
#
# Diversificar não é ter muitos tickers: oito ações que sobem e caem juntas
# protegem tanto quanto uma. A correlação mostra isso; o beta contra o
# Ibovespa estima quanto a carteira tende a cair numa queda do mercado. Tudo
# a partir das cotações já salvas - nenhuma chamada de API.
# --------------------------------------------------------------------------
CORRELACAO_JANELA_PREGOES = 120
CORRELACAO_ALTA = 0.7
CORRELACAO_MIN_RETORNOS = 20
CORRELACAO_MAX_ATIVOS_MATRIZ = 10
ESTRESSE_JANELA_PREGOES = 250
ESTRESSE_QUEDAS_IBOV_PCT = (10, 20, 30)


def _retornos_pareados(serie_a: dict, serie_b: dict) -> tuple[list[float], list[float]]:
    """Retornos diários de duas séries {data: preço} só nas datas que as duas têm (pregão a pregão)."""
    datas = sorted(set(serie_a) & set(serie_b))
    ret_a, ret_b = [], []
    for anterior, atual in zip(datas, datas[1:]):
        if serie_a[anterior] and serie_b[anterior]:
            ret_a.append(serie_a[atual] / serie_a[anterior] - 1)
            ret_b.append(serie_b[atual] / serie_b[anterior] - 1)
    return ret_a, ret_b


def _correlacao(xs: list[float], ys: list[float]) -> float | None:
    if len(xs) < CORRELACAO_MIN_RETORNOS:
        return None
    try:
        return statistics.correlation(xs, ys)
    except statistics.StatisticsError:  # série constante
        return None


def _beta(retornos_ativo: list[float], retornos_mercado: list[float]) -> float | None:
    if len(retornos_ativo) < CORRELACAO_MIN_RETORNOS:
        return None
    variancia = statistics.pvariance(retornos_mercado)
    if not variancia:
        return None
    media_a, media_m = statistics.fmean(retornos_ativo), statistics.fmean(retornos_mercado)
    covariancia = statistics.fmean((a - media_a) * (m - media_m) for a, m in zip(retornos_ativo, retornos_mercado))
    return covariancia / variancia


def _fundo_correlacao(valor: float | None) -> str:
    """Cor de fundo da célula da matriz: vermelho = andam juntos, azul = em sentidos opostos, mais forte quanto maior."""
    if valor is None:
        return "transparent"
    r, g, b = (230, 103, 103) if valor >= 0 else (57, 135, 229)
    return f"rgba({r}, {g}, {b}, {min(abs(valor), 1) * 0.75:.2f})"


def _series_precos(ativos: list, dias: int) -> dict:
    """{ticker: {data: preço}} dos últimos `dias` pregões de cada ativo."""
    series = {}
    for ativo in ativos:
        linhas = ativo.cotacoes.order_by("-data").values_list("data", "preco_fechamento")[:dias]
        series[ativo.ticker] = {data: float(preco) for data, preco in linhas}
    return series


def calcular_correlacao_e_estresse(posicoes: list) -> dict:
    """
    Correlação entre os ativos comprados (retornos diários dos últimos
    CORRELACAO_JANELA_PREGOES pregões, cada par nas datas que os dois têm),
    beta de cada um contra o Ibovespa, perda estimada da carteira em quedas
    do Ibovespa de ESTRESSE_QUEDAS_IBOV_PCT e o pior dia/semana/mês que a
    carteira ATUAL (mesmas quantidades) teria tido no último ano.
    """
    compradas = [p for p in posicoes if not p.apenas_reservado and p.quantidade > 0 and p.valor_atual]
    compradas.sort(key=lambda p: p.valor_atual, reverse=True)
    valor_total = float(sum(p.valor_atual for p in compradas))
    vazio = {"correlacao": None, "estresse": None}
    if not compradas or not valor_total:
        return vazio

    series = _series_precos([p.ativo for p in compradas], ESTRESSE_JANELA_PREGOES)
    tickers = [p.ativo.ticker for p in compradas]

    # ---- correlação (janela menor: o que importa é o comportamento recente)
    series_corr = {
        t: dict(sorted(s.items())[-CORRELACAO_JANELA_PREGOES:]) for t, s in series.items()
    }
    pares = []
    matriz_valores = {}
    for i, a in enumerate(tickers):
        for b in tickers[i + 1:]:
            c = _correlacao(*_retornos_pareados(series_corr[a], series_corr[b]))
            matriz_valores[(a, b)] = matriz_valores[(b, a)] = c
            if c is not None:
                pares.append({"a": a, "b": b, "valor": round(c, 2)})
    pares.sort(key=lambda par: par["valor"], reverse=True)

    correlacao = None
    if pares:
        na_matriz = tickers[:CORRELACAO_MAX_ATIVOS_MATRIZ]
        for a in na_matriz:
            matriz_valores[(a, a)] = 1.0
        correlacao = {
            "tickers": na_matriz,
            "linhas": [
                {
                    "ticker": a,
                    "celulas": [
                        {
                            "valor": round(matriz_valores[(a, b)], 2) if matriz_valores.get((a, b)) is not None else None,
                            "fundo": "transparent" if a == b else _fundo_correlacao(matriz_valores.get((a, b))),
                            "diagonal": a == b, "par": b,
                        }
                        for b in na_matriz
                    ],
                }
                for a in na_matriz
            ],
            "ocultos": len(tickers) - len(na_matriz),
            "pares_altos": [par for par in pares if par["valor"] >= CORRELACAO_ALTA],
            "media": round(statistics.fmean(par["valor"] for par in pares), 2),
            "limiar": CORRELACAO_ALTA,
            "pregoes": CORRELACAO_JANELA_PREGOES,
        }

    # ---- beta contra o Ibovespa e cenários de queda
    ibov = {
        data: float(valor) for data, valor in CotacaoIndice.objects.filter(indice=CotacaoIndice.IBOVESPA)
        .order_by("-data").values_list("data", "valor")[:ESTRESSE_JANELA_PREGOES]
    }
    betas, peso_com_beta, beta_carteira = [], 0.0, 0.0
    for p in compradas:
        peso = float(p.valor_atual) / valor_total
        b = _beta(*_retornos_pareados(series[p.ativo.ticker], ibov)) if ibov else None
        betas.append({"ticker": p.ativo.ticker, "beta": round(b, 2) if b is not None else None, "peso_pct": round(peso * 100, 1)})
        if b is not None:
            peso_com_beta += peso
            beta_carteira += peso * b
    beta_carteira = beta_carteira / peso_com_beta if peso_com_beta else None
    cenarios = [
        {
            "queda_ibov_pct": queda,
            "perda_pct": round(beta_carteira * queda, 1),
            "perda_valor": round(valor_total * beta_carteira * queda / 100, 2),
        }
        for queda in ESTRESSE_QUEDAS_IBOV_PCT
    ] if beta_carteira is not None else []

    # ---- piores períodos da carteira atual (mesmas quantidades) no histórico
    quantidades = {p.ativo.ticker: p.quantidade for p in compradas}
    datas_comuns = sorted(set.intersection(*(set(s) for s in series.values())))
    valores = [(d, sum(series[t][d] * q for t, q in quantidades.items())) for d in datas_comuns]
    piores = []
    for rotulo, pregoes in (("Pior dia", 1), ("Pior semana", 5), ("Pior mês", 21)):
        quedas = [
            (valores[i][1] / valores[i - pregoes][1] - 1, valores[i - pregoes][0], valores[i][0])
            for i in range(pregoes, len(valores)) if valores[i - pregoes][1]
        ]
        if quedas:
            pct, inicio, fim = min(quedas)
            piores.append({"rotulo": rotulo, "pct": round(pct * 100, 1), "inicio": inicio, "fim": fim})

    estresse = {
        "beta_carteira": round(beta_carteira, 2) if beta_carteira is not None else None,
        "cobertura_beta_pct": round(peso_com_beta * 100),
        "betas": betas,
        "cenarios": cenarios,
        "piores": piores,
        "pregoes_historico": len(valores),
        "pregoes_ibov": len(ibov),
    }
    return {"correlacao": correlacao, "estresse": estresse}


# --------------------------------------------------------------------------
# Gráfico comparativo: valor de compra x valor atual no histórico (por ativo)
# --------------------------------------------------------------------------
def _fmt_reais(valor: float) -> str:
    return f"R$ {valor:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")


def construir_comparativo_valores(
    posicoes: list[Posicao], dias: int = 30, largura: int = 480, altura: int = 180, padding: int = 40
) -> list[dict]:
    """
    Para cada posição realmente comprada (exclui reservas), monta um gráfico
    de linha com duas séries: o valor de compra (linha de referência, constante
    - é o quanto você pagou) e o valor atual ao longo do histórico de cotações
    (quantidade em carteira x preço de fechamento em cada pregão).

    Cada ativo tem seu próprio gráfico (duas linhas). Ativos sem histórico de
    cotações suficiente (menos de 2 pregões) ficam de fora.
    """
    graficos = []
    for p in posicoes:
        if p.apenas_reservado or p.quantidade <= 0:
            continue

        historico_desc = list(p.ativo.cotacoes.order_by("-data")[:dias])
        if len(historico_desc) < 2:
            continue
        historico = list(reversed(historico_desc))  # cronológico: mais antigo -> mais recente

        valor_compra = float(p.valor_investido)
        valores_atual = [float(p.quantidade) * float(c.preco_fechamento) for c in historico]

        maior_valor = max(max(valores_atual), valor_compra)
        menor_valor = min(min(valores_atual), valor_compra)
        faixa = (maior_valor - menor_valor) or (valor_compra or 1.0)
        maior_valor += faixa * 0.12
        menor_valor -= faixa * 0.12
        faixa = maior_valor - menor_valor

        y_topo = 16
        y_base = altura - padding
        plot_altura = y_base - y_topo
        plot_largura = largura - padding * 2
        passo_x = plot_largura / (len(historico) - 1)

        def escala_y(valor, _menor=menor_valor, _faixa=faixa, _base=y_base, _alt=plot_altura):
            return round(_base - ((valor - _menor) / _faixa) * _alt, 2)

        comandos_atual, pontos = [], []
        for i, (cotacao, valor) in enumerate(zip(historico, valores_atual)):
            x = round(padding + i * passo_x, 2)
            y = escala_y(valor)
            comandos_atual.append(f"{'M' if i == 0 else 'L'}{x} {y}")
            diferenca = valor - valor_compra
            pontos.append({
                "x": x,
                "y": y,
                "data_label": cotacao.data.strftime("%d/%m/%Y"),
                "valor_compra_label": _fmt_reais(valor_compra),
                "valor_atual_label": _fmt_reais(valor),
                "diferenca_label": f"{'+' if diferenca >= 0 else ''}{_fmt_reais(diferenca)}",
                "ganho": valor >= valor_compra,
            })

        y_compra = escala_y(valor_compra)

        eixo_y = []
        for fracao in (0, 1 / 3, 2 / 3, 1):
            valor = menor_valor + fracao * faixa
            y = escala_y(valor)
            eixo_y.append({"y": y, "label_y": round(y + 4, 2), "label": f"R$ {valor:,.0f}".replace(",", ".")})

        graficos.append({
            "ativo": p.ativo,
            "id_svg": f"pontos-comparativo-{p.ativo.id}",
            "largura": largura,
            "altura": altura,
            "padding": padding,
            "x_inicio": padding,
            "x_fim": largura - padding,
            "y_topo": y_topo,
            "y_base": y_base,
            "y_compra": y_compra,
            "path_compra": f"M{padding} {y_compra} L{largura - padding} {y_compra}",
            "path_atual": " ".join(comandos_atual),
            "pontos": pontos,
            "eixo_y": eixo_y,
            "ganho_atual": valores_atual[-1] >= valor_compra,
            "primeira_data": historico[0].data.strftime("%d/%m/%Y"),
            "ultima_data": historico[-1].data.strftime("%d/%m/%Y"),
        })

    return graficos


# --------------------------------------------------------------------------
# Exportação das posições em carteira (Excel / PDF)
# --------------------------------------------------------------------------
COLUNAS_EXPORTACAO_POSICOES = [
    "Ativo", "Quantidade", "Preço compra (R$)", "Valor investido (R$)",
    "Cotação atual (R$)", "Alvo (R$)", "Valor atual (R$)", "Lucro/Perda (R$)", "Lucro/Perda (%)",
    "Meta lucro (%)", "Meta perda (%)", "Tendência", "Dias em carteira",
]


def _linha_exportacao_posicao(p: Posicao) -> list:
    return [
        p.ativo.ticker,
        p.quantidade,
        float(p.preco_medio),
        float(p.valor_investido),
        float(p.preco_atual) if p.preco_atual is not None else None,
        float(p.preco_alvo) if p.preco_alvo is not None else None,
        float(p.valor_atual) if p.valor_atual is not None else None,
        float(p.lucro_perda_valor) if p.lucro_perda_valor is not None else None,
        float(p.lucro_perda_pct) if p.lucro_perda_pct is not None else None,
        float(p.meta_lucro_pct) if p.meta_lucro_pct is not None else None,
        float(p.meta_perda_pct) if p.meta_perda_pct is not None else None,
        p.tendencia or "",
        p.dias_desde_compra,
    ]


def gerar_excel_posicoes(
    posicoes: list[Posicao],
    comparativo_benchmark: dict | None = None,
    concentracao_setor: list[dict] | None = None,
    metricas_risco: dict | None = None,
) -> bytes:
    """
    Gera uma planilha .xlsx com as posições em carteira, pronta para
    download - a aba "Posições" com a grid principal, mais três abas
    espelhando as análises da tela (ver templates/core/posicoes.html):
    "Carteira x Mercado" (comparativo_benchmark), "Concentração por setor"
    e "Indicadores de risco". As três últimas são omitidas (recebem só uma
    linha de aviso) quando a análise correspondente ainda não tem dados.
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Posições"

    ws.append(COLUNAS_EXPORTACAO_POSICOES)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D1526")
        celula.alignment = Alignment(horizontal="center")

    for p in posicoes:
        ws.append(_linha_exportacao_posicao(p))

    for indice in range(1, len(COLUNAS_EXPORTACAO_POSICOES) + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 17
    ws.freeze_panes = "A2"

    def cabecalho(ws, colunas):
        ws.append(colunas)
        for celula in ws[ws.max_row]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")
        for indice in range(1, len(colunas) + 1):
            ws.column_dimensions[get_column_letter(indice)].width = 22

    ws_benchmark = wb.create_sheet("Carteira x Mercado")
    if comparativo_benchmark:
        c = comparativo_benchmark
        ws_benchmark.append([
            f"Desde {c['data_inicio'].strftime('%d/%m/%Y')} até {c['data_fim'].strftime('%d/%m/%Y')} "
            f"({c['dias_periodo']} dias)"
        ])
        ws_benchmark.append([])
        cabecalho(ws_benchmark, ["Indicador", "Retorno no período (%)"])
        ws_benchmark.append(["Sua carteira", float(c["retorno_carteira_pct"])])
        ws_benchmark.append([
            "Ibovespa", float(c["retorno_ibovespa_pct"]) if c["retorno_ibovespa_pct"] is not None else "sem dados"
        ])
        ws_benchmark.append([
            "CDI", float(c["retorno_cdi_pct"]) if c["retorno_cdi_pct"] is not None else "sem dados"
        ])
        ws_benchmark.append([])
        ws_benchmark.append(["Bateu o Ibovespa no período?", "Sim" if c["carteira_bateu_ibovespa"] else "Não"])
        ws_benchmark.append(["Bateu o CDI no período?", "Sim" if c["carteira_bateu_cdi"] else "Não"])
    else:
        ws_benchmark.append(["Sem dados suficientes para comparar com o mercado no momento."])

    ws_setor = wb.create_sheet("Concentração por setor")
    if concentracao_setor:
        cabecalho(ws_setor, ["Setor", "Valor (R$)", "% da carteira"])
        for item in concentracao_setor:
            ws_setor.append([item["setor"], float(item["valor"]), float(item["percentual"])])
    else:
        ws_setor.append(["Sem dados de concentração por setor no momento."])

    ws_risco = wb.create_sheet("Indicadores de risco")
    if metricas_risco and metricas_risco.get("por_ativo"):
        ws_risco.append([f"Índice de concentração (Herfindahl): {metricas_risco['indice_herfindahl']}"])
        ws_risco.append([])
        cabecalho(ws_risco, ["Ativo", "% da carteira", "Volatilidade anualizada (%)", "Drawdown máximo (%)"])
        for item in metricas_risco["por_ativo"]:
            ws_risco.append([
                item["ativo"].ticker,
                float(item["percentual_carteira"]) if item["percentual_carteira"] is not None else None,
                item["volatilidade_pct"],
                item["drawdown_pct"],
            ])
    else:
        ws_risco.append(["Sem indicadores de risco no momento."])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_posicoes(
    posicoes: list[Posicao],
    nome_usuario: str,
    comparativo_benchmark: dict | None = None,
    concentracao_setor: list[dict] | None = None,
    metricas_risco: dict | None = None,
) -> bytes:
    """
    Gera um PDF (paisagem) com as posições em carteira, pronto para
    download - a tabela principal, mais três seções espelhando as análises
    da tela (ver templates/core/posicoes.html): "Carteira x Mercado"
    (comparativo_benchmark), "Concentração por setor" e "Indicadores de
    risco". As três últimas mostram um aviso quando a análise
    correspondente ainda não tem dados, em vez de sumir da página.
    """
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()

    elementos = [
        Paragraph("BolsaTrader - Posições em Carteira", estilos["Title"]),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    def fmt(valor, sufixo="", quando_none="—"):
        return f"{valor:.2f}{sufixo}" if valor is not None else quando_none

    dados = [COLUNAS_EXPORTACAO_POSICOES]
    for p in posicoes:
        l = _linha_exportacao_posicao(p)
        dados.append([
            l[0], str(l[1]), fmt(l[2]), fmt(l[3]), fmt(l[4]), fmt(l[5]), fmt(l[6]), fmt(l[7]), fmt(l[8], "%"),
            fmt(l[9], "%", "padrão"), fmt(l[10], "%", "padrão"), l[11] or "—",
            "hoje" if l[12] == 0 else (f"{l[12]} dias" if l[12] is not None else "—"),
        ])

    if len(dados) == 1:
        elementos.append(Paragraph("Nenhuma posição em carteira.", estilos["Normal"]))
    else:
        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            # 13 colunas: com o padding lateral padrão (6) a tabela passa da margem da página
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.append(tabela)

    def estilo_tabela_secao():
        return TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])

    elementos.append(Spacer(1, 0.8 * cm))
    elementos.append(Paragraph("Sua carteira x mercado", estilos["Heading2"]))
    if comparativo_benchmark:
        c = comparativo_benchmark
        elementos.append(Paragraph(
            escape(
                f"Desde {c['data_inicio'].strftime('%d/%m/%Y')} até {c['data_fim'].strftime('%d/%m/%Y')} "
                f"({c['dias_periodo']} dias)"
            ),
            estilos["Normal"],
        ))
        elementos.append(Spacer(1, 0.2 * cm))
        dados_benchmark = [
            ["Indicador", "Retorno no período (%)"],
            ["Sua carteira", fmt(float(c["retorno_carteira_pct"]), "%")],
            ["Ibovespa", fmt(float(c["retorno_ibovespa_pct"]), "%") if c["retorno_ibovespa_pct"] is not None else "sem dados"],
            ["CDI", fmt(float(c["retorno_cdi_pct"]), "%") if c["retorno_cdi_pct"] is not None else "sem dados"],
        ]
        tabela_benchmark = Table(dados_benchmark, repeatRows=1, colWidths=[6 * cm, 6 * cm])
        tabela_benchmark.setStyle(estilo_tabela_secao())
        elementos.append(tabela_benchmark)
        elementos.append(Spacer(1, 0.2 * cm))
        elementos.append(Paragraph(
            ("✅" if c["carteira_bateu_ibovespa"] else "⚠️") + " Sua carteira "
            + ("bateu" if c["carteira_bateu_ibovespa"] else "não bateu") + " o Ibovespa no período. "
            + ("✅" if c["carteira_bateu_cdi"] else "⚠️") + " Sua carteira "
            + ("bateu" if c["carteira_bateu_cdi"] else "não bateu") + " o CDI no período.",
            estilos["Normal"],
        ))
    else:
        elementos.append(Paragraph("Sem dados suficientes para comparar com o mercado no momento.", estilos["Normal"]))

    elementos.append(Spacer(1, 0.8 * cm))
    elementos.append(Paragraph("Concentração por setor", estilos["Heading2"]))
    if concentracao_setor:
        dados_setor = [["Setor", "Valor (R$)", "% da carteira"]]
        for item in concentracao_setor:
            dados_setor.append([item["setor"], fmt(float(item["valor"])), fmt(float(item["percentual"]), "%")])
        tabela_setor = Table(dados_setor, repeatRows=1, colWidths=[8 * cm, 5 * cm, 5 * cm])
        tabela_setor.setStyle(estilo_tabela_secao())
        elementos.append(tabela_setor)
    else:
        elementos.append(Paragraph("Sem dados de concentração por setor no momento.", estilos["Normal"]))

    elementos.append(Spacer(1, 0.8 * cm))
    elementos.append(Paragraph("Indicadores de risco", estilos["Heading2"]))
    if metricas_risco and metricas_risco.get("por_ativo"):
        elementos.append(Paragraph(
            escape(f"Índice de concentração (Herfindahl): {metricas_risco['indice_herfindahl']}"),
            estilos["Normal"],
        ))
        elementos.append(Spacer(1, 0.2 * cm))
        dados_risco = [["Ativo", "% da carteira", "Volatilidade anualizada (%)", "Drawdown máximo (%)"]]
        for item in metricas_risco["por_ativo"]:
            dados_risco.append([
                item["ativo"].ticker,
                fmt(float(item["percentual_carteira"])) if item["percentual_carteira"] is not None else "—",
                fmt(item["volatilidade_pct"], "%", "aguardando histórico"),
                fmt(item["drawdown_pct"], "%", "aguardando histórico"),
            ])
        tabela_risco = Table(dados_risco, repeatRows=1)
        tabela_risco.setStyle(estilo_tabela_secao())
        elementos.append(tabela_risco)
    else:
        elementos.append(Paragraph("Sem indicadores de risco no momento.", estilos["Normal"]))

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Exportação das operações (Excel / PDF)
#
# Reflete as três grids da página "Minhas Operações" (ver core.views.
# operacao_lista): em carteira, reservadas e vendidas - cada uma com suas
# próprias colunas, igual à tela - em vez de uma tabela genérica única.
# --------------------------------------------------------------------------
COLUNAS_EM_CARTEIRA = [
    "Data", "Dias", "Ativo", "Qtd. comprada", "Preço compra (R$)", "Total compra (R$)",
    "Qtd. vendida", "Preço venda (R$)", "Total venda (R$)", "Saldo",
    "Lucro/Perda realizado (R$)", "Lucro/Perda realizado (%)", "Meta lucro (%)", "Meta perda (%)",
]

COLUNAS_RESERVADAS = [
    "Data", "Dias", "Ativo", "Quantidade", "Preço pretendido (R$)", "Total pretendido (R$)",
    "Variação desde a reserva (%)", "Meta lucro (%)", "Meta perda (%)",
]

COLUNAS_VENDIDAS = [
    "Data compra", "Data venda", "Dias em carteira", "Ativo", "Qtd. comprada",
    "Preço compra (R$)", "Total compra (R$)", "Preço venda (R$)", "Total venda (R$)",
    "Corretora (%)", "Valor líquido a receber (R$)",
    "Lucro/Perda realizado (R$)", "Lucro/Perda realizado (%)",
    "Lucro/Perda por dia (R$)", "Lucro/Perda por dia (%)",
    "Meta lucro (%)", "Meta perda (%)",
]

# Índices (0-based) das colunas numéricas que entram na linha de "TOTAL" no
# fim de cada tabela do relatório - preço unitário, percentuais e metas não
# somam (não fazem sentido agregados), só quantidades e valores em R$.
INDICES_SOMA_EM_CARTEIRA = {3, 5, 6, 8, 9, 10}
INDICES_SOMA_RESERVADAS = {3, 5}
INDICES_SOMA_VENDIDAS = {4, 6, 8, 10, 11}


def _linha_totais(linhas: list[list], indices_soma: set[int], num_colunas: int) -> list:
    """Linha de totais (soma das colunas numéricas indicadas) para o fim de uma tabela do relatório."""
    totais = [None] * num_colunas
    for indice in indices_soma:
        valores = [linha[indice] for linha in linhas if linha[indice] is not None]
        if valores:
            totais[indice] = round(sum(valores), 2)
    return totais


def _linha_em_carteira(op: Operacao) -> list:
    return [
        op.data_operacao,
        op.dias_desde_operacao,
        op.ativo.ticker,
        op.quantidade,
        float(op.preco_unitario),
        float(op.valor_total),
        op.quantidade_vendida,
        float(op.preco_venda) if op.preco_venda is not None else None,
        float(op.valor_total_vendido) if op.valor_total_vendido is not None else None,
        op.saldo,
        float(op.lucro_perda_realizado) if op.lucro_perda_realizado is not None else None,
        float(op.lucro_perda_pct_realizado) if op.lucro_perda_pct_realizado is not None else None,
        float(op.meta_lucro_pct) if op.meta_lucro_pct is not None else None,
        float(op.meta_perda_pct) if op.meta_perda_pct is not None else None,
    ]


def _linha_reservada(op: Operacao) -> list:
    return [
        op.data_operacao,
        op.dias_desde_operacao,
        op.ativo.ticker,
        op.quantidade,
        float(op.preco_unitario),
        float(op.valor_total),
        float(op.variacao_pct_reserva) if op.variacao_pct_reserva is not None else None,
        float(op.meta_lucro_pct) if op.meta_lucro_pct is not None else None,
        float(op.meta_perda_pct) if op.meta_perda_pct is not None else None,
    ]


def _linha_vendida(op: Operacao) -> list:
    return [
        op.data_operacao,
        op.data_venda,
        op.dias_em_carteira_ate_venda,
        op.ativo.ticker,
        op.quantidade,
        float(op.preco_unitario),
        float(op.valor_total),
        float(op.preco_venda) if op.preco_venda is not None else None,
        float(op.valor_total_vendido) if op.valor_total_vendido is not None else None,
        float(op.percentual_corretora) if op.percentual_corretora is not None else None,
        float(op.valor_liquido_vendido) if op.valor_liquido_vendido is not None else None,
        float(op.lucro_perda_realizado) if op.lucro_perda_realizado is not None else None,
        float(op.lucro_perda_pct_realizado) if op.lucro_perda_pct_realizado is not None else None,
        float(op.lucro_perda_por_dia) if op.lucro_perda_por_dia is not None else None,
        float(op.lucro_perda_pct_por_dia) if op.lucro_perda_pct_por_dia is not None else None,
        float(op.meta_lucro_pct) if op.meta_lucro_pct is not None else None,
        float(op.meta_perda_pct) if op.meta_perda_pct is not None else None,
    ]


def gerar_excel_operacoes(
    operacoes_compradas: list[Operacao],
    operacoes_vendidas: list[Operacao],
    operacoes_reservadas: list[Operacao],
    periodo_label_carteira: str,
    periodo_label_vendidas: str,
    periodo_label_reservadas: str,
    apenas_vendas: bool = False,
) -> bytes:
    """
    Gera uma planilha .xlsx com as operações do usuário, pronta para download -
    uma aba para cada grid da página (em carteira, vendidas, reservadas), cada
    uma com o período do seu próprio filtro indicado no topo (Vendidas filtra
    pela data da venda, as outras duas pela data da operação).

    Com apenas_vendas=True (relatório "somente vendas" da aba Vendidas), a
    planilha traz só a aba Vendidas.
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def montar_aba(ws, titulo: str, periodo_label: str, colunas: list[str], linhas: list[list], indices_soma: set[int]):
        ws.append([titulo])
        ws.append([periodo_label])
        ws.append([])
        ws.append(colunas)
        for celula in ws[4]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")

        for linha in linhas:
            ws.append(linha)

        if linhas:
            totais = _linha_totais(linhas, indices_soma, len(colunas))
            totais[0] = "TOTAL"
            ws.append(totais)
            for celula in ws[ws.max_row]:
                celula.font = Font(bold=True)
                celula.fill = PatternFill("solid", fgColor="E4E9F2")

        for indice in range(1, len(colunas) + 1):
            ws.column_dimensions[get_column_letter(indice)].width = 17
        ws.freeze_panes = "A5"

    if apenas_vendas:
        ws_vendidas = wb.active
        ws_vendidas.title = "Vendas"
        montar_aba(
            ws_vendidas, "Vendas", periodo_label_vendidas, COLUNAS_VENDIDAS,
            [_linha_vendida(op) for op in operacoes_vendidas], INDICES_SOMA_VENDIDAS,
        )
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    ws_carteira = wb.active
    ws_carteira.title = "Em carteira"
    montar_aba(
        ws_carteira, "Em carteira", periodo_label_carteira, COLUNAS_EM_CARTEIRA,
        [_linha_em_carteira(op) for op in operacoes_compradas], INDICES_SOMA_EM_CARTEIRA,
    )

    ws_vendidas = wb.create_sheet("Vendidas")
    montar_aba(
        ws_vendidas, "Vendidas", periodo_label_vendidas, COLUNAS_VENDIDAS,
        [_linha_vendida(op) for op in operacoes_vendidas], INDICES_SOMA_VENDIDAS,
    )

    ws_reservadas = wb.create_sheet("Reservadas")
    montar_aba(
        ws_reservadas, "Reservadas", periodo_label_reservadas, COLUNAS_RESERVADAS,
        [_linha_reservada(op) for op in operacoes_reservadas], INDICES_SOMA_RESERVADAS,
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_operacoes(
    operacoes_compradas: list[Operacao],
    operacoes_vendidas: list[Operacao],
    operacoes_reservadas: list[Operacao],
    periodo_label_carteira: str,
    periodo_label_vendidas: str,
    periodo_label_reservadas: str,
    nome_usuario: str,
    apenas_vendas: bool = False,
) -> bytes:
    """
    Gera um PDF (paisagem) com as operações do usuário, pronto para download -
    uma tabela para cada grid da página (em carteira, vendidas, reservadas),
    cada uma com o período do seu próprio filtro (Vendidas filtra pela data
    da venda, as outras duas pela data da operação) indicado logo abaixo do
    título da seção.

    Com apenas_vendas=True (relatório "somente vendas" da aba Vendidas), o
    PDF traz só a tabela de vendas.
    """
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()

    elementos = [
        Paragraph("BolsaTrader - Minhas Vendas" if apenas_vendas else "BolsaTrader - Minhas Operações", estilos["Title"]),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    def fmt(valor, sufixo="", quando_none="—"):
        return f"{valor:.2f}{sufixo}" if valor is not None else quando_none

    def fmt_data(valor):
        return valor.strftime("%d/%m/%Y") if valor is not None else "—"

    def fmt_dias(valor):
        if valor is None:
            return "—"
        return "hoje" if valor == 0 else f"{valor} dias"

    from reportlab.lib.styles import ParagraphStyle

    estilo_cabecalho = ParagraphStyle(
        "CabecalhoTabela", fontSize=6.5, leading=8, textColor=colors.white, alignment=1,
    )
    estilo_celula = ParagraphStyle("CelulaTabela", fontSize=6.5, leading=8, alignment=1)
    estilo_celula_total = ParagraphStyle(
        "CelulaTotal", fontSize=6.5, leading=8, alignment=1, fontName="Helvetica-Bold",
    )

    def adicionar_secao(
        titulo: str, periodo_label: str, colunas: list[str], linhas: list[list], formatadores, indices_soma: set[int]
    ):
        elementos.append(Paragraph(titulo, estilos["Heading2"]))
        elementos.append(Paragraph(escape(periodo_label), estilos["Normal"]))
        elementos.append(Spacer(1, 0.2 * cm))
        if not linhas:
            elementos.append(Paragraph("Nenhuma operação nesta situação no período.", estilos["Normal"]))
            elementos.append(Spacer(1, 0.4 * cm))
            return

        # células viram Paragraph (em vez de texto puro) para quebrar linha em
        # vez de alargar a coluna - com muitas colunas, texto puro faz a
        # tabela ficar mais larga que a página (estoura a margem direita).
        dados = [[Paragraph(str(c), estilo_cabecalho) for c in colunas]]
        for linha in linhas:
            dados.append([
                Paragraph(str(f(v)), estilo_celula) for f, v in zip(formatadores, linha)
            ])

        totais = _linha_totais(linhas, indices_soma, len(colunas))
        linha_totais = [Paragraph("TOTAL", estilo_celula_total)]
        for f, v in list(zip(formatadores, totais))[1:]:
            texto = f(v) if v is not None else "—"
            linha_totais.append(Paragraph(str(texto), estilo_celula_total))
        dados.append(linha_totais)

        largura_coluna = doc.width / len(colunas)
        tabela = Table(dados, colWidths=[largura_coluna] * len(colunas), repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#d7deec")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#eef2f7")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.append(tabela)
        elementos.append(Spacer(1, 0.6 * cm))

    identidade = lambda v: v if v is not None else "—"

    formatadores_vendidas = [
        fmt_data, fmt_data, fmt_dias, identidade, identidade, fmt, fmt, fmt, fmt,
        lambda v: fmt(v, "%"), fmt, fmt, lambda v: fmt(v, "%"), fmt, lambda v: fmt(v, "%"),
        lambda v: fmt(v, "%", "padrão"), lambda v: fmt(v, "%", "padrão"),
    ]

    if apenas_vendas:
        adicionar_secao(
            "Vendas", periodo_label_vendidas, COLUNAS_VENDIDAS,
            [_linha_vendida(op) for op in operacoes_vendidas], formatadores_vendidas, INDICES_SOMA_VENDIDAS,
        )
        doc.build(elementos)
        return buffer.getvalue()

    adicionar_secao(
        "Em carteira", periodo_label_carteira, COLUNAS_EM_CARTEIRA,
        [_linha_em_carteira(op) for op in operacoes_compradas],
        [fmt_data, fmt_dias, identidade, identidade, fmt, fmt, identidade, fmt, fmt, identidade,
         fmt, lambda v: fmt(v, "%"), lambda v: fmt(v, "%", "padrão"), lambda v: fmt(v, "%", "padrão")],
        INDICES_SOMA_EM_CARTEIRA,
    )
    adicionar_secao(
        "Vendidas", periodo_label_vendidas, COLUNAS_VENDIDAS,
        [_linha_vendida(op) for op in operacoes_vendidas], formatadores_vendidas, INDICES_SOMA_VENDIDAS,
    )
    adicionar_secao(
        "Reservadas", periodo_label_reservadas, COLUNAS_RESERVADAS,
        [_linha_reservada(op) for op in operacoes_reservadas],
        [fmt_data, fmt_dias, identidade, identidade, fmt, fmt,
         lambda v: fmt(v, "%"), lambda v: fmt(v, "%", "padrão"), lambda v: fmt(v, "%", "padrão")],
        INDICES_SOMA_RESERVADAS,
    )

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Relatórios de Minhas Operações de Cripto - espelham gerar_excel_operacoes/
# gerar_pdf_operacoes (ações), com as colunas da própria tela de cripto (sem
# corretora, sem lucro/perda por dia - campos que não existem em
# OperacaoCripto) e sem rótulo de período, já que a lista de cripto não tem
# filtro de data (ver core.views.cripto_operacao_lista).
# --------------------------------------------------------------------------
COLUNAS_EM_CARTEIRA_CRIPTO = [
    "Data", "Dias", "Moeda", "Quantidade", "Preço compra (R$)", "Total compra (R$)",
    "Qtd. vendida", "Saldo", "Meta lucro (%)", "Meta perda (%)",
]
COLUNAS_RESERVADAS_CRIPTO = [
    "Data", "Dias", "Moeda", "Quantidade", "Preço pretendido (R$)", "Total pretendido (R$)",
    "Variação desde a reserva (%)", "Meta lucro (%)", "Meta perda (%)",
]
COLUNAS_VENDIDAS_CRIPTO = [
    "Data compra", "Data venda", "Dias em carteira", "Moeda", "Quantidade",
    "Preço compra (R$)", "Total compra (R$)", "Preço venda (R$)", "Total venda (R$)",
    "Lucro/Perda realizado (R$)", "Lucro/Perda realizado (%)", "Meta lucro (%)", "Meta perda (%)",
]

INDICES_SOMA_EM_CARTEIRA_CRIPTO = {3, 5, 6, 7}
INDICES_SOMA_RESERVADAS_CRIPTO = {3, 5}
INDICES_SOMA_VENDIDAS_CRIPTO = {4, 6, 8, 9}


def _linha_em_carteira_cripto(op: OperacaoCripto) -> list:
    return [
        op.data_operacao,
        op.dias_desde_operacao,
        op.cripto_ativo.coin,
        float(op.quantidade),
        float(op.preco_unitario),
        float(op.valor_total),
        float(op.quantidade_vendida),
        float(op.saldo),
        float(op.meta_lucro_pct) if op.meta_lucro_pct is not None else None,
        float(op.meta_perda_pct) if op.meta_perda_pct is not None else None,
    ]


def _linha_reservada_cripto(op: OperacaoCripto) -> list:
    return [
        op.data_operacao,
        op.dias_desde_operacao,
        op.cripto_ativo.coin,
        float(op.quantidade),
        float(op.preco_unitario),
        float(op.valor_total),
        float(op.variacao_pct_reserva) if op.variacao_pct_reserva is not None else None,
        float(op.meta_lucro_pct) if op.meta_lucro_pct is not None else None,
        float(op.meta_perda_pct) if op.meta_perda_pct is not None else None,
    ]


def _linha_vendida_cripto(op: OperacaoCripto) -> list:
    return [
        op.data_operacao,
        op.data_venda,
        op.dias_em_carteira_ate_venda,
        op.cripto_ativo.coin,
        float(op.quantidade),
        float(op.preco_unitario),
        float(op.valor_total),
        float(op.preco_venda) if op.preco_venda is not None else None,
        float(op.valor_total_vendido) if op.valor_total_vendido is not None else None,
        float(op.lucro_perda_realizado) if op.lucro_perda_realizado is not None else None,
        float(op.lucro_perda_pct_realizado) if op.lucro_perda_pct_realizado is not None else None,
        float(op.meta_lucro_pct) if op.meta_lucro_pct is not None else None,
        float(op.meta_perda_pct) if op.meta_perda_pct is not None else None,
    ]


def gerar_excel_operacoes_cripto(
    operacoes_compradas: list[OperacaoCripto],
    operacoes_vendidas: list[OperacaoCripto],
    operacoes_reservadas: list[OperacaoCripto],
) -> bytes:
    """Gera uma planilha .xlsx com as operações de criptomoeda do usuário - espelha gerar_excel_operacoes (ações), sem rótulo de período (a lista de cripto não tem filtro de data)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def montar_aba(ws, titulo: str, colunas: list[str], linhas: list[list], indices_soma: set[int]):
        ws.append([titulo])
        ws.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])
        ws.append(colunas)
        for celula in ws[4]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")

        for linha in linhas:
            ws.append(linha)

        if linhas:
            totais = _linha_totais(linhas, indices_soma, len(colunas))
            totais[0] = "TOTAL"
            ws.append(totais)
            for celula in ws[ws.max_row]:
                celula.font = Font(bold=True)
                celula.fill = PatternFill("solid", fgColor="E4E9F2")

        for indice in range(1, len(colunas) + 1):
            ws.column_dimensions[get_column_letter(indice)].width = 17
        ws.freeze_panes = "A5"

    ws_carteira = wb.active
    ws_carteira.title = "Em carteira"
    montar_aba(
        ws_carteira, "Em carteira", COLUNAS_EM_CARTEIRA_CRIPTO,
        [_linha_em_carteira_cripto(op) for op in operacoes_compradas], INDICES_SOMA_EM_CARTEIRA_CRIPTO,
    )

    ws_vendidas = wb.create_sheet("Vendidas")
    montar_aba(
        ws_vendidas, "Vendidas", COLUNAS_VENDIDAS_CRIPTO,
        [_linha_vendida_cripto(op) for op in operacoes_vendidas], INDICES_SOMA_VENDIDAS_CRIPTO,
    )

    ws_reservadas = wb.create_sheet("Reservadas")
    montar_aba(
        ws_reservadas, "Reservadas", COLUNAS_RESERVADAS_CRIPTO,
        [_linha_reservada_cripto(op) for op in operacoes_reservadas], INDICES_SOMA_RESERVADAS_CRIPTO,
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_operacoes_cripto(
    operacoes_compradas: list[OperacaoCripto],
    operacoes_vendidas: list[OperacaoCripto],
    operacoes_reservadas: list[OperacaoCripto],
    nome_usuario: str,
) -> bytes:
    """Gera um PDF (paisagem) com as operações de criptomoeda do usuário - espelha gerar_pdf_operacoes (ações), sem rótulo de período."""
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()

    elementos = [
        Paragraph("BolsaTrader - Minhas Operações de Cripto", estilos["Title"]),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    def fmt(valor, sufixo="", quando_none="—"):
        return f"{valor:.8f}{sufixo}" if valor is not None else quando_none

    def fmt_moeda(valor, quando_none="—"):
        return f"{valor:.2f}" if valor is not None else quando_none

    def fmt_data(valor):
        return valor.strftime("%d/%m/%Y") if valor is not None else "—"

    def fmt_dias(valor):
        if valor is None:
            return "—"
        return "hoje" if valor == 0 else f"{valor} dias"

    identidade = lambda v: v if v is not None else "—"

    estilo_cabecalho = ParagraphStyle(
        "CabecalhoTabelaCripto", fontSize=6.5, leading=8, textColor=colors.white, alignment=1,
    )
    estilo_celula = ParagraphStyle("CelulaTabelaCripto", fontSize=6.5, leading=8, alignment=1)
    estilo_celula_total = ParagraphStyle(
        "CelulaTotalCripto", fontSize=6.5, leading=8, alignment=1, fontName="Helvetica-Bold",
    )

    def adicionar_secao(titulo: str, colunas: list[str], linhas: list[list], formatadores, indices_soma: set[int]):
        elementos.append(Paragraph(titulo, estilos["Heading2"]))
        elementos.append(Spacer(1, 0.2 * cm))
        if not linhas:
            elementos.append(Paragraph("Nenhuma operação de criptomoeda nesta situação.", estilos["Normal"]))
            elementos.append(Spacer(1, 0.4 * cm))
            return

        dados = [[Paragraph(str(c), estilo_cabecalho) for c in colunas]]
        for linha in linhas:
            dados.append([Paragraph(str(f(v)), estilo_celula) for f, v in zip(formatadores, linha)])

        totais = _linha_totais(linhas, indices_soma, len(colunas))
        linha_totais = [Paragraph("TOTAL", estilo_celula_total)]
        for f, v in list(zip(formatadores, totais))[1:]:
            texto = f(v) if v is not None else "—"
            linha_totais.append(Paragraph(str(texto), estilo_celula_total))
        dados.append(linha_totais)

        largura_coluna = doc.width / len(colunas)
        tabela = Table(dados, colWidths=[largura_coluna] * len(colunas), repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#d7deec")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#eef2f7")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.append(tabela)
        elementos.append(Spacer(1, 0.6 * cm))

    adicionar_secao(
        "Em carteira", COLUNAS_EM_CARTEIRA_CRIPTO,
        [_linha_em_carteira_cripto(op) for op in operacoes_compradas],
        [fmt_data, fmt_dias, identidade, fmt, fmt_moeda, fmt_moeda, fmt, fmt,
         lambda v: fmt_moeda(v, "padrão"), lambda v: fmt_moeda(v, "padrão")],
        INDICES_SOMA_EM_CARTEIRA_CRIPTO,
    )
    adicionar_secao(
        "Vendidas", COLUNAS_VENDIDAS_CRIPTO,
        [_linha_vendida_cripto(op) for op in operacoes_vendidas],
        [fmt_data, fmt_data, fmt_dias, identidade, fmt, fmt_moeda, fmt_moeda, fmt_moeda, fmt_moeda,
         fmt_moeda, lambda v: fmt_moeda(v, "%"), lambda v: fmt_moeda(v, "padrão"), lambda v: fmt_moeda(v, "padrão")],
        INDICES_SOMA_VENDIDAS_CRIPTO,
    )
    adicionar_secao(
        "Reservadas", COLUNAS_RESERVADAS_CRIPTO,
        [_linha_reservada_cripto(op) for op in operacoes_reservadas],
        [fmt_data, fmt_dias, identidade, fmt, fmt_moeda, fmt_moeda,
         lambda v: fmt_moeda(v, "%"), lambda v: fmt_moeda(v, "padrão"), lambda v: fmt_moeda(v, "padrão")],
        INDICES_SOMA_RESERVADAS_CRIPTO,
    )

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Gráfico de acompanhamento (linha simples com o histórico de fechamento)
# --------------------------------------------------------------------------
def construir_grafico_cotacoes(
    historico_desc: list[Cotacao], largura: int = 560, altura: int = 160, padding: int = 26,
    id_svg_sufixo: str = "",
) -> dict | None:
    """
    Monta os dados de um gráfico de linha (SVG) com o histórico de preços de
    fechamento de um ativo, para o acompanhamento visual na página de Análise
    de Mercado. Recebe o histórico já carregado (mais recente primeiro, mesma
    lista usada na tabela) para não repetir consulta ao banco.

    `id_svg_sufixo` evita ids de SVG duplicados na mesma página quando o
    mesmo ativo aparece em mais de um gráfico (ex: o gráfico "desde a
    compra" e o gráfico "desde a reserva" do mesmo ativo, em Histórico e
    Gráficos) - ids duplicados quebrariam o crosshair/tooltip (ver grafico.js,
    que busca os pontos por getElementById).

    Retorna None quando não há pontos suficientes para traçar uma linha.
    """
    cotacoes = list(reversed(historico_desc))  # cronológico: mais antigo -> mais recente
    if len(cotacoes) < 2:
        return None

    precos = [float(c.preco_fechamento) for c in cotacoes]
    preco_min, preco_max = min(precos), max(precos)
    faixa = (preco_max - preco_min) or 1.0

    plot_largura = largura - (padding * 2)
    plot_altura = altura - (padding * 2)
    passo_x = plot_largura / (len(cotacoes) - 1)

    pontos = []
    comandos_path = []
    for i, (cotacao, preco) in enumerate(zip(cotacoes, precos)):
        x = round(padding + i * passo_x, 2)
        y = round(padding + (1 - (preco - preco_min) / faixa) * plot_altura, 2)
        comandos_path.append(f"{'M' if i == 0 else 'L'}{x} {y}")

        variacao = cotacao.variacao_dia_pct
        pontos.append({
            "x": x,
            "y": y,
            "data_label": cotacao.data.strftime("%d/%m/%Y"),
            "preco_label": f"R$ {preco:.2f}".replace(".", ","),
            "variacao_label": f"{'+' if variacao and variacao > 0 else ''}{variacao}%".replace(".", ",")
            if variacao is not None
            else None,
            "variacao_positiva": bool(variacao and variacao >= 0),
        })

    return {
        "id_svg": f"pontos-cotacoes-{cotacoes[0].ativo_id}{id_svg_sufixo}",
        "largura": largura,
        "altura": altura,
        "padding": padding,
        "x_inicio": padding,
        "x_fim": largura - padding,
        "y_topo": padding,
        "y_base": altura - padding,
        "path_d": " ".join(comandos_path),
        "pontos": pontos,
        "faixa_min_label": f"R$ {preco_min:.2f}".replace(".", ","),
        "faixa_max_label": f"R$ {preco_max:.2f}".replace(".", ","),
        "tendencia_alta": precos[-1] >= precos[0],
    }


# --------------------------------------------------------------------------
# Análise simples de tendência de mercado (alta / baixa / neutro)
# --------------------------------------------------------------------------
def _tendencia_a_partir_de_precos_desc(precos_desc: list, janela_curta: int = 3, janela_longa: int = 10) -> str:
    """
    Núcleo puro (sem acesso ao banco) da análise de tendência: recebe preços
    já em ordem do mais recente para o mais antigo (no máximo `janela_longa`
    itens) e classifica comparando a média móvel curta com a longa. Extraído
    de analisar_tendencia para ser reaproveitado pelo backtesting (ver
    backtest_sinais_robo), que precisa simular a tendência "como ela era" em
    cada dia do histórico, não só a de hoje.
    """
    if len(precos_desc) < 2:
        return "DADOS_INSUFICIENTES"

    precos = [Decimal(p) for p in precos_desc]
    curta = precos[: min(janela_curta, len(precos))]
    media_curta = sum(curta) / len(curta)
    media_longa = sum(precos) / len(precos)

    if media_longa == 0:
        return "NEUTRO"

    diferenca_pct = ((media_curta - media_longa) / media_longa) * 100

    if diferenca_pct >= Decimal("0.5"):
        return "ALTA"
    if diferenca_pct <= Decimal("-0.5"):
        return "BAIXA"
    return "NEUTRO"


def analisar_tendencia(ativo: Ativo, janela_curta: int = 3, janela_longa: int = 10) -> str:
    """
    Analisa o histórico de cotações diárias de um ativo e classifica a
    tendência recente comparando a média móvel curta com a média móvel longa
    (uma aproximação simples de "possível alta/baixa" para dar suporte à
    decisão do investidor - não é recomendação de investimento).
    """
    historico = list(
        ativo.cotacoes.order_by("-data").values_list("preco_fechamento", flat=True)[:janela_longa]
    )
    return _tendencia_a_partir_de_precos_desc(historico, janela_curta, janela_longa)


TENDENCIA_LABELS = {
    "ALTA": "Possível alta",
    "BAIXA": "Possível baixa",
    "NEUTRO": "Estável",
    "DADOS_INSUFICIENTES": "Aguardando histórico",
}


# --------------------------------------------------------------------------
# Indicadores técnicos (RSI e MACD) e sinal consolidado de compra/venda
# --------------------------------------------------------------------------
def calcular_rsi(precos: list[float], periodo: int = 14) -> float | None:
    """
    RSI (Índice de Força Relativa) pelo método de suavização de Wilder, a
    partir de uma lista de preços de fechamento em ordem cronológica (mais
    antigo primeiro). Retorna None sem histórico suficiente (precisa de pelo
    menos periodo + 1 preços). Acima de 70 = sobrecomprado (sinal de venda);
    abaixo de 30 = sobrevendido (sinal de compra).
    """
    if len(precos) < periodo + 1:
        return None

    ganhos, perdas = [], []
    for anterior, atual in zip(precos, precos[1:]):
        variacao = atual - anterior
        ganhos.append(max(variacao, 0.0))
        perdas.append(max(-variacao, 0.0))

    media_ganho = sum(ganhos[:periodo]) / periodo
    media_perda = sum(perdas[:periodo]) / periodo

    for i in range(periodo, len(ganhos)):
        media_ganho = (media_ganho * (periodo - 1) + ganhos[i]) / periodo
        media_perda = (media_perda * (periodo - 1) + perdas[i]) / periodo

    if media_perda == 0:
        return 100.0
    rs = media_ganho / media_perda
    return round(100 - (100 / (1 + rs)), 2)


def _classificar_rsi(rsi: float | None) -> str:
    if rsi is None:
        return "DADOS_INSUFICIENTES"
    if rsi >= 70:
        return "SOBRECOMPRADO"
    if rsi <= 30:
        return "SOBREVENDIDO"
    return "NEUTRO"


RSI_LABELS = {
    "SOBRECOMPRADO": "Sobrecomprado (RSI alto)",
    "SOBREVENDIDO": "Sobrevendido (RSI baixo)",
    "NEUTRO": "Neutro",
    "DADOS_INSUFICIENTES": "Aguardando histórico",
}
# classe CSS (badge-alta/badge-baixa/badge-neutro) de cada rótulo - RSI baixo é
# sinal de compra (verde/"alta"), RSI alto é sinal de venda (vermelho/"baixa")
RSI_CLASSES = {
    "SOBRECOMPRADO": "baixa",
    "SOBREVENDIDO": "alta",
    "NEUTRO": "neutro",
    "DADOS_INSUFICIENTES": "neutro",
}
# traduz o vocabulário do RSI (SOBRECOMPRADO/SOBREVENDIDO) para o vocabulário
# de voto ALTA/BAIXA/NEUTRO/DADOS_INSUFICIENTES usado pelo Scanner Técnico
# (ver core.services.escanear_ativo_precos) - o rótulo exibido continua sendo
# o mais descritivo (RSI_LABELS/RSI_CLASSES), só a contagem de votos usa isto.
RSI_CHAVE_PARA_VOTO_SCANNER = {
    "SOBRECOMPRADO": "BAIXA",
    "SOBREVENDIDO": "ALTA",
    "NEUTRO": "NEUTRO",
    "DADOS_INSUFICIENTES": "DADOS_INSUFICIENTES",
}


def _media_movel_exponencial(valores: list[float], periodo: int) -> list[float]:
    """EMA de `valores`, iniciada com a média simples dos primeiros `periodo` itens."""
    if len(valores) < periodo:
        return []
    fator = 2 / (periodo + 1)
    ema = [sum(valores[:periodo]) / periodo]
    for valor in valores[periodo:]:
        ema.append(valor * fator + ema[-1] * (1 - fator))
    return ema


def calcular_macd(
    precos: list[float], periodo_curto: int = 12, periodo_longo: int = 26, periodo_sinal: int = 9
) -> dict | None:
    """
    MACD (Moving Average Convergence Divergence): linha MACD = EMA curta -
    EMA longa; linha de sinal = EMA do MACD. Histograma positivo (MACD acima
    da linha de sinal) = cruzamento de alta; negativo = cruzamento de baixa.
    Retorna None sem histórico suficiente.
    """
    if len(precos) < periodo_longo + periodo_sinal:
        return None

    ema_curta = _media_movel_exponencial(precos, periodo_curto)
    ema_longa = _media_movel_exponencial(precos, periodo_longo)
    # ema_curta[i] corresponde ao preço de índice (periodo_curto - 1 + i) na
    # lista original; ema_longa[j] ao índice (periodo_longo - 1 + j). Alinha
    # as duas séries pelo índice original antes de subtrair.
    deslocamento = periodo_longo - periodo_curto
    linha_macd = [c - l for c, l in zip(ema_curta[deslocamento:], ema_longa)]

    linha_sinal = _media_movel_exponencial(linha_macd, periodo_sinal)
    if not linha_sinal:
        return None

    macd_atual = linha_macd[-1]
    sinal_atual = linha_sinal[-1]
    histograma = macd_atual - sinal_atual

    return {
        "macd": round(macd_atual, 4),
        "sinal": round(sinal_atual, 4),
        "histograma": round(histograma, 4),
        "cruzamento": "ALTA" if histograma > 0 else ("BAIXA" if histograma < 0 else "NEUTRO"),
    }


MACD_LABELS = {
    "ALTA": "Cruzamento de alta",
    "BAIXA": "Cruzamento de baixa",
    "NEUTRO": "Neutro",
    "DADOS_INSUFICIENTES": "Aguardando histórico",
}
MACD_CLASSES = {
    "ALTA": "alta",
    "BAIXA": "baixa",
    "NEUTRO": "neutro",
    "DADOS_INSUFICIENTES": "neutro",
}

SINAL_GERAL_LABELS = {"COMPRA": "Sinal de compra", "VENDA": "Sinal de venda", "NEUTRO": "Neutro"}
SINAL_GERAL_CLASSES = {"COMPRA": "alta", "VENDA": "baixa", "NEUTRO": "neutro"}


def analisar_indicadores_tecnicos_precos(
    precos_cronologico: list[float], janela_curta: int = 3, janela_longa: int = 10
) -> dict:
    """
    Núcleo puro (sem acesso ao banco) da análise técnica consolidada: recebe
    os preços de fechamento em ordem cronológica (mais antigo primeiro) "até
    aquele momento" e devolve tendência, RSI, MACD e o sinal geral. Extraído
    de analisar_indicadores_tecnicos para ser reaproveitado pelo backtesting
    (ver backtest_sinais_robo), que recalcula o mesmo sinal em cada dia do
    passado usando só os preços disponíveis até aquele dia - exatamente a
    mesma lógica usada ao vivo, sem duplicar as regras de voto.
    """
    tendencia = _tendencia_a_partir_de_precos_desc(
        list(reversed(precos_cronologico))[:janela_longa], janela_curta, janela_longa
    )

    rsi = calcular_rsi(precos_cronologico)
    rsi_label_chave = _classificar_rsi(rsi)

    macd = calcular_macd(precos_cronologico)
    macd_label_chave = macd["cruzamento"] if macd else "DADOS_INSUFICIENTES"

    votos_compra = 0
    votos_venda = 0
    if tendencia == "ALTA":
        votos_compra += 1
    elif tendencia == "BAIXA":
        votos_venda += 1
    if rsi_label_chave == "SOBREVENDIDO":
        votos_compra += 1
    elif rsi_label_chave == "SOBRECOMPRADO":
        votos_venda += 1
    if macd_label_chave == "ALTA":
        votos_compra += 1
    elif macd_label_chave == "BAIXA":
        votos_venda += 1

    if votos_compra > votos_venda:
        sinal_geral = "COMPRA"
    elif votos_venda > votos_compra:
        sinal_geral = "VENDA"
    else:
        sinal_geral = "NEUTRO"

    return {
        "tendencia": tendencia,
        "tendencia_label": TENDENCIA_LABELS.get(tendencia, tendencia),
        "tendencia_classe": tendencia.lower() if tendencia != "DADOS_INSUFICIENTES" else "neutro",
        "rsi": rsi,
        "rsi_label": RSI_LABELS[rsi_label_chave],
        "rsi_classe": RSI_CLASSES[rsi_label_chave],
        "macd": macd,
        "macd_label": MACD_LABELS[macd_label_chave],
        "macd_classe": MACD_CLASSES[macd_label_chave],
        "sinal_geral": sinal_geral,
        "sinal_geral_label": SINAL_GERAL_LABELS[sinal_geral],
        "sinal_geral_classe": SINAL_GERAL_CLASSES[sinal_geral],
        "votos_compra": votos_compra,
        "votos_venda": votos_venda,
    }


def analisar_indicadores_tecnicos(ativo: Ativo, janela_curta: int = 3, janela_longa: int = 10) -> dict:
    """
    Consolida a análise técnica de um ativo: tendência (cruzamento de médias
    móveis já existente), RSI e MACD - cada um com seu próprio sinal - mais um
    sinal geral simples (quantos indicadores apontam compra x quantos apontam
    venda; empate ou indicadores insuficientes = neutro). É só uma referência
    de apoio à decisão, não uma recomendação de investimento.
    """
    historico = list(
        ativo.cotacoes.order_by("data").values_list("preco_fechamento", flat=True)
    )
    precos = [float(p) for p in historico]
    return analisar_indicadores_tecnicos_precos(precos, janela_curta, janela_longa)


# --------------------------------------------------------------------------
# Scanner Técnico da carteira (ações compradas)
#
# Combina IFR + médias móveis + MACD + volume + volatilidade + tendência num
# só painel por ativo comprado. Diferença deliberada do sinal_geral do robô
# consultor (analisar_indicadores_tecnicos_precos, que sempre elege um
# "vencedor" por votação, inclusive em caso de empate): aqui, sempre que pelo
# menos um indicador direcional aponta pra cada lado ao mesmo tempo, o
# veredito é explicitamente CONFLITANTE, em vez de fingir que existe uma
# previsão certa. Só fecha PREDOMINIO_ALTA/PREDOMINIO_BAIXA quando todos os
# indicadores direcionais que não ficaram neutros/sem dados concordam entre
# si. Volatilidade não entra nessa votação - é risco (magnitude), não direção.
# --------------------------------------------------------------------------
MEDIA_MOVEL_CURTA_PADRAO = 9
MEDIA_MOVEL_LONGA_PADRAO = 21


def calcular_medias_moveis(
    precos_cronologico: list[float], curta: int = MEDIA_MOVEL_CURTA_PADRAO, longa: int = MEDIA_MOVEL_LONGA_PADRAO
) -> dict | None:
    """
    Médias móveis simples curta (padrão: 9 pregões) e longa (padrão: 21
    pregões) a partir de preços em ordem cronológica (mais antigo primeiro).
    None sem histórico suficiente para a média longa.
    """
    if len(precos_cronologico) < longa:
        return None

    media_curta = sum(precos_cronologico[-curta:]) / curta
    media_longa = sum(precos_cronologico[-longa:]) / longa
    preco_atual = precos_cronologico[-1]

    return {
        "media_curta": round(media_curta, 2),
        "media_longa": round(media_longa, 2),
        "preco_atual": round(preco_atual, 2),
        "preco_acima_da_longa": preco_atual > media_longa,
        "preco_abaixo_da_longa": preco_atual < media_longa,
        "curta_acima_da_longa": media_curta > media_longa,
        "curta_abaixo_da_longa": media_curta < media_longa,
    }


def _classificar_medias_moveis(info: dict | None) -> str:
    """
    ALTA só quando o preço atual E a média curta concordam (as duas acima da
    longa); BAIXA quando as duas concordam abaixo. Qualquer combinação mista
    (ex: preço já voltou pra cima da longa, mas a curta ainda não cruzou) é um
    sinal internamente conflitante desse próprio indicador - fica NEUTRO em
    vez de forçar um lado.
    """
    if info is None:
        return "DADOS_INSUFICIENTES"
    if info["preco_acima_da_longa"] and info["curta_acima_da_longa"]:
        return "ALTA"
    if info["preco_abaixo_da_longa"] and info["curta_abaixo_da_longa"]:
        return "BAIXA"
    return "NEUTRO"


MEDIAS_MOVEIS_LABELS = {
    "ALTA": "Preço e médias em alinhamento de alta",
    "BAIXA": "Preço e médias em alinhamento de baixa",
    "NEUTRO": "Sem alinhamento claro",
    "DADOS_INSUFICIENTES": "Aguardando histórico",
}
MEDIAS_MOVEIS_CLASSES = {
    "ALTA": "alta", "BAIXA": "baixa", "NEUTRO": "neutro", "DADOS_INSUFICIENTES": "neutro",
}


VOLUME_JANELA_PADRAO = 20


def analisar_volume_precos(
    volumes_cronologico: list, precos_cronologico: list[float], janela: int = VOLUME_JANELA_PADRAO
) -> dict:
    """
    Compara o volume do último pregão com a média dos anteriores (janela
    padrão: 20 pregões) e cruza com a direção do preço no mesmo dia: volume
    acima da média (margem de 10%, pra não reagir a ruído) confirmando alta é
    força compradora; volume acima da média numa queda ("volume aumentando na
    queda") é força vendedora.

    Sem pelo menos 6 pregões com volume registrado, fica "DADOS_INSUFICIENTES"
    em vez de fingir um sinal - o campo Cotacao.volume foi adicionado depois
    de várias cotações já existirem sem ele (ver backfill_historico_cotacoes),
    então esse indicador só passa a valer conforme o histórico com volume vai
    se acumulando.
    """
    pares = [(v, p) for v, p in zip(volumes_cronologico, precos_cronologico) if v is not None]
    if len(pares) < 6:
        return {"classificacao": "DADOS_INSUFICIENTES", "volume_atual": None, "media_volume": None, "variacao_pct": None}

    volumes_validos = [v for v, _ in pares]
    volume_atual = volumes_validos[-1]
    anteriores = volumes_validos[-(janela + 1):-1]
    media_volume = sum(anteriores) / len(anteriores)

    variacao_pct = round((volume_atual - media_volume) / media_volume * 100, 2) if media_volume else None
    acima_da_media = media_volume > 0 and volume_atual > media_volume * 1.1

    preco_atual = precos_cronologico[-1]
    preco_anterior = precos_cronologico[-2] if len(precos_cronologico) >= 2 else None
    preco_caiu = preco_anterior is not None and preco_atual < preco_anterior
    preco_subiu = preco_anterior is not None and preco_atual > preco_anterior

    if acima_da_media and preco_caiu:
        classificacao = "BAIXA"
    elif acima_da_media and preco_subiu:
        classificacao = "ALTA"
    else:
        classificacao = "NEUTRO"

    return {
        "classificacao": classificacao,
        "volume_atual": volume_atual,
        "media_volume": round(media_volume, 0),
        "variacao_pct": variacao_pct,
    }


VOLUME_LABELS = {
    "ALTA": "Volume acima da média confirmando a alta",
    "BAIXA": "Volume aumentando na queda",
    "NEUTRO": "Volume dentro do normal",
    "DADOS_INSUFICIENTES": "Sem histórico de volume suficiente ainda",
}
VOLUME_CLASSES = {
    "ALTA": "alta", "BAIXA": "baixa", "NEUTRO": "neutro", "DADOS_INSUFICIENTES": "neutro",
}


VOLATILIDADE_ALTA_LIMIAR_PCT = 40.0
VOLATILIDADE_BAIXA_LIMIAR_PCT = 15.0


def _classificar_volatilidade(volatilidade_pct: float | None) -> str:
    """Não é direcional (não vota alta nem baixa) - é só um alerta de risco/contexto, independente do veredito."""
    if volatilidade_pct is None:
        return "DADOS_INSUFICIENTES"
    if volatilidade_pct >= VOLATILIDADE_ALTA_LIMIAR_PCT:
        return "ALTA"
    if volatilidade_pct <= VOLATILIDADE_BAIXA_LIMIAR_PCT:
        return "BAIXA"
    return "MODERADA"


VOLATILIDADE_LABELS = {
    "ALTA": "Alta volatilidade - oscila bastante",
    "MODERADA": "Volatilidade moderada",
    "BAIXA": "Baixa volatilidade - mais estável",
    "DADOS_INSUFICIENTES": "Aguardando histórico",
}
# não é direção de preço - "ALTA volatilidade" usa a cor de alerta (mesma do
# "BAIXA" nos outros indicadores) e "BAIXA volatilidade" usa a cor "positiva"
# (menos risco), pra não ser lido como um voto de alta/baixa do preço.
VOLATILIDADE_CLASSES = {
    "ALTA": "baixa", "MODERADA": "neutro", "BAIXA": "alta", "DADOS_INSUFICIENTES": "neutro",
}

def calcular_atr(maximas: list, minimas: list, fechamentos: list[float], periodo: int = 14) -> float | None:
    """
    ATR (Average True Range) pela suavização de Wilder, a partir de máxima,
    mínima e fechamento em ordem cronológica (mais antigo primeiro) - mede o
    "tamanho médio" da oscilação diária em R$. Não vota direção (é contexto
    de risco, igual volatilidade_pct), mas serve pra dimensionar stop/alvo do
    plano de trade sugerido (ver sugerir_plano_trade).

    Tolera máxima/mínima faltando em pregões antigos (campos adicionados
    depois de várias cotações já existirem, ver backfill_historico_cotacoes)
    - descarta esses pregões da conta em vez de quebrar, igual
    analisar_volume_precos faz com volume ausente. Precisa de pelo menos
    `periodo` pregões com máxima E mínima para retornar um valor.
    """
    trs = []
    fechamento_anterior = None
    for maxima, minima, fechamento in zip(maximas, minimas, fechamentos):
        if maxima is not None and minima is not None:
            maxima, minima = float(maxima), float(minima)
            if fechamento_anterior is None:
                tr = maxima - minima
            else:
                tr = max(
                    maxima - minima,
                    abs(maxima - fechamento_anterior),
                    abs(minima - fechamento_anterior),
                )
            trs.append(tr)
        fechamento_anterior = fechamento

    if len(trs) < periodo:
        return None

    atr = sum(trs[:periodo]) / periodo
    for tr in trs[periodo:]:
        atr = (atr * (periodo - 1) + tr) / periodo
    return round(atr, 2)


def calcular_suporte_resistencia(minimas: list, maximas: list, preco_atual: float, janela: int = 20) -> dict | None:
    """
    Suporte = menor mínima e resistência = maior máxima dos `janela` pregões
    ANTERIORES ao atual (o próprio pregão de hoje fica de fora da conta, pra
    não contaminar o nível com o rompimento que ele mesmo pode estar fazendo
    agora). None sem pelo menos `janela` pregões anteriores com máxima E
    mínima registradas (ver backfill_historico_cotacoes).
    """
    pares = [
        (float(minima), float(maxima))
        for minima, maxima in zip(minimas[:-1], maximas[:-1])
        if minima is not None and maxima is not None
    ]
    if len(pares) < janela:
        return None

    janela_pares = pares[-janela:]
    suporte = min(minima for minima, _ in janela_pares)
    resistencia = max(maxima for _, maxima in janela_pares)

    return {
        "suporte": round(suporte, 2),
        "resistencia": round(resistencia, 2),
        "rompeu_resistencia": preco_atual > resistencia,
        "perdeu_suporte": preco_atual < suporte,
    }


def _classificar_suporte_resistencia(info: dict | None) -> str:
    if info is None:
        return "DADOS_INSUFICIENTES"
    if info["rompeu_resistencia"]:
        return "ALTA"
    if info["perdeu_suporte"]:
        return "BAIXA"
    return "NEUTRO"


SUPORTE_RESISTENCIA_LABELS = {
    "ALTA": "Rompeu a resistência recente",
    "BAIXA": "Perdeu o suporte recente",
    "NEUTRO": "Dentro do canal entre suporte e resistência",
    "DADOS_INSUFICIENTES": "Aguardando histórico",
}
SUPORTE_RESISTENCIA_CLASSES = {
    "ALTA": "alta", "BAIXA": "baixa", "NEUTRO": "neutro", "DADOS_INSUFICIENTES": "neutro",
}


def sugerir_plano_trade(
    preco_atual: float, veredito: str, suporte_resistencia: dict | None, atr: float | None,
) -> dict | None:
    """
    Plano de trade sugerido (entrada/stop/alvo/relação risco-retorno) - só
    calculado quando o veredito do Scanner tem uma direção predominante
    (PREDOMINIO_ALTA/PREDOMINIO_BAIXA); em CONFLITANTE/NEUTRO/DADOS_
    INSUFICIENTES não há direção pra basear entrada nenhuma. Usa o suporte/
    resistência como stop/alvo quando disponível, caindo pra um buffer
    baseado no ATR (ou 2% do preço, se o ATR também não estiver disponível
    ainda) quando não há suporte/resistência calculado.

    A carteira é long-only (sem venda a descoberto - ver
    Operacao.TIPO_CHOICES): no lado PREDOMINIO_BAIXA, o plano não sugere
    "abrir uma venda", e sim usar os níveis como referência para decidir se
    sai (total ou parcialmente) de uma posição já comprada.
    """
    if veredito not in ("PREDOMINIO_ALTA", "PREDOMINIO_BAIXA"):
        return None

    entrada = preco_atual
    buffer = atr if atr else entrada * 0.02

    if veredito == "PREDOMINIO_ALTA":
        stop = suporte_resistencia["suporte"] if suporte_resistencia else entrada - buffer * 1.5
        stop = min(stop, entrada - buffer * 0.5)
        tem_resistencia_acima = suporte_resistencia and suporte_resistencia["resistencia"] > entrada
        alvo = suporte_resistencia["resistencia"] if tem_resistencia_acima else entrada + buffer * 2
        acao_label = "Considere comprar / reforçar a posição"
        observacao = None
    else:
        stop = suporte_resistencia["resistencia"] if suporte_resistencia else entrada + buffer * 1.5
        stop = max(stop, entrada + buffer * 0.5)
        tem_suporte_abaixo = suporte_resistencia and suporte_resistencia["suporte"] < entrada
        alvo = suporte_resistencia["suporte"] if tem_suporte_abaixo else entrada - buffer * 2
        acao_label = "Considere reduzir ou sair da posição"
        observacao = (
            "Como a carteira não opera venda a descoberto, use estes níveis como "
            "referência para decidir se sai (total ou parcialmente) da posição já comprada."
        )

    stop = round(max(stop, 0.01), 2)
    alvo = round(max(alvo, 0.01), 2)
    risco = abs(entrada - stop)
    retorno = abs(alvo - entrada)

    return {
        "acao_label": acao_label,
        "observacao": observacao,
        "entrada": round(entrada, 2),
        "stop": stop,
        "alvo": alvo,
        "relacao_risco_retorno": round(retorno / risco, 2) if risco else None,
    }


VEREDITO_SCANNER_LABELS = {
    "PREDOMINIO_ALTA": "Predomínio de sinais de alta",
    "PREDOMINIO_BAIXA": "Predomínio de sinais de baixa",
    "CONFLITANTE": "Indicadores conflitantes - sem sinal claro",
    "NEUTRO": "Neutro - nenhum indicador aponta direção",
    "DADOS_INSUFICIENTES": "Aguardando histórico suficiente",
}
VEREDITO_SCANNER_CLASSES = {
    # CONFLITANTE fica sozinho com o âmbar (badge-neutro) - é o estado que
    # mais precisa chamar atenção nessa tela; NEUTRO/DADOS_INSUFICIENTES usam
    # o cinza (badge-dados-insuficientes) para não competir visualmente com
    # ele, já que os dois não têm nada de "conflito" pra sinalizar.
    "PREDOMINIO_ALTA": "alta", "PREDOMINIO_BAIXA": "baixa", "CONFLITANTE": "neutro",
    "NEUTRO": "dados-insuficientes", "DADOS_INSUFICIENTES": "dados-insuficientes",
}


def _veredito_scanner(chaves: list[str]) -> str:
    """
    A regra do veredito do Scanner Técnico, isolada dos indicadores em si pra
    poder ser testada direto: CONFLITANTE sempre que há pelo menos um "ALTA" E
    um "BAIXA" ao mesmo tempo entre as chaves (não importa quantos "NEUTRO"/
    "DADOS_INSUFICIENTES" também existam) - só fecha PREDOMINIO_ALTA/
    PREDOMINIO_BAIXA quando todas as chaves que apontaram uma direção
    concordam entre si. DADOS_INSUFICIENTES só quando TODAS as chaves são
    "DADOS_INSUFICIENTES" (nenhum indicador tem histórico suficiente ainda).
    """
    votos_alta = chaves.count("ALTA")
    votos_baixa = chaves.count("BAIXA")
    votos_sem_dados = chaves.count("DADOS_INSUFICIENTES")

    if votos_sem_dados == len(chaves):
        return "DADOS_INSUFICIENTES"
    if votos_alta > 0 and votos_baixa > 0:
        return "CONFLITANTE"
    if votos_alta > 0:
        return "PREDOMINIO_ALTA"
    if votos_baixa > 0:
        return "PREDOMINIO_BAIXA"
    return "NEUTRO"


def escanear_ativo_precos(
    precos_cronologico: list[float], volumes_cronologico: list | None = None,
    janela_curta: int = 3, janela_longa: int = 10,
    maximas_cronologico: list | None = None, minimas_cronologico: list | None = None,
) -> dict:
    """
    Núcleo puro (sem acesso ao banco) do Scanner Técnico: combina IFR (RSI),
    médias móveis, MACD, volume, suporte/resistência e tendência de curto
    prazo num só veredito. Ver o comentário da seção acima para a regra do
    veredito (CONFLITANTE sempre que houver indicadores discordando entre
    si). ATR e volatilidade são contexto de risco, não entram na votação.
    Quando o veredito tem uma direção predominante, também sugere um plano de
    trade (entrada/stop/alvo/relação risco-retorno - ver sugerir_plano_trade).
    """
    if volumes_cronologico is None:
        volumes_cronologico = [None] * len(precos_cronologico)
    if maximas_cronologico is None:
        maximas_cronologico = [None] * len(precos_cronologico)
    if minimas_cronologico is None:
        minimas_cronologico = [None] * len(precos_cronologico)

    rsi = calcular_rsi(precos_cronologico)
    rsi_chave = _classificar_rsi(rsi)

    medias = calcular_medias_moveis(precos_cronologico)
    medias_chave = _classificar_medias_moveis(medias)

    macd = calcular_macd(precos_cronologico)
    macd_chave = macd["cruzamento"] if macd else "DADOS_INSUFICIENTES"

    volume_info = analisar_volume_precos(volumes_cronologico, precos_cronologico)
    volume_chave = volume_info["classificacao"]

    preco_atual = precos_cronologico[-1] if precos_cronologico else None
    suporte_resistencia = (
        calcular_suporte_resistencia(minimas_cronologico, maximas_cronologico, preco_atual)
        if preco_atual is not None else None
    )
    suporte_resistencia_chave = _classificar_suporte_resistencia(suporte_resistencia)

    tendencia_chave = _tendencia_a_partir_de_precos_desc(
        list(reversed(precos_cronologico))[:janela_longa], janela_curta, janela_longa
    )

    volatilidade_metricas = _volatilidade_e_drawdown_precos(precos_cronologico)
    volatilidade_chave = _classificar_volatilidade(volatilidade_metricas["volatilidade_pct"])

    atr = calcular_atr(maximas_cronologico, minimas_cronologico, precos_cronologico)

    indicadores = [
        {
            "nome": "IFR (RSI)", "chave": RSI_CHAVE_PARA_VOTO_SCANNER[rsi_chave],
            "label": RSI_LABELS[rsi_chave], "classe": RSI_CLASSES[rsi_chave],
        },
        {
            "nome": "Médias móveis", "chave": medias_chave,
            "label": MEDIAS_MOVEIS_LABELS[medias_chave], "classe": MEDIAS_MOVEIS_CLASSES[medias_chave],
        },
        {"nome": "MACD", "chave": macd_chave, "label": MACD_LABELS[macd_chave], "classe": MACD_CLASSES[macd_chave]},
        {"nome": "Volume", "chave": volume_chave, "label": VOLUME_LABELS[volume_chave], "classe": VOLUME_CLASSES[volume_chave]},
        {
            "nome": "Suporte e resistência", "chave": suporte_resistencia_chave,
            "label": SUPORTE_RESISTENCIA_LABELS[suporte_resistencia_chave],
            "classe": SUPORTE_RESISTENCIA_CLASSES[suporte_resistencia_chave],
        },
        {
            "nome": "Tendência de curto prazo", "chave": tendencia_chave,
            "label": TENDENCIA_LABELS.get(tendencia_chave, tendencia_chave),
            "classe": tendencia_chave.lower() if tendencia_chave != "DADOS_INSUFICIENTES" else "neutro",
        },
    ]

    chaves = [item["chave"] for item in indicadores]
    votos_alta = chaves.count("ALTA")
    votos_baixa = chaves.count("BAIXA")
    votos_sem_dados = chaves.count("DADOS_INSUFICIENTES")
    veredito = _veredito_scanner(chaves)

    plano_trade = (
        sugerir_plano_trade(preco_atual, veredito, suporte_resistencia, atr)
        if preco_atual is not None else None
    )

    return {
        "indicadores": indicadores,
        "rsi": rsi,
        "medias_moveis": medias,
        "macd": macd,
        "volume": volume_info,
        "suporte_resistencia": suporte_resistencia,
        "suporte_resistencia_label": SUPORTE_RESISTENCIA_LABELS[suporte_resistencia_chave],
        "suporte_resistencia_classe": SUPORTE_RESISTENCIA_CLASSES[suporte_resistencia_chave],
        "atr": atr,
        "volatilidade_pct": volatilidade_metricas["volatilidade_pct"],
        "drawdown_pct": volatilidade_metricas["drawdown_pct"],
        "volatilidade_chave": volatilidade_chave,
        "volatilidade_label": VOLATILIDADE_LABELS[volatilidade_chave],
        "volatilidade_classe": VOLATILIDADE_CLASSES[volatilidade_chave],
        "tendencia_label": TENDENCIA_LABELS.get(tendencia_chave, tendencia_chave),
        "votos_alta": votos_alta,
        "votos_baixa": votos_baixa,
        "votos_sem_dados": votos_sem_dados,
        "veredito": veredito,
        "veredito_label": VEREDITO_SCANNER_LABELS[veredito],
        "veredito_classe": VEREDITO_SCANNER_CLASSES[veredito],
        "plano_trade": plano_trade,
    }


def escanear_ativo(ativo: Ativo, janela_curta: int = 3, janela_longa: int = 10) -> dict:
    """Versão do Scanner Técnico que busca preços, volumes, máximas e mínimas já salvos (Cotacao) de um ativo."""
    cotacoes = list(
        ativo.cotacoes.order_by("data").values_list("preco_fechamento", "volume", "maxima", "minima")
    )
    precos = [float(preco) for preco, _, _, _ in cotacoes]
    volumes = [volume for _, volume, _, _ in cotacoes]
    maximas = [maxima for _, _, maxima, _ in cotacoes]
    minimas = [minima for _, _, _, minima in cotacoes]
    resultado = escanear_ativo_precos(precos, volumes, janela_curta, janela_longa, maximas, minimas)
    resultado["ativo"] = ativo
    return resultado


def escanear_carteira(usuario) -> list[dict]:
    """
    Scanner Técnico de todos os ativos com posição em carteira do usuário -
    tanto os realmente comprados (saldo > 0) quanto os apenas reservados
    (watchlist, sem saldo comprado) - ordenados por ticker - ver
    core.views.scanner_tecnico. calcular_posicoes já descarta posições
    totalmente zeradas (nem comprada, nem reservada), então não é preciso
    filtrar de novo aqui além de ordenar.
    """
    posicoes = sorted(calcular_posicoes(usuario), key=lambda p: p.ativo.ticker)
    resultados = []
    for posicao in posicoes:
        resultado = escanear_ativo(posicao.ativo)
        resultado["apenas_reservado"] = posicao.apenas_reservado
        resultados.append(resultado)
    return resultados


COLUNAS_SCANNER = [
    "Ativo", "Reservado", "Veredito", "IFR (RSI)", "Médias móveis", "MACD", "Volume",
    "Suporte e resistência", "Tendência", "Volatilidade (30d, %)", "ATR (14d, R$)",
    "Suporte (R$)", "Resistência (R$)", "Entrada", "Stop", "Alvo", "Relação R:R",
]


def _linha_scanner(item: dict) -> list:
    indicadores = {i["nome"]: i["label"] for i in item["indicadores"]}
    plano = item.get("plano_trade")
    sr = item.get("suporte_resistencia")
    return [
        item["ativo"].ticker,
        "Sim" if item.get("apenas_reservado") else "Não",
        item["veredito_label"],
        indicadores.get("IFR (RSI)", "—"),
        indicadores.get("Médias móveis", "—"),
        indicadores.get("MACD", "—"),
        indicadores.get("Volume", "—"),
        indicadores.get("Suporte e resistência", "—"),
        indicadores.get("Tendência de curto prazo", "—"),
        item["volatilidade_pct"] if item["volatilidade_pct"] is not None else "—",
        item["atr"] if item["atr"] is not None else "—",
        sr["suporte"] if sr else "—",
        sr["resistencia"] if sr else "—",
        plano["entrada"] if plano else "—",
        plano["stop"] if plano else "—",
        plano["alvo"] if plano else "—",
        f"1:{plano['relacao_risco_retorno']}" if plano and plano.get("relacao_risco_retorno") else "—",
    ]


def _ranking_tecnico_scanner(resultados_scanner: list[dict]) -> list[dict]:
    """
    Ranking técnico dos ativos acompanhados para os relatórios Scanner + IA:
    mesma pontuação (ver _score_scanner) e ordem (maior primeiro) do ranking
    da tela Análise de Mercado, mas calculado direto do Scanner Técnico - não
    depende da IA, então sai no relatório mesmo sem IA configurada.
    """
    ranking = [
        {
            "ticker": item["ativo"].ticker,
            "score": _score_scanner(item),
            "veredito_label": item["veredito_label"],
        }
        for item in resultados_scanner
    ]
    ranking = [linha for linha in ranking if linha["score"] is not None]
    return sorted(ranking, key=lambda linha: linha["score"], reverse=True)


def _sem_emoji(texto: str) -> str:
    """Tira emojis/símbolos (✅, ⚠, 🛑...) de um texto pro PDF - a fonte padrão do reportlab não os desenha."""
    return re.sub(r"[☀-➿️\U0001F000-\U0001FAFF]", "", texto).strip()


def _texto_ativo_analise_carteira(a: dict) -> dict:
    """Valores já formatados de uma linha de "Ativo por ativo" (ver analisar_carteira_comprada) para os relatórios."""
    return {
        "ativo": a["ticker"] + (f" ({a['setor']})" if a["setor"] else ""),
        "peso": f"{a['peso_pct']}%" if a["peso_pct"] is not None else "—",
        "resultado": f"{a['lucro_pct']:+.2f}%" if a["lucro_pct"] is not None else "—",
        "scanner": f"{a['veredito_label']} ({a['score']}/100)",
        "recomendacao": a["recomendacao_label"],
        "motivos": " ".join(a["motivos"]),
    }


def _abas_excel_analise_carteira(wb, analise: dict | None, analise_ia: dict | None, erro_ia: str | None) -> None:
    """Abas "Análise da carteira" (resumo + ativo por ativo) e "Análise da carteira (IA)" (+ próximos passos) dos relatórios Excel."""
    from openpyxl.styles import Alignment, Font, PatternFill

    def _cabecalho(ws, valores):
        ws.append(valores)
        for celula in ws[ws.max_row]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")

    def _titulo(ws, texto):
        ws.append([texto])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)

    gerado_em = f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
    ws = wb.create_sheet("Análise da carteira")
    _titulo(ws, "Análise dos ativos comprados")
    ws.append([gerado_em])
    ws.append([
        "Calculada pelo sistema com as cotações e o histórico já salvos (resultado e metas, Scanner Técnico, "
        "volatilidade, peso, qualidade, setores, correlação, beta, Ibovespa/CDI) - não constitui recomendação de investimento.",
    ])
    ws.append([])
    if not analise:
        ws.append(["Nenhuma ação comprada para analisar."])
    else:
        ws.append(["Saúde da carteira", f"{analise['nota']}/100 - {analise['classificacao']}"])
        ws.append(["Ativos analisados", analise["total"]])
        ws.append(["Em lucro / em prejuízo", f"{analise['em_lucro']} / {analise['em_perda']}"])
        if analise["melhor"]:
            ws.append(["Melhor posição", f"{analise['melhor']['ticker']} {analise['melhor']['lucro_pct']:+.2f}%"])
        if analise["pior"]:
            ws.append(["Pior posição", f"{analise['pior']['ticker']} {analise['pior']['lucro_pct']:+.2f}%"])
        ws.append([])
        _titulo(ws, "✓ Pontos fortes")
        for texto in analise["pontos_fortes"] or ["Nenhum destaque positivo no momento."]:
            ws.append(["", texto])
        _titulo(ws, "⚠ Pontos de atenção")
        for texto in analise["pontos_atencao"] or ["Nada preocupante encontrado."]:
            ws.append(["", texto])
        ws.append([])
        _titulo(ws, "Ativo por ativo")
        _cabecalho(ws, ["Ativo", "% da carteira", "Resultado", "Scanner", "Recomendação", "Por quê"])
        for a in analise["ativos"]:
            t = _texto_ativo_analise_carteira(a)
            ws.append([t["ativo"], t["peso"], t["resultado"], t["scanner"], t["recomendacao"], t["motivos"]])
            ws.cell(row=ws.max_row, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        for coluna, largura in zip("ABCDEF", (24, 16, 12, 26, 34, 90)):
            ws.column_dimensions[coluna].width = largura

    ws_ia = wb.create_sheet("Análise da carteira (IA)")
    _titulo(ws_ia, "Análise da carteira por IA")
    ws_ia.append([gerado_em])
    ws_ia.append(["A IA só comenta os números calculados pelo sistema - não constitui recomendação de investimento."])
    ws_ia.append([])
    if not analise:
        ws_ia.append(["Nenhuma ação comprada para analisar."])
    elif erro_ia:
        ws_ia.append([erro_ia])
    elif not analise_ia:
        ws_ia.append(["A Análise da carteira por IA não está configurada neste sistema."])
    elif analise_ia["carteira"] or analise_ia["ativos"] or analise_ia["dicas"]:
        if analise_ia["carteira"]:
            ws_ia.append(["Carteira", analise_ia["carteira"]["acao"], analise_ia["carteira"]["comentario"]])
            ws_ia.append([])
        if analise_ia["ativos"]:
            _cabecalho(ws_ia, ["Ativo", "Leitura da IA", "Comentário"])
            for linha in analise_ia["ativos"]:
                ws_ia.append([linha["chave"], linha["acao"], linha["comentario"]])
            ws_ia.append([])
        if analise_ia["dicas"]:
            _titulo(ws_ia, "Próximos passos sugeridos")
            for dica in analise_ia["dicas"]:
                ws_ia.append(["•", "", dica["comentario"]])
        ws_ia.column_dimensions["A"].width = 14
        ws_ia.column_dimensions["B"].width = 14
        ws_ia.column_dimensions["C"].width = 100
    else:
        ws_ia.append([analise_ia["texto"]])


def _secoes_pdf_analise_carteira(
    elementos: list, analise: dict | None, analise_ia: dict | None, erro_ia: str | None, estilos, estilo_celula,
) -> None:
    """Seções "Análise da carteira" (resumo, ativo por ativo, IA e próximos passos) dos relatórios PDF."""
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    estilo_tabela = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ])

    def _p(texto, estilo=estilo_celula):
        return Paragraph(escape(_sem_emoji(texto)), estilo)

    elementos.append(Spacer(1, 0.7 * cm))
    elementos.append(Paragraph("Análise da carteira", estilos["Heading1"]))
    elementos.append(Paragraph("Análise dos ativos comprados", estilos["Heading2"]))
    if not analise:
        elementos.append(Paragraph("Nenhuma ação comprada para analisar.", estilos["Normal"]))
        return
    elementos.append(Paragraph(
        "Calculada pelo sistema com as cotações e o histórico já salvos: resultado e metas de cada posição, Scanner "
        "Técnico, volatilidade, peso na carteira, qualidade fundamentalista, setores, correlação, beta e comparação "
        "com Ibovespa/CDI. Não é recomendação de investimento.",
        estilo_celula,
    ))
    elementos.append(Spacer(1, 0.2 * cm))
    resumo = [
        ["Saúde da carteira", f"{analise['nota']}/100 - {analise['classificacao']}"],
        ["Ativos analisados", str(analise["total"])],
        ["Em lucro / em prejuízo", f"{analise['em_lucro']} / {analise['em_perda']}"],
    ]
    if analise["melhor"]:
        resumo.append(["Melhor posição", f"{analise['melhor']['ticker']} {analise['melhor']['lucro_pct']:+.2f}%"])
    if analise["pior"]:
        resumo.append(["Pior posição", f"{analise['pior']['ticker']} {analise['pior']['lucro_pct']:+.2f}%"])
    tabela_resumo = Table(resumo, colWidths=[5 * cm, 8 * cm], hAlign="LEFT")
    tabela_resumo.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2f7")),
    ]))
    elementos.append(tabela_resumo)

    for titulo, itens, vazio in (
        ("Pontos fortes", analise["pontos_fortes"], "Nenhum destaque positivo no momento."),
        ("Pontos de atenção", analise["pontos_atencao"], "Nada preocupante encontrado."),
    ):
        elementos.append(Paragraph(titulo, estilos["Heading3"]))
        for texto in itens or [vazio]:
            elementos.append(_p(f"• {texto}"))

    elementos.append(Paragraph("Ativo por ativo", estilos["Heading2"]))
    dados = [["Ativo", "% da carteira", "Resultado", "Scanner", "Recomendação", "Por quê"]]
    for a in analise["ativos"]:
        t = _texto_ativo_analise_carteira(a)
        dados.append([
            _p(t["ativo"]), t["peso"], t["resultado"], _p(t["scanner"]), _p(t["recomendacao"]), _p(t["motivos"]),
        ])
    tabela = Table(dados, colWidths=[3.5 * cm, 2.2 * cm, 2 * cm, 3.5 * cm, 3.8 * cm, 12.3 * cm], repeatRows=1)
    tabela.setStyle(estilo_tabela)
    elementos.append(tabela)

    elementos.append(Spacer(1, 0.5 * cm))
    elementos.append(Paragraph("Análise da carteira por IA", estilos["Heading2"]))
    if erro_ia:
        elementos.append(Paragraph(escape(erro_ia), estilos["Normal"]))
    elif not analise_ia:
        elementos.append(Paragraph("A Análise da carteira por IA não está configurada neste sistema.", estilos["Normal"]))
    elif analise_ia["carteira"] or analise_ia["ativos"] or analise_ia["dicas"]:
        if analise_ia["carteira"]:
            elementos.append(_p(f"[{analise_ia['carteira']['acao']}] {analise_ia['carteira']['comentario']}", estilos["Normal"]))
            elementos.append(Spacer(1, 0.3 * cm))
        if analise_ia["ativos"]:
            dados_ia = [["Ativo", "Leitura da IA", "Comentário"]]
            for linha in analise_ia["ativos"]:
                dados_ia.append([linha["chave"], linha["acao"], _p(linha["comentario"])])
            tabela_ia = Table(dados_ia, colWidths=[3 * cm, 3 * cm, 21.3 * cm], repeatRows=1)
            tabela_ia.setStyle(estilo_tabela)
            elementos.append(tabela_ia)
        if analise_ia["dicas"]:
            elementos.append(Paragraph("Próximos passos sugeridos", estilos["Heading3"]))
            for dica in analise_ia["dicas"]:
                elementos.append(_p(f"• {dica['comentario']}", estilos["Normal"]))
    else:
        elementos.append(_p(analise_ia["texto"], estilos["Normal"]))
    elementos.append(Spacer(1, 0.2 * cm))
    elementos.append(Paragraph(
        "A IA só comenta os números calculados pelo sistema acima. Não é recomendação de investimento.", estilo_celula,
    ))


def gerar_excel_scanner_e_analise_ia(
    resultados_scanner: list[dict], analise_ia: dict | None, erro_ia: str | None,
    sugestoes_ia: dict | None = None, erro_sugestoes_ia: str | None = None,
    analise_carteira: dict | None = None, analise_carteira_ia: dict | None = None,
    erro_analise_carteira_ia: str | None = None,
) -> bytes:
    """
    Gera uma planilha .xlsx com seis abas: o Scanner Técnico completo (cada
    indicador, suporte/resistência, ATR e o plano de trade sugerido) de todo
    ativo em carteira, o Ranking técnico dos ativos acompanhados (ver
    _ranking_tecnico_scanner), a Análise da B3 hoje por IA (ver gerar_analise_b3_ia),
    as Sugestões de hoje por IA (ver gerar_sugestoes_ia), quando configuradas,
    e a Análise da carteira - local (ver analisar_carteira_comprada) e por IA
    (ver gerar_analise_carteira_ia) - `analise_ia`/`erro_ia`/`sugestoes_ia`/...
    vêm de chamadas já feitas por quem chama esta função (não recalcula
    sozinha, pra não repetir a mesma consulta à IA de graça).
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Scanner Técnico"
    ws.append(["Scanner Técnico"])
    ws.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws.append(["Referência de apoio à leitura técnica - não constitui recomendação de investimento."])
    ws.append([])
    ws.append(COLUNAS_SCANNER)
    for celula in ws[5]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D1526")
        celula.alignment = Alignment(horizontal="center")

    for item in resultados_scanner:
        ws.append(_linha_scanner(item))

    for indice in range(1, len(COLUNAS_SCANNER) + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 16
    ws.freeze_panes = "A6"

    ws_ranking = wb.create_sheet("Ranking técnico")
    ws_ranking.append(["Ranking técnico dos ativos acompanhados"])
    ws_ranking.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws_ranking.append([
        "Ordenado pela pontuação do Scanner Técnico (0 a 100, calculada localmente) - "
        "não é uma previsão nem recomendação de investimento.",
    ])
    ws_ranking.append([])
    ranking = _ranking_tecnico_scanner(resultados_scanner)
    if ranking:
        ws_ranking.append(["Posição", "Ativo", "Pontuação", "Veredito"])
        for celula in ws_ranking[5]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")
        for posicao, linha in enumerate(ranking, start=1):
            ws_ranking.append([posicao, linha["ticker"], f"{linha['score']}/100", linha["veredito_label"]])
        ws_ranking.column_dimensions["A"].width = 10
        ws_ranking.column_dimensions["B"].width = 14
        ws_ranking.column_dimensions["C"].width = 12
        ws_ranking.column_dimensions["D"].width = 22
    else:
        ws_ranking.append(["Nenhum ativo com pontuação técnica para ranquear."])

    ws2 = wb.create_sheet("Análise da B3 (IA)")
    ws2.append(["Análise da B3 hoje (IA)"])
    ws2.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws2.append([
        "Comentário gerado por IA a partir de dados já calculados pelo sistema - "
        "não constitui recomendação de investimento.",
    ])
    ws2.append([])
    if erro_ia:
        ws2.append([erro_ia])
    elif not analise_ia:
        ws2.append(["A Análise da B3 hoje por IA não está configurada neste sistema."])
    elif analise_ia["linhas"]:
        ws2.append(["Ativo", "Situação", "Comentário"])
        for celula in ws2[5]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")
        for linha in analise_ia["linhas"]:
            ws2.append([
                "B3 (mercado geral)" if linha["eh_mercado_geral"] else linha["ticker"],
                linha["situacao"],
                linha["comentario"],
            ])
        ws2.column_dimensions["A"].width = 22
        ws2.column_dimensions["B"].width = 14
        ws2.column_dimensions["C"].width = 80
    else:
        ws2.append([analise_ia["texto"]])

    ws3 = wb.create_sheet("Sugestões de hoje (IA)")
    ws3.append(["Sugestões de hoje (IA)"])
    ws3.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws3.append([
        "Sugestões geradas por IA a partir das maiores altas/baixas do dia - "
        "não constitui recomendação de investimento.",
    ])
    ws3.append([])
    if erro_sugestoes_ia:
        ws3.append([erro_sugestoes_ia])
    elif not sugestoes_ia:
        ws3.append(["As Sugestões de hoje por IA não estão configuradas neste sistema."])
    elif sugestoes_ia["linhas"]:
        ws3.append(["Ativo", "Justificativa"])
        for celula in ws3[5]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center")
        for linha in sugestoes_ia["linhas"]:
            ws3.append([linha["ticker"], linha["justificativa"]])
        ws3.column_dimensions["A"].width = 16
        ws3.column_dimensions["B"].width = 90
    else:
        ws3.append([sugestoes_ia["texto"]])

    _abas_excel_analise_carteira(wb, analise_carteira, analise_carteira_ia, erro_analise_carteira_ia)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_scanner_e_analise_ia(
    resultados_scanner: list[dict], analise_ia: dict | None, erro_ia: str | None, nome_usuario: str,
    sugestoes_ia: dict | None = None, erro_sugestoes_ia: str | None = None,
    analise_carteira: dict | None = None, analise_carteira_ia: dict | None = None,
    erro_analise_carteira_ia: str | None = None,
) -> bytes:
    """
    Gera um PDF com o Scanner Técnico (veredito, suporte/resistência, ATR e
    plano de trade sugerido) de todo ativo em carteira, o Ranking técnico dos
    ativos acompanhados, a Análise da B3 hoje por IA, as Sugestões de hoje
    por IA, quando configuradas, e a Análise da carteira (ativo por ativo, IA
    e próximos passos) - mesma observação de gerar_excel_scanner_e_analise_ia sobre
    os parâmetros de IA já virem prontos de quem chama.
    """
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()
    estilo_celula = ParagraphStyle("celula_scanner_ia", parent=estilos["Normal"], fontSize=8, leading=10)

    elementos = [
        Paragraph(
            f"BolsaTrader - Scanner Técnico e Análise da B3 {timezone.localdate().strftime('%d/%m/%Y')}",
            estilos["Title"],
        ),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')} - "
            "referência de apoio à decisão, não constitui recomendação de investimento.",
            estilos["Normal"],
        ),
        Spacer(1, 0.4 * cm),
        Paragraph("Scanner Técnico", estilos["Heading2"]),
    ]

    if not resultados_scanner:
        elementos.append(Paragraph("Nenhum ativo em carteira para escanear.", estilos["Normal"]))
    else:
        colunas_pdf = ["Ativo", "Reservado", "Veredito", "Suporte", "Resistência", "ATR", "Entrada", "Stop", "Alvo", "R:R"]
        dados = [colunas_pdf]
        for item in resultados_scanner:
            plano = item.get("plano_trade")
            sr = item.get("suporte_resistencia")
            dados.append([
                item["ativo"].ticker,
                "Sim" if item.get("apenas_reservado") else "Não",
                item["veredito_label"],
                f"R$ {sr['suporte']}" if sr else "—",
                f"R$ {sr['resistencia']}" if sr else "—",
                f"R$ {item['atr']}" if item["atr"] is not None else "—",
                f"R$ {plano['entrada']}" if plano else "—",
                f"R$ {plano['stop']}" if plano else "—",
                f"R$ {plano['alvo']}" if plano else "—",
                f"1:{plano['relacao_risco_retorno']}" if plano and plano.get("relacao_risco_retorno") else "—",
            ])
        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela)

    elementos.append(Spacer(1, 0.7 * cm))
    elementos.append(Paragraph("Ranking técnico dos ativos acompanhados", estilos["Heading2"]))
    elementos.append(Paragraph(
        "Ordenado pela pontuação do Scanner Técnico (0 a 100, calculada localmente a partir de IFR, médias "
        "móveis, MACD, volume, suporte/resistência e tendência) - não é uma previsão, só reflete quantos "
        "indicadores apontam para cada lado agora.",
        estilo_celula,
    ))
    elementos.append(Spacer(1, 0.2 * cm))
    ranking = _ranking_tecnico_scanner(resultados_scanner)
    if ranking:
        dados_ranking = [["Posição", "Ativo", "Pontuação", "Veredito"]]
        for posicao, linha in enumerate(ranking, start=1):
            dados_ranking.append([f"{posicao}º", linha["ticker"], f"{linha['score']}/100", linha["veredito_label"]])
        tabela_ranking = Table(dados_ranking, colWidths=[2 * cm, 3.5 * cm, 3 * cm, 5 * cm], repeatRows=1, hAlign="LEFT")
        tabela_ranking.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela_ranking)
    else:
        elementos.append(Paragraph("Nenhum ativo com pontuação técnica para ranquear.", estilos["Normal"]))

    elementos.append(Spacer(1, 0.7 * cm))
    elementos.append(Paragraph("Análise da B3 hoje (IA)", estilos["Heading2"]))

    if erro_ia:
        elementos.append(Paragraph(escape(erro_ia), estilos["Normal"]))
    elif not analise_ia:
        elementos.append(Paragraph(
            "A Análise da B3 hoje por IA não está configurada neste sistema.", estilos["Normal"],
        ))
    elif analise_ia["linhas"]:
        dados_ia = [["Ativo", "Situação", "Comentário"]]
        for linha in analise_ia["linhas"]:
            ativo_label = "B3 (mercado geral)" if linha["eh_mercado_geral"] else linha["ticker"]
            dados_ia.append([
                ativo_label,
                linha["situacao"],
                Paragraph(escape(linha["comentario"]), estilo_celula),
            ])
        tabela_ia = Table(dados_ia, colWidths=[3.5 * cm, 2.5 * cm, 18 * cm], repeatRows=1)
        tabela_ia.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (0, -1), 8),
            ("FONTSIZE", (1, 0), (1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela_ia)
    else:
        elementos.append(Paragraph(escape(analise_ia["texto"]), estilos["Normal"]))

    elementos.append(Spacer(1, 0.7 * cm))
    elementos.append(Paragraph("Sugestões de hoje (IA)", estilos["Heading2"]))

    if erro_sugestoes_ia:
        elementos.append(Paragraph(escape(erro_sugestoes_ia), estilos["Normal"]))
    elif not sugestoes_ia:
        elementos.append(Paragraph(
            "As Sugestões de hoje por IA não estão configuradas neste sistema.", estilos["Normal"],
        ))
    elif sugestoes_ia["linhas"]:
        dados_sugestoes = [["Ativo", "Justificativa"]]
        for linha in sugestoes_ia["linhas"]:
            dados_sugestoes.append([linha["ticker"], Paragraph(escape(linha["justificativa"]), estilo_celula)])
        tabela_sugestoes = Table(dados_sugestoes, colWidths=[3.5 * cm, 20.5 * cm], repeatRows=1)
        tabela_sugestoes.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (0, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela_sugestoes)
    else:
        elementos.append(Paragraph(escape(sugestoes_ia["texto"]), estilos["Normal"]))

    _secoes_pdf_analise_carteira(
        elementos, analise_carteira, analise_carteira_ia, erro_analise_carteira_ia, estilos, estilo_celula,
    )

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Backtesting do robô consultor
#
# Responde à pergunta que o robô sozinho não responde: "se eu tivesse
# seguido todo sinal de compra/venda que o robô já deu para este ativo,
# quantas vezes eu teria acertado?" - percorre o histórico de Cotacao já
# salvo recalculando, em cada pregão, o mesmo sinal que analisar_indicadores_
# tecnicos_precos calcularia "ao vivo" naquele dia (usando só os preços
# disponíveis até ali, nunca os futuros), e confere contra o retorno real
# `dias_retorno` pregões à frente. É só uma estatística sobre o passado, não
# garante desempenho futuro - a mesma ressalva do robô ao vivo.
# --------------------------------------------------------------------------
# Abaixo disso, a taxa de acerto é mais sorte que estatística.
BACKTEST_MINIMO_SINAIS_CONFIAVEL = 20
# Diferença mínima (pontos percentuais) de taxa de acerto sobre a referência
# para dizer que o robô "supera" (ou "fica abaixo de") comprar em qualquer dia.
BACKTEST_MARGEM_VANTAGEM_PP = 5
# Retorno médio por operação (pontos percentuais) abaixo da referência a
# partir do qual o robô "fica abaixo" mesmo acertando mais vezes - acertar
# muito e ganhar pouco, com perdas grandes nos erros, não compensa.
BACKTEST_MARGEM_RETORNO_PP = 0.5


def _veredito_backtest(sinais_compra: int, vantagem_pp, retorno_compra, retorno_base) -> str:
    """
    Compara os sinais de COMPRA do robô com a referência (comprar em qualquer
    pregão do mesmo período): POUCOS_DADOS, SUPERA (acerta bem mais e ganha
    pelo menos o mesmo), ABAIXO (acerta bem menos, ou ganha bem menos por
    operação) ou EMPATA (diferença pequena demais para separar do acaso).
    """
    if sinais_compra < BACKTEST_MINIMO_SINAIS_CONFIAVEL or vantagem_pp is None:
        return "POUCOS_DADOS"
    diferenca_retorno = retorno_compra - retorno_base
    if vantagem_pp <= -BACKTEST_MARGEM_VANTAGEM_PP or diferenca_retorno <= -BACKTEST_MARGEM_RETORNO_PP:
        return "ABAIXO"
    if vantagem_pp >= BACKTEST_MARGEM_VANTAGEM_PP and diferenca_retorno >= 0:
        return "SUPERA"
    return "EMPATA"


def backtest_sinais_robo(
    ativo: Ativo, dias_retorno: int = 5, janela_curta: int = 3, janela_longa: int = 10,
    minimo_precos: int = 15, custo_pct: float | None = None,
) -> dict:
    """
    Simula, dia a dia, os sinais que o robô consultor teria dado para
    `ativo` no passado (com base só no histórico de Cotacao já salvo) e mede
    a taxa de acerto: um sinal de COMPRA "acerta" quando o preço `dias_retorno`
    pregões depois está mais alto; um sinal de VENDA "acerta" quando está mais
    baixo. Sinais NEUTRO não entram na contagem (não haveria ação a avaliar).

    Os retornos já descontam `custo_pct` (compra + venda: corretagem e
    emolumentos; padrão BACKTEST_CUSTO_PCT do .env). E a taxa de acerto da
    COMPRA é comparada com a REFERÊNCIA: comprar em qualquer pregão do mesmo
    período, sem sinal nenhum. Em mercado de alta quase toda compra "acerta" -
    o robô só tem valor se acertar mais que isso (ver _veredito_backtest).

    `minimo_precos` é o menor histórico a partir do qual já vale a pena
    simular um sinal (por padrão, o mínimo do RSI - com menos que isso o
    sinal geral tende a ficar sempre neutro por falta de dados); o MACD
    completo só passa a contar a partir de 35 pregões acumulados, mas isso é
    tratado automaticamente por analisar_indicadores_tecnicos_precos (fica
    "aguardando histórico" e simplesmente não vota até lá).
    """
    if custo_pct is None:
        custo_pct = settings.BACKTEST_CUSTO_PCT
    cotacoes = list(ativo.cotacoes.order_by("data"))
    precos = [float(c.preco_fechamento) for c in cotacoes]
    datas = [c.data for c in cotacoes]
    total_pregoes = len(precos)

    sinais = []
    retornos_referencia = []
    for i in range(minimo_precos - 1, total_pregoes):
        if i + dias_retorno >= total_pregoes:
            break  # não há pregões suficientes à frente pra conferir o resultado deste sinal ainda

        preco_no_sinal = precos[i]
        if not preco_no_sinal:
            continue
        retorno_bruto = (precos[i + dias_retorno] - preco_no_sinal) / preco_no_sinal * 100
        # referência: comprar neste pregão sem olhar sinal nenhum
        retornos_referencia.append(retorno_bruto - custo_pct)

        indicadores = analisar_indicadores_tecnicos_precos(precos[: i + 1], janela_curta, janela_longa)
        sinal = indicadores["sinal_geral"]
        if sinal not in ("COMPRA", "VENDA"):
            continue

        # o custo pesa contra nos dois sentidos: a compra precisa subir mais
        # que o custo, a venda (zerar/evitar posição) precisa cair mais que ele
        retorno_pct = retorno_bruto - custo_pct if sinal == "COMPRA" else retorno_bruto + custo_pct
        acerto = (sinal == "COMPRA" and retorno_pct > 0) or (sinal == "VENDA" and retorno_pct < 0)
        sinais.append({
            "data": datas[i], "sinal": sinal, "preco": round(preco_no_sinal, 2),
            "retorno_pct": round(retorno_pct, 2), "acerto": acerto,
        })

    referencia_taxa_acerto = (
        round(sum(1 for r in retornos_referencia if r > 0) / len(retornos_referencia) * 100, 1)
        if retornos_referencia else None
    )
    referencia_retorno_medio = (
        round(sum(retornos_referencia) / len(retornos_referencia), 2) if retornos_referencia else None
    )
    comum = {
        "ativo": ativo,
        "dias_retorno": dias_retorno,
        "custo_pct": custo_pct,
        "total_pregoes": total_pregoes,
        "data_inicial": datas[0] if datas else None,
        "data_final": datas[-1] if datas else None,
        "referencia_taxa_acerto_pct": referencia_taxa_acerto,
        "referencia_retorno_medio_pct": referencia_retorno_medio,
        "referencia_total": len(retornos_referencia),
        # somas usadas pelo resumo geral de todos os ativos (resumir_backtests)
        "acertos_compra": sum(1 for s in sinais if s["sinal"] == "COMPRA" and s["acerto"]),
        "soma_retorno_compra": sum(s["retorno_pct"] for s in sinais if s["sinal"] == "COMPRA"),
        "acertos_referencia": sum(1 for r in retornos_referencia if r > 0),
        "soma_retorno_referencia": sum(retornos_referencia),
    }

    if not sinais:
        return {
            **comum,
            "total_sinais": 0,
            "sinais_compra": 0, "sinais_venda": 0, "taxa_acerto_pct": None,
            "taxa_acerto_compra_pct": None, "taxa_acerto_venda_pct": None,
            "retorno_medio_compra_pct": None, "retorno_medio_venda_pct": None,
            "vantagem_compra_pp": None, "veredito": "POUCOS_DADOS",
            "sinais": [], "dados_insuficientes": total_pregoes < minimo_precos + dias_retorno,
        }

    def _media_pct(itens):
        valores = [s["retorno_pct"] for s in itens]
        return round(sum(valores) / len(valores), 2) if valores else None

    def _taxa_acerto(itens):
        return round(sum(1 for s in itens if s["acerto"]) / len(itens) * 100, 1) if itens else None

    sinais_compra = [s for s in sinais if s["sinal"] == "COMPRA"]
    sinais_venda = [s for s in sinais if s["sinal"] == "VENDA"]
    taxa_compra = _taxa_acerto(sinais_compra)
    retorno_compra = _media_pct(sinais_compra)
    vantagem = (
        round(taxa_compra - referencia_taxa_acerto, 1)
        if taxa_compra is not None and referencia_taxa_acerto is not None else None
    )

    return {
        **comum,
        "total_sinais": len(sinais),
        "sinais_compra": len(sinais_compra),
        "sinais_venda": len(sinais_venda),
        "taxa_acerto_pct": _taxa_acerto(sinais),
        "taxa_acerto_compra_pct": taxa_compra,
        "taxa_acerto_venda_pct": _taxa_acerto(sinais_venda),
        "retorno_medio_compra_pct": retorno_compra,
        "retorno_medio_venda_pct": _media_pct(sinais_venda),
        "vantagem_compra_pp": vantagem,
        "veredito": _veredito_backtest(len(sinais_compra), vantagem, retorno_compra, referencia_retorno_medio),
        # mais recentes primeiro, limitado pra não pesar a página com anos de sinais
        "sinais": list(reversed(sinais))[:30],
        "dados_insuficientes": False,
    }


def resumir_backtests(resultados: list[dict]) -> dict | None:
    """
    Junta os backtests de todos os ativos num resultado só: todos os sinais
    de COMPRA do robô contra todos os pregões de referência. Responde "no
    geral, seguir o robô é melhor que comprar em qualquer dia?" com bem mais
    amostras que cada ativo sozinho. None se nenhum pregão foi avaliado.
    """
    total_referencia = sum(r["referencia_total"] for r in resultados)
    if not total_referencia:
        return None
    total_compra = sum(r["sinais_compra"] for r in resultados)
    taxa_compra = (
        round(sum(r["acertos_compra"] for r in resultados) / total_compra * 100, 1) if total_compra else None
    )
    retorno_compra = (
        round(sum(r["soma_retorno_compra"] for r in resultados) / total_compra, 2) if total_compra else None
    )
    taxa_referencia = round(sum(r["acertos_referencia"] for r in resultados) / total_referencia * 100, 1)
    retorno_referencia = round(sum(r["soma_retorno_referencia"] for r in resultados) / total_referencia, 2)
    vantagem = round(taxa_compra - taxa_referencia, 1) if taxa_compra is not None else None
    return {
        "ativos": sum(1 for r in resultados if r["referencia_total"]),
        "sinais_compra": total_compra,
        "taxa_acerto_compra_pct": taxa_compra,
        "retorno_medio_compra_pct": retorno_compra,
        "referencia_total": total_referencia,
        "referencia_taxa_acerto_pct": taxa_referencia,
        "referencia_retorno_medio_pct": retorno_referencia,
        "vantagem_compra_pp": vantagem,
        "veredito": _veredito_backtest(total_compra, vantagem, retorno_compra, retorno_referencia),
    }


# --------------------------------------------------------------------------
# Calculadora de risco da compra (tela Nova operação)
#
# Gestão de risco pesa mais no resultado que acertar a direção: com perdas
# pequenas e controladas, dá para ganhar dinheiro acertando menos da metade
# das vezes. A calculadora responde, ANTES de comprar: onde fica o stop,
# quanto posso perder nele, quantas ações comprar para não arriscar mais que
# RISCO_POR_OPERACAO_PCT do patrimônio, e se a compra concentra demais a
# carteira num setor. Usa só dados já salvos - não consome a API.
# --------------------------------------------------------------------------
# múltiplo do ATR para o stop sugerido: 2x a oscilação média diária fica
# fora do "ruído" normal do papel
STOP_MULTIPLO_ATR = 2
# sem ATR (histórico curto ou sem máxima/mínima), stop a 5% da entrada
STOP_PCT_SEM_ATR = 5
# alvo sugerido: ganho de 2x o risco (relação risco-retorno 1:2)
ALVO_MULTIPLO_RISCO = 2


def calcular_risco_compra(
    usuario, ticker: str, preco_entrada: float | None = None, stop: float | None = None,
    risco_pct: float | None = None,
) -> dict:
    """
    Calculadora de risco de uma compra de `ticker` para `usuario`. Sem
    `preco_entrada`, usa o último fechamento salvo; sem `stop`, sugere
    entrada - 2x ATR (ou -5% sem ATR), nunca abaixo do suporte se o suporte
    estiver mais perto. Devolve {"erro": ...} quando não dá para calcular.
    """
    ticker = (ticker or "").strip().upper()
    risco_pct = float(risco_pct if risco_pct is not None else settings.RISCO_POR_OPERACAO_PCT)
    ativo = Ativo.objects.filter(ticker=ticker).first()

    scanner = escanear_ativo(ativo) if ativo else None
    if preco_entrada is None:
        ultima = ativo.ultima_cotacao() if ativo else None
        preco_entrada = float(ultima.preco_fechamento) if ultima else None
    if not preco_entrada or preco_entrada <= 0:
        return {"erro": f"Sem preço para {ticker or 'o ativo'} - informe o preço unitário."}

    atr = scanner["atr"] if scanner else None
    suporte = (scanner["suporte_resistencia"] or {}).get("suporte") if scanner else None

    stop_sugerido = preco_entrada - STOP_MULTIPLO_ATR * atr if atr else preco_entrada * (1 - STOP_PCT_SEM_ATR / 100)
    base_stop = "ATR" if atr else "PERCENTUAL"
    if suporte and stop_sugerido < suporte < preco_entrada:
        # um pouco abaixo do suporte: rompido o suporte, a tese da compra caiu
        stop_sugerido = suporte * 0.99
        base_stop = "SUPORTE"
    stop_sugerido = round(max(stop_sugerido, 0.01), 2)

    stop_usado = stop if stop is not None else stop_sugerido
    if stop_usado >= preco_entrada:
        return {"erro": "O stop precisa ficar abaixo do preço de entrada."}
    risco_por_acao = preco_entrada - stop_usado

    # patrimônio = dinheiro em conta + valor de mercado do que já está comprado
    posicoes = [p for p in calcular_posicoes(usuario) if p.quantidade > 0]
    saldo = float(saldo_conta_corrente(obter_ou_criar_conta_corrente(usuario)))
    valor_posicoes = sum(float(p.valor_atual or p.valor_investido) for p in posicoes)
    patrimonio = max(saldo, 0) + valor_posicoes

    risco_maximo = patrimonio * risco_pct / 100
    qtd_por_risco = int(risco_maximo // risco_por_acao) if risco_maximo > 0 else 0
    # sem saldo registrado na conta corrente, a quantidade sai só pelo risco
    # (muita gente não lança o caixa no sistema) - com saldo, também não passa dele
    qtd_por_saldo = int(saldo // preco_entrada) if saldo > 0 else None
    quantidade = qtd_por_risco if qtd_por_saldo is None else min(qtd_por_risco, qtd_por_saldo)
    limitada_pelo_saldo = qtd_por_saldo is not None and qtd_por_saldo < qtd_por_risco

    setor = ativo.setor if ativo else ""
    exposicao_setor_pct = None
    exposicao_setor_apos_pct = None
    if setor and patrimonio > 0:
        valor_setor = sum(float(p.valor_atual or p.valor_investido) for p in posicoes if p.ativo.setor == setor)
        exposicao_setor_pct = round(valor_setor / patrimonio * 100, 1)
        exposicao_setor_apos_pct = round((valor_setor + quantidade * preco_entrada) / patrimonio * 100, 1)

    avisos = []
    if patrimonio <= 0:
        avisos.append("Sem patrimônio registrado: informe o saldo inicial em Conta Corrente para calcular a quantidade.")
    elif qtd_por_saldo is None:
        avisos.append(
            "Conta corrente sem saldo registrado: a quantidade considera só o risco - confira se você tem "
            "dinheiro disponível (e lance o saldo em Conta Corrente para um cálculo mais preciso)."
        )
    elif qtd_por_saldo == 0:
        avisos.append("Saldo da conta corrente insuficiente para comprar 1 ação a esse preço.")
    if patrimonio > 0 and qtd_por_risco == 0:
        avisos.append(
            f"Nem 1 ação cabe no risco de {risco_pct:g}% do patrimônio com esse stop - "
            "o stop está longe demais para o tamanho da carteira."
        )
    if exposicao_setor_apos_pct is not None and exposicao_setor_apos_pct > settings.LIMITE_SETOR_PCT:
        avisos.append(
            f"Concentração: com esta compra, {exposicao_setor_apos_pct}% do patrimônio fica no setor "
            f"{setor} (limite recomendado: {settings.LIMITE_SETOR_PCT:g}%)."
        )
    if scanner and scanner["veredito"] == "PREDOMINIO_BAIXA":
        avisos.append("O Scanner Técnico aponta predomínio de BAIXA neste ativo agora.")

    backtest = backtest_sinais_robo(ativo) if ativo else None
    alvo = preco_entrada + ALVO_MULTIPLO_RISCO * risco_por_acao
    fundamentos = FundamentosAtivo.objects.filter(ticker=ticker).first()
    if fundamentos and fundamentos.classificacao == "BAIXA":
        avisos.append(
            f"Qualidade fundamentalista BAIXA (nota {fundamentos.nota}) - veja Qualidade das Empresas."
        )

    return {
        "erro": None,
        "ticker": ticker,
        "preco_entrada": round(preco_entrada, 2),
        "atr": round(atr, 2) if atr else None,
        "suporte": round(suporte, 2) if suporte else None,
        "stop_sugerido": stop_sugerido,
        "base_stop": base_stop,
        "stop": round(stop_usado, 2),
        "queda_ate_stop_pct": round(risco_por_acao / preco_entrada * 100, 2),
        "alvo": round(alvo, 2),
        "risco_pct": risco_pct,
        "patrimonio": round(patrimonio, 2),
        "saldo": round(saldo, 2),
        "risco_maximo": round(risco_maximo, 2),
        "quantidade": quantidade,
        "limitada_pelo_saldo": limitada_pelo_saldo,
        "valor_compra": round(quantidade * preco_entrada, 2),
        "perda_no_stop": round(quantidade * risco_por_acao, 2),
        "ganho_no_alvo": round(quantidade * (alvo - preco_entrada), 2),
        "setor": setor,
        "exposicao_setor_pct": exposicao_setor_pct,
        "exposicao_setor_apos_pct": exposicao_setor_apos_pct,
        "veredito_scanner": scanner["veredito_label"] if scanner else None,
        "nota_qualidade": float(fundamentos.nota) if fundamentos and fundamentos.nota is not None else None,
        "backtest_veredito": backtest["veredito"] if backtest else None,
        "backtest_sinais_compra": backtest["sinais_compra"] if backtest else 0,
        "avisos": avisos,
    }


# --------------------------------------------------------------------------
# "Análise da B3 hoje" por IA (tela Análise de Mercado)
#
# A IA só interpreta e comenta dados que o sistema já calculou sozinho
# (maiores altas/baixas do dia e os indicadores do Scanner Técnico de cada
# ativo acompanhado - comprado ou reservado, ver escanear_carteira) - não
# inventa preço, indicador, pontuação nem notícia nenhuma. Qualquer provedor
# compatível com o formato de chat completions da OpenAI funciona (ver
# OPENAI_BASE_URL no .env) - não precisa ser a OpenAI paga, provedores como
# Groq ou OpenRouter têm camada gratuita com esse mesmo formato de API.
# --------------------------------------------------------------------------
IA_TIMEOUT_SEGUNDOS = 45
IA_SITUACAO_CLASSES = {"ALTA": "alta", "BAIXA": "baixa", "NEUTRO": "neutro"}


class IAError(Exception):
    """Erro ao consultar a IA da "Análise da B3 hoje" - chave não configurada, rede, ou resposta inesperada."""


def ia_configurada() -> bool:
    """True quando há uma chave de API configurada (OPENAI_API_KEY no .env) para a Análise da B3 hoje por IA."""
    return bool(settings.OPENAI_API_KEY)


def _score_scanner(item: dict) -> int:
    """
    Pontuação 0-100 derivada dos votos do Scanner Técnico (ver
    escanear_ativo_precos): 50 é neutro; cada indicador de alta soma e cada
    indicador de baixa subtrai o mesmo peso (50 / total de indicadores),
    arredondado e limitado a [0, 100]. Puramente determinístico - calculado
    localmente, nunca pela IA (ver _prompt_analise_b3_ia/gerar_analise_b3_ia),
    pra não depender da IA "inventar" uma nota.
    """
    total = len(item["indicadores"])
    if not total:
        return 50
    bruto = 50 + (item["votos_alta"] - item["votos_baixa"]) * (50 / total)
    return round(min(100, max(0, bruto)))


def _prompt_analise_b3_ia(scanner_itens: list[dict], maiores_altas: list[dict], maiores_baixas: list[dict]) -> str:
    """
    Monta o prompt com o contexto do pregão de hoje já calculado localmente,
    pedindo a resposta em linhas "TICKER|SITUACAO|COMENTARIO" (ver
    _parsear_resposta_ia_em_grade) em vez de texto corrido. `scanner_itens`
    é a saída de escanear_carteira - mesmos indicadores (IFR, médias móveis,
    MACD, volume, tendência) e pontuação (ver _score_scanner) mostrados na
    tela Scanner Técnico, cobrindo tanto os ativos comprados quanto os só
    reservados.
    """
    linhas_altas = "\n".join(f"- {a['stock']} ({a['name']}): +{a['change']}%" for a in maiores_altas[:10]) or "sem dados"
    linhas_baixas = "\n".join(f"- {a['stock']} ({a['name']}): {a['change']}%" for a in maiores_baixas[:10]) or "sem dados"

    def _linha_ativo(item):
        indicadores = {i["nome"]: i["label"] for i in item["indicadores"]}
        rsi_valor = f" ({item['rsi']})" if item.get("rsi") is not None else ""
        return (
            f"- {item['ativo'].ticker}: IFR {indicadores.get('IFR (RSI)', 'sem dados')}{rsi_valor}, "
            f"médias móveis {indicadores.get('Médias móveis', 'sem dados')}, "
            f"tendência {indicadores.get('Tendência de curto prazo', 'sem dados')}, "
            f"volume {indicadores.get('Volume', 'sem dados')}, "
            f"veredito do scanner {item['veredito_label']}, pontuação técnica {_score_scanner(item)}/100"
        )

    linhas_ativos = "\n".join(_linha_ativo(item) for item in scanner_itens) or "o usuário ainda não acompanha nenhum ativo"

    return (
        "Você é um analista comentando o pregão de hoje da B3 (bolsa brasileira) para um investidor pessoa "
        "física. Use SOMENTE os dados abaixo, já calculados pelo sistema - não invente preços, indicadores, "
        "pontuações nem notícias que não estão aqui.\n\n"
        f"Maiores altas do dia:\n{linhas_altas}\n\n"
        f"Maiores baixas do dia:\n{linhas_baixas}\n\n"
        f"Indicadores técnicos e pontuação do Scanner Técnico dos ativos acompanhados pelo usuário "
        f"(comprados ou reservados):\n{linhas_ativos}\n\n"
        "Responda APENAS com linhas no formato exato abaixo (sem cabeçalho, sem markdown, sem texto antes ou "
        "depois):\nTICKER|SITUACAO|COMENTARIO\n\n"
        "- Uma linha \"MERCADO|SITUACAO|comentário\" resumindo o pregão geral da B3 hoje.\n"
        "- Uma linha para cada ativo acompanhado pelo usuário, na mesma ordem listada acima.\n"
        "- SITUACAO deve ser exatamente uma destas palavras: ALTA, BAIXA ou NEUTRO.\n"
        "- COMENTARIO em português, direto ao ponto, no máximo 160 caracteres.\n"
        "- Isto não é recomendação de investimento, só uma leitura dos dados apresentados."
    )


def _parsear_resposta_ia_em_grade(texto: str) -> list[dict]:
    """
    Converte a resposta da IA (linhas "TICKER|SITUACAO|COMENTARIO") numa lista
    de dicts prontos pra tabela. Linhas fora do formato esperado são
    ignoradas silenciosamente (a IA às vezes decora a resposta com algo a
    mais, apesar do pedido) - se nenhuma linha for aproveitável, quem chama
    trata a lista vazia mostrando o texto bruto como alternativa.
    """
    linhas = []
    for bruta in texto.strip().splitlines():
        partes = [p.strip() for p in bruta.split("|")]
        if len(partes) != 3:
            continue
        ticker, situacao, comentario = partes
        situacao = situacao.upper()
        if situacao not in IA_SITUACAO_CLASSES or not ticker or not comentario:
            continue
        linhas.append({
            "ticker": ticker,
            "eh_mercado_geral": ticker.upper() == "MERCADO",
            "situacao": situacao,
            "situacao_classe": IA_SITUACAO_CLASSES[situacao],
            "comentario": comentario,
        })
    return linhas


def _enriquecer_linhas_com_scanner(linhas: list[dict], scanner_itens: list[dict]) -> list[dict]:
    """
    Anexa a cada linha já parseada da resposta da IA (ver
    _parsear_resposta_ia_em_grade) os indicadores e a pontuação do Scanner
    Técnico do mesmo ticker (ver escanear_carteira/_score_scanner) - dados
    100% calculados localmente, nunca pela IA. A linha "MERCADO" (pregão
    geral) e qualquer ticker que a IA tenha citado fora da lista enviada
    ficam com esses campos em branco (score None).
    """
    por_ticker = {item["ativo"].ticker: item for item in scanner_itens}
    for linha in linhas:
        item = por_ticker.get(linha["ticker"])
        if item is None:
            linha["score"] = None
            continue
        indicadores = {i["nome"]: i for i in item["indicadores"]}
        linha["rsi"] = item.get("rsi")
        linha["ifr_label"] = indicadores.get("IFR (RSI)", {}).get("label")
        linha["ifr_classe"] = indicadores.get("IFR (RSI)", {}).get("classe")
        linha["medias_label"] = indicadores.get("Médias móveis", {}).get("label")
        linha["medias_classe"] = indicadores.get("Médias móveis", {}).get("classe")
        linha["tendencia_label"] = indicadores.get("Tendência de curto prazo", {}).get("label")
        linha["tendencia_classe"] = indicadores.get("Tendência de curto prazo", {}).get("classe")
        linha["volume_label"] = indicadores.get("Volume", {}).get("label")
        linha["volume_classe"] = indicadores.get("Volume", {}).get("classe")
        linha["veredito_label"] = item["veredito_label"]
        linha["veredito_classe"] = item["veredito_classe"]
        linha["score"] = _score_scanner(item)
    return linhas


def gerar_analise_b3_ia(scanner_itens: list[dict], maiores_altas: list[dict], maiores_baixas: list[dict]) -> dict:
    """
    Consulta a IA configurada (ver ia_configurada) pra comentar o pregão de
    hoje e os ativos acompanhados, devolvendo o texto bruto, as linhas já
    organizadas (ver _parsear_resposta_ia_em_grade) enriquecidas com IFR/
    médias/tendência/volume/pontuação do Scanner Técnico (ver
    _enriquecer_linhas_com_scanner) pra montar a tabela na tela Análise de
    Mercado, e um ranking dos ativos por pontuação técnica (maior primeiro,
    calculado localmente - não é a IA que ordena). Levanta IAError se a IA
    não estiver configurada ou a chamada falhar (rede, chave inválida,
    resposta fora do esperado etc.).
    """
    if not ia_configurada():
        raise IAError("A Análise da B3 hoje por IA ainda não está configurada neste sistema (falta a chave de API).")

    prompt = _prompt_analise_b3_ia(scanner_itens, maiores_altas, maiores_baixas)
    url = f"{settings.OPENAI_BASE_URL.rstrip('/')}/chat/completions"
    headers = {"Authorization": f"Bearer {settings.OPENAI_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": settings.OPENAI_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.4,
    }
    try:
        resposta = requests.post(url, json=payload, headers=headers, timeout=IA_TIMEOUT_SEGUNDOS)
        resposta.raise_for_status()
        texto = resposta.json()["choices"][0]["message"]["content"]
    except (requests.RequestException, KeyError, IndexError, TypeError, ValueError) as exc:
        raise IAError("Não foi possível consultar a IA agora. Tente novamente em instantes.") from exc

    linhas = _enriquecer_linhas_com_scanner(_parsear_resposta_ia_em_grade(texto), scanner_itens)
    ranking = sorted(
        (linha for linha in linhas if not linha["eh_mercado_geral"] and linha["score"] is not None),
        key=lambda linha: linha["score"], reverse=True,
    )

    return {"texto": texto.strip(), "linhas": linhas, "ranking": ranking}


def _prompt_sugestoes_ia(maiores_altas: list[dict], maiores_baixas: list[dict]) -> str:
    """
    Monta o prompt do botão "Sugestões de hoje (IA)": pede pra IA escolher, só
    dentre as maiores altas/baixas do dia já calculadas pelo sistema (nunca
    inventando um ativo ou preço fora dessa lista), de 3 a 5 que merecem
    atenção - resposta em linhas "TICKER|JUSTIFICATIVA" (ver
    _parsear_sugestoes_ia_em_grade).
    """
    linhas_altas = "\n".join(f"- {a['stock']} ({a['name']}): +{a['change']}%" for a in maiores_altas[:10]) or "sem dados"
    linhas_baixas = "\n".join(f"- {a['stock']} ({a['name']}): {a['change']}%" for a in maiores_baixas[:10]) or "sem dados"

    return (
        "Você é um analista comentando o pregão de hoje da B3 (bolsa brasileira) para um investidor pessoa "
        "física. Use SOMENTE os ativos e variações abaixo, já calculados pelo sistema - não invente nenhum "
        "ativo, preço ou variação que não esteja nesta lista.\n\n"
        f"Maiores altas do dia:\n{linhas_altas}\n\n"
        f"Maiores baixas do dia:\n{linhas_baixas}\n\n"
        "Escolha de 3 a 5 ativos DESTA LISTA que merecem atenção para os próximos dias e responda APENAS com "
        "linhas no formato exato abaixo (sem cabeçalho, sem markdown, sem texto antes ou depois):\n"
        "TICKER|JUSTIFICATIVA\n\n"
        "- Uma linha por ativo escolhido, só ativos que aparecem na lista acima.\n"
        "- JUSTIFICATIVA em português, direto ao ponto, no máximo 160 caracteres, baseada só na variação do "
        "dia apresentada.\n"
        "- Isto não é recomendação de investimento, só uma leitura das maiores movimentações do dia."
    )


def _parsear_sugestoes_ia_em_grade(texto: str) -> list[dict]:
    """
    Converte a resposta do prompt de sugestões (linhas "TICKER|JUSTIFICATIVA")
    numa lista de dicts prontos pra tabela - mesmo espírito tolerante de
    _parsear_resposta_ia_em_grade: linha fora do formato esperado é ignorada
    em vez de quebrar a tela.
    """
    linhas = []
    for bruta in texto.strip().splitlines():
        partes = [p.strip() for p in bruta.split("|", 1)]
        if len(partes) != 2:
            continue
        ticker, justificativa = partes
        if not ticker or not justificativa:
            continue
        linhas.append({"ticker": ticker.upper(), "justificativa": justificativa})
    return linhas


def gerar_sugestoes_ia(maiores_altas: list[dict], maiores_baixas: list[dict]) -> dict:
    """
    Consulta a IA configurada (ver ia_configurada) pra sugerir, sob demanda,
    de 3 a 5 ativos entre as maiores altas/baixas do dia que merecem atenção
    - botão "🤖 Sugestões de hoje (IA)" na tela Análise de Mercado. Mesma
    ideia de gerar_analise_b3_ia (a IA só interpreta dados que o sistema já
    calculou), levanta IAError nas mesmas condições.
    """
    if not ia_configurada():
        raise IAError("As Sugestões de hoje por IA ainda não estão configuradas neste sistema (falta a chave de API).")

    prompt = _prompt_sugestoes_ia(maiores_altas, maiores_baixas)
    url = f"{settings.OPENAI_BASE_URL.rstrip('/')}/chat/completions"
    headers = {"Authorization": f"Bearer {settings.OPENAI_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": settings.OPENAI_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.4,
    }
    try:
        resposta = requests.post(url, json=payload, headers=headers, timeout=IA_TIMEOUT_SEGUNDOS)
        resposta.raise_for_status()
        texto = resposta.json()["choices"][0]["message"]["content"]
    except (requests.RequestException, KeyError, IndexError, TypeError, ValueError) as exc:
        raise IAError("Não foi possível consultar a IA agora. Tente novamente em instantes.") from exc

    return {"texto": texto.strip(), "linhas": _parsear_sugestoes_ia_em_grade(texto)}


# --------------------------------------------------------------------------
# "Análise da carteira" (aba da tela Posições em Carteira)
#
# Duas camadas: analisar_carteira_comprada é 100% local (resultado, metas,
# Scanner Técnico, fundamentos, concentração, correlação, beta, benchmark) e
# sai na hora em que a tela abre; gerar_analise_carteira_ia manda esse mesmo
# retrato pra IA comentar - a IA só interpreta, não inventa número nenhum.
# --------------------------------------------------------------------------
RECOMENDACOES_CARTEIRA = {
    "STOP": ("🛑 Stop atingido — revisar", "baixa"),
    "REALIZAR": ("💰 Meta atingida — realizar lucro", "alta"),
    "PROTEGER": ("📈 Meta atingida em alta — subir o stop", "alta"),
    "ATENCAO": ("⚠ Atenção", "neutro"),
    "MANTER": ("✅ Manter", "alta"),
    "ACOMPANHAR": ("👀 Acompanhar", "neutro"),
}
IA_ACOES_CARTEIRA = {"MANTER": "alta", "AUMENTAR": "alta", "REALIZAR": "neutro", "ATENCAO": "neutro", "REDUZIR": "baixa"}


def _recomendacao_posicao(lucro_pct, meta_perda, meta_lucro_atingida: bool, veredito: str, qualidade: str) -> str:
    """Regra (ordem importa) da recomendação local de cada posição - ver RECOMENDACOES_CARTEIRA."""
    if lucro_pct is not None and lucro_pct <= meta_perda:
        return "STOP"
    if meta_lucro_atingida:
        return "PROTEGER" if veredito == "PREDOMINIO_ALTA" else "REALIZAR"
    if veredito == "PREDOMINIO_BAIXA" or (qualidade == "BAIXA" and lucro_pct is not None and lucro_pct < 0):
        return "ATENCAO"
    if veredito == "PREDOMINIO_ALTA":
        return "MANTER"
    return "ACOMPANHAR"


def analisar_carteira_comprada(
    posicoes: list, metricas_risco: dict | None = None, concentracao_setor: list | None = None,
    correlacao: dict | None = None, estresse: dict | None = None, comparativo_benchmark: dict | None = None,
) -> dict | None:
    """
    Análise completa das posições compradas (reservas não entram): para cada
    ativo, resultado, distância até a meta de lucro e até o stop, Scanner
    Técnico (veredito, pontuação, IFR, suporte/resistência), volatilidade,
    peso na carteira e nota de qualidade, com uma recomendação e os motivos;
    para a carteira, uma nota de saúde 0-100 com pontos fortes e de atenção
    (concentração por ativo/setor, correlação, beta, ativos em baixa ou no
    stop, comparação com Ibovespa/CDI). Os dicts opcionais são os que a tela
    já calculou (ver core.views.posicoes) - quando não vêm, são calculados
    aqui. None se não há nada comprado.
    """
    compradas = [p for p in posicoes if not p.apenas_reservado and p.quantidade > 0]
    if not compradas:
        return None
    if metricas_risco is None:
        metricas_risco = calcular_metricas_risco(posicoes)
    if concentracao_setor is None:
        concentracao_setor = calcular_concentracao_setor(posicoes)

    pesos = {i["ativo"].ticker: i for i in metricas_risco["por_ativo"]}
    fundamentos = {f.ticker: f for f in FundamentosAtivo.objects.filter(ticker__in=[p.ativo.ticker for p in compradas])}

    ativos = []
    for p in compradas:
        ticker = p.ativo.ticker
        scanner = escanear_ativo(p.ativo)
        score = _score_scanner(scanner)
        fund = fundamentos.get(ticker)
        qualidade = fund.classificacao if fund else "SEM_DADOS"
        risco = pesos.get(ticker, {})
        peso_pct = float(risco["percentual_carteira"]) if risco.get("percentual_carteira") is not None else None
        lucro_pct = float(p.lucro_perda_pct) if p.lucro_perda_pct is not None else None
        meta_lucro = float(abs(p.meta_lucro_pct if p.meta_lucro_pct is not None else Decimal(str(settings.META_LUCRO_PADRAO))))
        meta_perda = -float(abs(p.meta_perda_pct if p.meta_perda_pct is not None else Decimal(str(settings.META_PERDA_PADRAO))))
        rsi = scanner.get("rsi")
        sr = scanner.get("suporte_resistencia") or {}

        motivos = []
        if lucro_pct is not None:
            dias = p.dias_desde_compra
            motivos.append(
                f"Resultado de {lucro_pct:+.2f}%"
                + (f" em {dias} dia{'s' if dias != 1 else ''} de carteira" if dias is not None else "")
                + (f" (R$ {float(p.lucro_perda_valor):+.2f})" if p.lucro_perda_valor is not None else "") + "."
            )
            if lucro_pct < meta_lucro:
                motivos.append(f"Faltam {meta_lucro - lucro_pct:.2f} p.p. para a meta de lucro (+{meta_lucro:g}%).")
            if lucro_pct > meta_perda:
                motivos.append(f"Está {lucro_pct - meta_perda:.2f} p.p. acima do stop ({meta_perda:g}%).")
        else:
            motivos.append("Sem cotação para apurar o resultado - atualize as cotações.")
        motivos.append(f"Scanner Técnico: {scanner['veredito_label']} (pontuação {score}/100).")
        if rsi is not None and rsi >= 70:
            motivos.append(f"IFR {rsi:.0f}: sobrecomprado - o preço subiu rápido e pode corrigir.")
        elif rsi is not None and rsi <= 30:
            motivos.append(f"IFR {rsi:.0f}: sobrevendido - caiu muito em pouco tempo.")
        if sr.get("suporte") or sr.get("resistencia"):
            motivos.append(
                "Suporte " + (f"R$ {sr['suporte']:.2f}" if sr.get("suporte") else "—")
                + " · resistência " + (f"R$ {sr['resistencia']:.2f}" if sr.get("resistencia") else "—") + "."
            )
        if scanner["volatilidade_chave"] == "ALTA":
            motivos.append(f"Volatilidade alta ({scanner['volatilidade_pct']}% ao ano) - oscila bastante.")
        if peso_pct is not None and peso_pct >= float(LIMIAR_CONCENTRACAO_ALERTA_PCT):
            motivos.append(f"Sozinho é {peso_pct:.1f}% da carteira - posição concentrada.")
        if fund and fund.nota is not None:
            motivos.append(f"Qualidade fundamentalista {qualidade.lower().replace('media', 'média')} (nota {fund.nota}/10).")

        codigo = _recomendacao_posicao(lucro_pct, meta_perda, p.meta_lucro_atingida, scanner["veredito"], qualidade)
        label, classe = RECOMENDACOES_CARTEIRA[codigo]
        ativos.append({
            "ticker": ticker,
            "nome": p.ativo.nome or "",
            "setor": p.ativo.setor or "",
            "peso_pct": round(peso_pct, 1) if peso_pct is not None else None,
            "lucro_pct": round(lucro_pct, 2) if lucro_pct is not None else None,
            "lucro_valor": float(p.lucro_perda_valor) if p.lucro_perda_valor is not None else None,
            "dias": p.dias_desde_compra,
            "preco_medio": float(p.preco_medio),
            "preco_atual": float(p.preco_atual) if p.preco_atual is not None else None,
            "meta_lucro_pct": meta_lucro,
            "meta_perda_pct": meta_perda,
            "veredito": scanner["veredito"],
            "veredito_label": scanner["veredito_label"],
            "veredito_classe": scanner["veredito_classe"],
            "score": score,
            "rsi": round(rsi, 1) if rsi is not None else None,
            "volatilidade_pct": scanner["volatilidade_pct"],
            "nota_qualidade": float(fund.nota) if fund and fund.nota is not None else None,
            "recomendacao": codigo,
            "recomendacao_label": label,
            "recomendacao_classe": classe,
            "motivos": motivos,
        })
    ativos.sort(key=lambda a: a["peso_pct"] or 0, reverse=True)

    # ---- visão da carteira inteira
    nota, fortes, atencao = 100, [], []
    com_resultado = [a for a in ativos if a["lucro_pct"] is not None]
    em_lucro = [a for a in com_resultado if a["lucro_pct"] > 0]
    em_perda = [a for a in com_resultado if a["lucro_pct"] < 0]
    if len(ativos) < 3:
        nota -= 10
        atencao.append(f"Só {len(ativos)} ativo{'s' if len(ativos) > 1 else ''} comprado{'s' if len(ativos) > 1 else ''}: pouca diversificação.")
    if metricas_risco.get("concentracao_alta"):
        nota -= 15
        atencao.append(f"{metricas_risco['maior_posicao_ativo'].ticker} sozinho é {metricas_risco['maior_posicao_pct']}% da carteira.")
    setores_altos = [s for s in concentracao_setor if s["setor"] != "Sem setor" and float(s["percentual"]) > settings.LIMITE_SETOR_PCT]
    if setores_altos:
        nota -= min(10 * len(setores_altos), 20)
        atencao.append(
            "Setor acima do limite recomendado de " + f"{settings.LIMITE_SETOR_PCT:g}%: "
            + ", ".join(f"{s['setor']} ({s['percentual']}%)" for s in setores_altos) + "."
        )
    elif len([s for s in concentracao_setor if s["setor"] != "Sem setor"]) >= 3:
        fortes.append(f"Carteira espalhada em {len(concentracao_setor)} setores, nenhum acima de {settings.LIMITE_SETOR_PCT:g}%.")
    if correlacao and correlacao.get("pares_altos"):
        nota -= min(5 * len(correlacao["pares_altos"]), 15)
        atencao.append(
            "Ativos que andam muito juntos (não diversificam entre si): "
            + ", ".join(f"{par['a']} & {par['b']} ({par['valor']})" for par in correlacao["pares_altos"][:4]) + "."
        )
    elif correlacao:
        fortes.append(f"Nenhum par de ativos com correlação alta (média da carteira {correlacao['media']}).")
    beta = estresse.get("beta_carteira") if estresse else None
    if beta is not None and beta > 1.2:
        nota -= 10
        atencao.append(f"Beta {beta}: a carteira tende a cair mais que o Ibovespa nas quedas.")
    elif beta is not None and beta < 0.9:
        fortes.append(f"Beta {beta}: a carteira tende a oscilar menos que o Ibovespa.")
    em_baixa = [a for a in ativos if a["veredito"] == "PREDOMINIO_BAIXA"]
    if em_baixa:
        nota -= round(20 * len(em_baixa) / len(ativos))
        atencao.append("Scanner Técnico com predomínio de baixa em: " + ", ".join(a["ticker"] for a in em_baixa) + ".")
    em_alta = [a for a in ativos if a["veredito"] == "PREDOMINIO_ALTA"]
    if em_alta:
        fortes.append("Scanner Técnico com predomínio de alta em: " + ", ".join(a["ticker"] for a in em_alta) + ".")
    no_stop = [a for a in ativos if a["recomendacao"] == "STOP"]
    if no_stop:
        nota -= min(10 * len(no_stop), 20)
        atencao.append("Já passaram do stop: " + ", ".join(f"{a['ticker']} ({a['lucro_pct']:+.2f}%)" for a in no_stop) + ".")
    na_meta = [a for a in ativos if a["recomendacao"] in ("REALIZAR", "PROTEGER")]
    if na_meta:
        fortes.append("Meta de lucro atingida: " + ", ".join(f"{a['ticker']} ({a['lucro_pct']:+.2f}%)" for a in na_meta) + ".")
    qualidade_baixa = [a for a in ativos if a["nota_qualidade"] is not None and a["nota_qualidade"] < 4]
    if qualidade_baixa:
        nota -= 5
        atencao.append("Qualidade fundamentalista baixa: " + ", ".join(a["ticker"] for a in qualidade_baixa) + ".")
    if comparativo_benchmark:
        if comparativo_benchmark.get("retorno_ibovespa_pct") is not None:
            if comparativo_benchmark.get("carteira_bateu_ibovespa"):
                fortes.append(f"Carteira ({comparativo_benchmark['retorno_carteira_pct']}%) bateu o Ibovespa ({comparativo_benchmark['retorno_ibovespa_pct']}%) no período.")
            else:
                nota -= 5
                atencao.append(f"Carteira ({comparativo_benchmark['retorno_carteira_pct']}%) abaixo do Ibovespa ({comparativo_benchmark['retorno_ibovespa_pct']}%) no período.")
        if comparativo_benchmark.get("retorno_cdi_pct") is not None and not comparativo_benchmark.get("carteira_bateu_cdi"):
            nota -= 5
            atencao.append(f"Carteira abaixo do CDI ({comparativo_benchmark['retorno_cdi_pct']}%) no período - a renda fixa rendeu mais.")
    nota = max(0, min(100, nota))
    if nota >= 75:
        classificacao, classe = "Saudável", "alta"
    elif nota >= 50:
        classificacao, classe = "Atenção moderada", "neutro"
    else:
        classificacao, classe = "Precisa de cuidado", "baixa"

    return {
        "ativos": ativos,
        "nota": nota,
        "classificacao": classificacao,
        "classe": classe,
        "pontos_fortes": fortes,
        "pontos_atencao": atencao,
        "total": len(ativos),
        "em_lucro": len(em_lucro),
        "em_perda": len(em_perda),
        "melhor": max(com_resultado, key=lambda a: a["lucro_pct"]) if com_resultado else None,
        "pior": min(com_resultado, key=lambda a: a["lucro_pct"]) if com_resultado else None,
        "beta": beta,
    }


def _prompt_analise_carteira_ia(analise: dict) -> str:
    """Prompt da "Análise da carteira (IA)": manda o retrato local (ver analisar_carteira_comprada) e pede linhas "CHAVE|ACAO|COMENTARIO"."""
    def _fmt(valor, sufixo=""):
        return f"{valor}{sufixo}" if valor is not None else "sem dados"

    linhas_ativos = "\n".join(
        f"- {a['ticker']} ({a['nome'] or 'sem nome'}, setor {a['setor'] or 'desconhecido'}): "
        f"{_fmt(a['peso_pct'], '%')} da carteira, resultado {_fmt(a['lucro_pct'], '%')} "
        f"em {_fmt(a['dias'])} dias, preço médio R$ {a['preco_medio']:.2f}, cotação {_fmt(a['preco_atual'])}, "
        f"meta +{a['meta_lucro_pct']:g}% / stop {a['meta_perda_pct']:g}%, scanner {a['veredito_label']} "
        f"({a['score']}/100), IFR {_fmt(a['rsi'])}, volatilidade {_fmt(a['volatilidade_pct'], '% a.a.')}, "
        f"nota de qualidade {_fmt(a['nota_qualidade'], '/10')}, regra do sistema: {a['recomendacao_label']}"
        for a in analise["ativos"]
    )
    fortes = "\n".join(f"- {t}" for t in analise["pontos_fortes"]) or "- nenhum destacado"
    atencao = "\n".join(f"- {t}" for t in analise["pontos_atencao"]) or "- nenhum destacado"

    return (
        "Você é um analista revisando a carteira de ações (B3) de um investidor pessoa física. Use SOMENTE os "
        "dados abaixo, já calculados pelo sistema - não invente preços, indicadores, notícias nem números que "
        "não estão aqui.\n\n"
        f"Nota de saúde da carteira (calculada pelo sistema): {analise['nota']}/100 - {analise['classificacao']}. "
        f"{analise['em_lucro']} ativo(s) em lucro e {analise['em_perda']} em prejuízo, beta {_fmt(analise['beta'])}.\n\n"
        f"Ativos comprados:\n{linhas_ativos}\n\n"
        f"Pontos fortes:\n{fortes}\n\nPontos de atenção:\n{atencao}\n\n"
        "Responda APENAS com linhas no formato exato abaixo (sem cabeçalho, sem markdown, sem texto antes ou "
        "depois):\nCHAVE|ACAO|COMENTARIO\n\n"
        "- Uma linha \"CARTEIRA|ACAO|comentário\" com o diagnóstico geral (até 350 caracteres).\n"
        "- Uma linha por ativo, com o TICKER como CHAVE, na mesma ordem listada acima (comentário até 220 "
        "caracteres, citando os dados que justificam).\n"
        "- De 2 a 4 linhas \"DICA|ACAO|próximo passo prático para a carteira\" (até 200 caracteres).\n"
        "- ACAO deve ser exatamente uma destas palavras: MANTER, AUMENTAR, REALIZAR, ATENCAO ou REDUZIR.\n"
        "- Comentários em português, diretos ao ponto.\n"
        "- Isto não é recomendação de investimento, só uma leitura dos dados apresentados."
    )


def _parsear_analise_carteira_ia(texto: str) -> dict:
    """Separa a resposta da IA em diagnóstico geral, linhas por ativo e dicas - linha fora do formato é ignorada."""
    resultado = {"carteira": None, "ativos": [], "dicas": []}
    for bruta in texto.strip().splitlines():
        partes = [p.strip().strip("*") for p in bruta.split("|")]
        if len(partes) != 3:
            continue
        chave, acao, comentario = partes
        acao = unicodedata.normalize("NFKD", acao.upper()).encode("ascii", "ignore").decode()
        if acao not in IA_ACOES_CARTEIRA or not chave or not comentario:
            continue
        linha = {"chave": chave.upper(), "acao": acao, "acao_classe": IA_ACOES_CARTEIRA[acao], "comentario": comentario}
        if linha["chave"] == "CARTEIRA":
            resultado["carteira"] = linha
        elif linha["chave"] == "DICA":
            resultado["dicas"].append(linha)
        else:
            resultado["ativos"].append(linha)
    return resultado


def gerar_analise_carteira_ia(analise: dict) -> dict:
    """
    Consulta a IA configurada (ver ia_configurada) pra comentar a análise
    local da carteira comprada (ver analisar_carteira_comprada) - botão
    "Analisar agora" da aba Análise da carteira. Levanta IAError nas mesmas
    condições de gerar_analise_b3_ia.
    """
    if not ia_configurada():
        raise IAError("A Análise da carteira por IA ainda não está configurada neste sistema (falta a chave de API).")

    url = f"{settings.OPENAI_BASE_URL.rstrip('/')}/chat/completions"
    headers = {"Authorization": f"Bearer {settings.OPENAI_API_KEY}", "Content-Type": "application/json"}
    payload = {
        "model": settings.OPENAI_MODEL,
        "messages": [{"role": "user", "content": _prompt_analise_carteira_ia(analise)}],
        "temperature": 0.4,
    }
    try:
        resposta = requests.post(url, json=payload, headers=headers, timeout=IA_TIMEOUT_SEGUNDOS)
        resposta.raise_for_status()
        texto = resposta.json()["choices"][0]["message"]["content"]
    except (requests.RequestException, KeyError, IndexError, TypeError, ValueError) as exc:
        raise IAError("Não foi possível consultar a IA agora. Tente novamente em instantes.") from exc

    return {"texto": texto.strip(), **_parsear_analise_carteira_ia(texto)}


# --------------------------------------------------------------------------
# Notificações pelo Telegram
# --------------------------------------------------------------------------
TELEGRAM_TIMEOUT_SEGUNDOS = 10


def credenciais_telegram(usuario) -> tuple[str, str] | None:
    """
    (token do bot, chat de destino) cadastrados pelo usuário em "Minha Conta"
    (accounts.models.PerfilUsuario) - cada usuário usa o seu próprio bot.
    None quando ele não tem perfil ou falta um dos dois campos.
    """
    perfil = getattr(usuario, "perfil", None)
    if perfil is None or not (perfil.telegram_bot_token and perfil.telegram_chat_id):
        return None
    return perfil.telegram_bot_token, perfil.telegram_chat_id


def enviar_notificacao_telegram(usuario, mensagem: str) -> bool:
    """
    Envia uma mensagem de texto para o usuário pelo bot do Telegram que ele
    cadastrou (ver credenciais_telegram). Retorna True se o Telegram
    confirmou o envio.

    Nunca levanta exceção: uma falha aqui (sem internet, token inválido) não
    pode derrubar o ciclo de cotações nem a tela que gerou o alerta - só fica
    registrada no log. O log não inclui a exceção crua porque o texto dela
    traz a URL da requisição, e a URL contém o token do bot.
    """
    credenciais = credenciais_telegram(usuario)
    if credenciais is None:
        return False
    token, chat_id = credenciais

    url = f"https://api.telegram.org/bot{token}/sendMessage"
    payload = {"chat_id": chat_id, "text": mensagem}
    try:
        resposta = requests.post(url, data=payload, timeout=TELEGRAM_TIMEOUT_SEGUNDOS)
        dados = resposta.json()
    except (requests.RequestException, ValueError) as exc:
        logger.warning("Falha ao enviar notificação pelo Telegram (%s).", type(exc).__name__)
        return False

    if not dados.get("ok"):
        logger.warning(
            "Telegram recusou a notificação (HTTP %s): %s", resposta.status_code, dados.get("description"),
        )
        return False
    return True


def notificar_alertas_telegram(usuario, alertas: list[Alerta]) -> bool:
    """
    Avisa pelo Telegram os alertas recém-criados de um usuário, todos numa
    única mensagem. Como os alertas não se repetem no mesmo dia para o mesmo
    ativo/tipo, cada um é avisado uma vez só.

    Só vão para o Telegram a meta de lucro atingida de um ativo comprado
    (Alerta.LUCRO) e o sinal de compra do robô consultor (Alerta.SINAL_COMPRA)
    - meta de perda, tendência e sinal de venda continuam sendo gerados, mas
    ficam só na tela de alertas do sistema.
    """
    alertas = [alerta for alerta in alertas if alerta.tipo in (Alerta.LUCRO, Alerta.SINAL_COMPRA)]
    if not alertas or credenciais_telegram(usuario) is None:
        return False

    linhas = [f"🔔 BolsaTrader - {usuario.get_username()}"]
    linhas += [f"• {alerta.mensagem}" for alerta in alertas]
    return enviar_notificacao_telegram(usuario, "\n".join(linhas))


TELEGRAM_LIMITE_CARACTERES = 4000  # o Telegram recusa mensagens acima de 4096


def _enviar_linhas_telegram(usuario, titulo: str, linhas: list[str]) -> bool:
    """Monta "título + linhas" numa só mensagem (cortada no limite do Telegram) e envia."""
    texto = "\n".join([f"{titulo} - {usuario.get_username()}", *linhas])
    if len(texto) > TELEGRAM_LIMITE_CARACTERES:
        texto = texto[:TELEGRAM_LIMITE_CARACTERES - 1] + "…"
    return enviar_notificacao_telegram(usuario, texto)


def notificar_compras_scanner_telegram(usuario, resultados_scanner: list[dict] | None = None) -> list[Alerta]:
    """
    Aviso de compra do Scanner Técnico pelo Telegram: para cada ativo
    acompanhado (comprado ou reservado) com veredito PREDOMINIO_ALTA e plano
    de trade, quando a última cotação está na entrada sugerida ou abaixo
    dela. Cada ativo é avisado no máximo uma vez por dia (registra um
    Alerta.SCANNER_COMPRA, que também aparece na tela de alertas).

    `resultados_scanner` é a saída de escanear_carteira - quem já calculou
    (a tela do Scanner) passa a lista pronta pra não escanear duas vezes.
    """
    if credenciais_telegram(usuario) is None:
        return []
    if resultados_scanner is None:
        resultados_scanner = escanear_carteira(usuario)

    hoje = timezone.localdate()
    novos_alertas = []
    for item in resultados_scanner:
        plano = item.get("plano_trade")
        if item["veredito"] != "PREDOMINIO_ALTA" or not plano:
            continue
        ativo = item["ativo"]
        cotacao = ativo.ultima_cotacao()
        if cotacao is None or float(cotacao.preco_fechamento) > plano["entrada"]:
            continue
        if Alerta.objects.filter(
            usuario=usuario, ativo=ativo, tipo=Alerta.SCANNER_COMPRA, criado_em__date=hoje,
        ).exists():
            continue

        rr = f", R:R 1:{plano['relacao_risco_retorno']}" if plano.get("relacao_risco_retorno") else ""
        situacao = "reservado" if item.get("apenas_reservado") else "em carteira"
        mensagem = (
            f"{ativo.ticker} ({situacao}): cotação R$ {cotacao.preco_fechamento} na entrada "
            f"R$ {plano['entrada']} - stop R$ {plano['stop']}, alvo R$ {plano['alvo']}{rr}."
        )
        novos_alertas.append(Alerta.objects.create(
            usuario=usuario, ativo=ativo, tipo=Alerta.SCANNER_COMPRA, mensagem=mensagem[:255],
        ))

    if novos_alertas:
        _enviar_linhas_telegram(
            usuario, "🛒 Scanner Técnico: oportunidade de compra",
            [f"• {alerta.mensagem}" for alerta in novos_alertas]
            + ["Sinal técnico, não é recomendação de investimento."],
        )
    return novos_alertas


def notificar_reservas_atingidas_telegram(usuario) -> list[Alerta]:
    """
    Avisa pelo Telegram quando a última cotação de um ativo reservado
    (Operacao RESERVAR) cai até o preço da reserva ou abaixo - um aviso por
    ativo por dia (Alerta.RESERVA_ATINGIDA), usando o maior preço reservado
    quando há mais de uma reserva do mesmo ativo. A reserva automática que o
    robô consultor criou hoje fica de fora: ela nasce com a cotação do momento
    como preço, e o sinal de compra do robô já foi avisado.
    """
    if credenciais_telegram(usuario) is None:
        return []

    hoje = timezone.localdate()
    reservas = (
        Operacao.objects.filter(usuario=usuario, tipo=Operacao.RESERVAR)
        .exclude(data_operacao=hoje, observacao=OBSERVACAO_RESERVA_ROBO)
        .select_related("ativo")
    )
    preco_reservado_por_ativo = {}
    for reserva in reservas:
        atual = preco_reservado_por_ativo.get(reserva.ativo)
        if atual is None or reserva.preco_unitario > atual:
            preco_reservado_por_ativo[reserva.ativo] = reserva.preco_unitario

    novos_alertas = []
    for ativo, preco_reservado in sorted(preco_reservado_por_ativo.items(), key=lambda par: par[0].ticker):
        cotacao = ativo.ultima_cotacao()
        if cotacao is None or cotacao.preco_fechamento > preco_reservado:
            continue
        if Alerta.objects.filter(
            usuario=usuario, ativo=ativo, tipo=Alerta.RESERVA_ATINGIDA, criado_em__date=hoje,
        ).exists():
            continue
        novos_alertas.append(Alerta.objects.create(
            usuario=usuario, ativo=ativo, tipo=Alerta.RESERVA_ATINGIDA,
            mensagem=(
                f"{ativo.ticker}: cotação R$ {cotacao.preco_fechamento} chegou ao preço da reserva "
                f"(R$ {preco_reservado})."
            ),
        ))

    if novos_alertas:
        _enviar_linhas_telegram(
            usuario, "🎯 Reserva atingida: hora de decidir a compra",
            [f"• {alerta.mensagem}" for alerta in novos_alertas],
        )
    return novos_alertas


def notificar_analise_ia_telegram(usuario, analise_ia: dict) -> bool:
    """
    Envia pelo Telegram o resultado do botão "Análise da B3 hoje (IA)": o
    resumo do mercado e os ativos que a IA classificou como ALTA (candidatos
    a compra), com a pontuação do Scanner quando houver. Nada é enviado se
    nenhum ativo saiu como ALTA.
    """
    linhas_ia = analise_ia.get("linhas") or []
    altas = [l for l in linhas_ia if not l["eh_mercado_geral"] and l["situacao"] == "ALTA"]
    if not altas or credenciais_telegram(usuario) is None:
        return False

    linhas = [f"B3 hoje: {l['comentario']}" for l in linhas_ia if l["eh_mercado_geral"]]
    for linha in altas:
        score = f" ({linha['score']}/100)" if linha.get("score") is not None else ""
        linhas.append(f"• {linha['ticker']}{score}: {linha['comentario']}")
    linhas.append("Leitura da IA, não é recomendação de investimento.")
    return _enviar_linhas_telegram(usuario, "🤖 Análise da B3 hoje (IA): sinais de compra", linhas)


def notificar_sugestoes_ia_telegram(usuario, sugestoes_ia: dict, maiores_altas: list[dict], maiores_baixas: list[dict]) -> bool:
    """
    Envia pelo Telegram as "Sugestões de hoje (IA)", com a cotação e a
    variação do dia de cada ativo (tiradas das maiores altas/baixas que o
    sistema já buscou - não vêm da IA).
    """
    sugestoes = sugestoes_ia.get("linhas") or []
    if not sugestoes or credenciais_telegram(usuario) is None:
        return False

    por_ticker = {a.get("stock", "").upper(): a for a in [*maiores_altas, *maiores_baixas]}
    linhas = []
    for sugestao in sugestoes:
        dados = por_ticker.get(sugestao["ticker"], {})
        detalhes = []
        if dados.get("close") is not None:
            detalhes.append(f"R$ {dados['close']}")
        if dados.get("change") is not None:
            detalhes.append(f"{dados['change']:+.2f}%" if isinstance(dados["change"], (int, float)) else f"{dados['change']}%")
        sufixo = f" ({', '.join(detalhes)})" if detalhes else ""
        linhas.append(f"• {sugestao['ticker']}{sufixo}: {sugestao['justificativa']}")
    linhas.append("Sugestões da IA, não é recomendação de investimento.")
    return _enviar_linhas_telegram(usuario, "💡 Sugestões de hoje (IA)", linhas)


def enviar_ia_diaria_telegram(usuario, cache: dict) -> bool:
    """
    Envio automático (uma vez por dia útil, a partir de
    settings.TELEGRAM_IA_DIARIA_HORA) da Análise da B3 hoje e das Sugestões
    de hoje (IA) pelo Telegram, chamado pelo ciclo de cotações (ver
    executar_ciclo_atualizacao_cotacoes). O envio do dia fica marcado num
    Alerta.IA_DIARIA; se a IA ou a brapi falharem, não marca, e o próximo
    ciclo tenta de novo.

    `cache` é um dict compartilhado entre os usuários do mesmo ciclo: as
    maiores altas/baixas e as sugestões (iguais pra todos) são buscadas uma
    vez só - só a análise depende da carteira de cada um.
    """
    agora = timezone.localtime()
    if (
        not settings.TELEGRAM_IA_DIARIA or not ia_configurada()
        or agora.weekday() >= 5 or agora.hour < settings.TELEGRAM_IA_DIARIA_HORA
        or credenciais_telegram(usuario) is None
        or Alerta.objects.filter(usuario=usuario, tipo=Alerta.IA_DIARIA, criado_em__date=agora.date()).exists()
    ):
        return False

    try:
        if "variacoes" not in cache:
            cache["variacoes"] = buscar_maiores_variacoes(limite=10)
        maiores_altas, maiores_baixas = cache["variacoes"]
        if "sugestoes" not in cache:
            cache["sugestoes"] = gerar_sugestoes_ia(maiores_altas, maiores_baixas)
        analise = gerar_analise_b3_ia(escanear_carteira(usuario), maiores_altas, maiores_baixas)
    except (BrapiError, IAError) as exc:
        logger.warning("IA diária pelo Telegram não enviada para %s: %s", usuario.get_username(), exc)
        return False

    enviou_analise = notificar_analise_ia_telegram(usuario, analise)
    enviou_sugestoes = notificar_sugestoes_ia_telegram(usuario, cache["sugestoes"], maiores_altas, maiores_baixas)
    Alerta.objects.create(
        usuario=usuario, tipo=Alerta.IA_DIARIA,
        mensagem="Análise da B3 hoje e Sugestões de hoje (IA) enviadas automaticamente pelo Telegram.",
    )
    return enviou_analise or enviou_sugestoes


# --------------------------------------------------------------------------
# Geração de alertas/lembretes
# --------------------------------------------------------------------------
def gerar_alertas_para_usuario(usuario) -> list[Alerta]:
    """
    Verifica as posições do usuário contra as metas de lucro/perda definidas
    e cria alertas quando alguma meta é atingida (evita duplicar alertas do
    mesmo dia para o mesmo ativo/tipo).
    """
    hoje = timezone.localdate()
    novos_alertas = []

    for posicao in calcular_posicoes(usuario):
        if posicao.lucro_perda_pct is None:
            continue

        meta_lucro = abs(
            posicao.meta_lucro_pct if posicao.meta_lucro_pct is not None else Decimal(
                str(settings.META_LUCRO_PADRAO)
            )
        )
        # a meta de perda é sempre uma queda (valor negativo) - normaliza caso
        # tenha sido salva como positiva (dado antigo, de antes da validação)
        meta_perda = -abs(
            posicao.meta_perda_pct if posicao.meta_perda_pct is not None else Decimal(
                str(settings.META_PERDA_PADRAO)
            )
        )

        tipo_disparado = None
        if posicao.lucro_perda_pct >= meta_lucro:
            tipo_disparado = Alerta.LUCRO
            mensagem = (
                f"{posicao.ativo.ticker}: meta de lucro atingida "
                f"({posicao.lucro_perda_pct}% ≥ {meta_lucro}%)."
            )
        elif posicao.lucro_perda_pct <= meta_perda:
            tipo_disparado = Alerta.PERDA
            mensagem = (
                f"{posicao.ativo.ticker}: meta de perda atingida "
                f"({posicao.lucro_perda_pct}% ≤ {meta_perda}%)."
            )

        if tipo_disparado:
            ja_existe = Alerta.objects.filter(
                usuario=usuario,
                ativo=posicao.ativo,
                tipo=tipo_disparado,
                criado_em__date=hoje,
            ).exists()
            if not ja_existe:
                alerta = Alerta.objects.create(
                    usuario=usuario,
                    ativo=posicao.ativo,
                    tipo=tipo_disparado,
                    mensagem=mensagem,
                    percentual=posicao.lucro_perda_pct,
                )
                novos_alertas.append(alerta)

        if posicao.tendencia in ("ALTA", "BAIXA"):
            ja_existe_tendencia = Alerta.objects.filter(
                usuario=usuario,
                ativo=posicao.ativo,
                tipo=Alerta.TENDENCIA,
                criado_em__date=hoje,
            ).exists()
            if not ja_existe_tendencia:
                rotulo = TENDENCIA_LABELS.get(posicao.tendencia, posicao.tendencia)
                alerta = Alerta.objects.create(
                    usuario=usuario,
                    ativo=posicao.ativo,
                    tipo=Alerta.TENDENCIA,
                    mensagem=f"{posicao.ativo.ticker}: sinal de {rotulo.lower()} identificado no mercado.",
                )
                novos_alertas.append(alerta)

    notificar_alertas_telegram(usuario, novos_alertas)
    return novos_alertas


def ativo_ja_comprado(usuario, ativo: Ativo) -> bool:
    """True quando o usuário já tem saldo comprado (em carteira) deste ativo."""
    return any(
        op.saldo > 0
        for op in Operacao.objects.filter(usuario=usuario, ativo=ativo, tipo=Operacao.COMPRA)
    )


def ativos_distintos_comprados(usuario) -> set[int]:
    """
    IDs dos ativos que o usuário tem efetivamente em carteira (saldo comprado
    > 0 em pelo menos um lote) - usado para aplicar o limite de ativos
    diferentes em carteira (MAX_ATIVOS_EM_CARTEIRA no .env, ver
    core.forms.OperacaoForm e core.forms.ConfirmarCompraForm).
    """
    ids = set()
    for op in Operacao.objects.filter(usuario=usuario, tipo=Operacao.COMPRA):
        if op.saldo > 0:
            ids.add(op.ativo_id)
    return ids


def gerar_sinais_robo_para_usuario(usuario) -> list[Alerta]:
    """
    "Robô consultor": aplica os indicadores técnicos já calculados pelo
    sistema (tendência por média móvel, RSI, MACD - ver
    analisar_indicadores_tecnicos) a cada ativo que o usuário acompanha
    (comprado ou reservado) e cria um alerta quando o sinal geral consolidado
    aponta claramente compra ou venda - evita duplicar o mesmo sinal no
    mesmo dia para o mesmo ativo.

    Sinal de COMPRA é ignorado pra ativos que o usuário já tem em carteira
    (saldo comprado > 0) - a sugestão de compra é pra ajudar a decidir sobre
    ativos que ele ainda não tem, não pra quem já está comprado (esse caso já
    é coberto pelo alerta de meta de lucro/perda). Quando o sinal de compra
    se aplica, além do alerta, cadastra automaticamente uma reserva (Operacao
    tipo RESERVAR, ver _reservar_automaticamente) pro ativo - a sugestão vira
    um item de verdade na grid "Reservadas" de Minhas Operações, onde o
    usuário decide se efetiva a compra (o robô não executa nenhuma ordem
    sozinho). Sinal de VENDA continua sendo avaliado normalmente mesmo pra
    quem já está comprado - é justamente o caso que interessa avisar.

    É só um apoio técnico à decisão - não constitui recomendação de
    investimento.
    """
    hoje = timezone.localdate()
    novos_alertas = []

    ativos = Ativo.objects.filter(operacoes__usuario=usuario).distinct()
    for ativo in ativos:
        indicadores = analisar_indicadores_tecnicos(ativo)
        sinal = indicadores["sinal_geral"]
        if sinal not in ("COMPRA", "VENDA"):
            continue

        if sinal == "COMPRA" and ativo_ja_comprado(usuario, ativo):
            continue

        tipo = Alerta.SINAL_COMPRA if sinal == "COMPRA" else Alerta.SINAL_VENDA
        ja_existe = Alerta.objects.filter(
            usuario=usuario, ativo=ativo, tipo=tipo, criado_em__date=hoje,
        ).exists()
        if not ja_existe:
            mensagem = (
                f"{ativo.ticker}: robô identificou sinal de {sinal.lower()} "
                f"({indicadores['tendencia_label'].lower()}, RSI: {indicadores['rsi_label'].lower()}, "
                f"MACD: {indicadores['macd_label'].lower()}) - sinal técnico, não é recomendação de investimento."
            )
            alerta = Alerta.objects.create(usuario=usuario, ativo=ativo, tipo=tipo, mensagem=mensagem)
            novos_alertas.append(alerta)

        if sinal == "COMPRA":
            _reservar_automaticamente(usuario, ativo)

    notificar_alertas_telegram(usuario, novos_alertas)
    return novos_alertas


OBSERVACAO_RESERVA_ROBO = "Reserva automática do robô consultor - sinal técnico de compra (RSI/MACD/tendência)."


def _reservar_automaticamente(usuario, ativo: Ativo) -> None:
    """
    Cadastra uma reserva (intenção de compra) pro ativo quando o robô
    consultor identifica um sinal de compra, usando a cotação mais recente
    como preço pretendido - assim aparece como item acionável na grid
    "Reservadas" de Minhas Operações (o usuário decide se compra, editando o
    preço/quantidade se quiser, ou excluindo a reserva).

    Não duplica: se o usuário já tem uma reserva em aberto pra esse ativo,
    não mexe nela.
    """
    ja_reservado = Operacao.objects.filter(usuario=usuario, ativo=ativo, tipo=Operacao.RESERVAR).exists()
    if ja_reservado:
        return

    cotacao = ativo.ultima_cotacao()
    if cotacao is None:
        return  # sem cotação ainda não há preço pretendido pra registrar

    Operacao.objects.create(
        usuario=usuario,
        ativo=ativo,
        tipo=Operacao.RESERVAR,
        quantidade=1,
        preco_unitario=cotacao.preco_fechamento,
        data_operacao=timezone.localdate(),
        observacao=OBSERVACAO_RESERVA_ROBO,
    )


def limpar_alertas_antigos() -> int:
    """
    Remove os avisos/lembretes (Alerta) de TODOS os usuários que não são de
    hoje - mantém só os do dia atual. Retorna quantos foram apagados.
    """
    hoje = timezone.localdate()
    apagados, _ = Alerta.objects.exclude(criado_em__date=hoje).delete()
    return apagados


# --------------------------------------------------------------------------
# Relatório de indicação do robô consultor (Excel / PDF)
#
# Reflete a tabela "Indicadores técnicos" da página Análise de Mercado (ver
# core.views._sinais_robo): o sinal geral, a tendência, o RSI e o MACD de
# cada ativo acompanhado, mais se o ativo já está em carteira (contexto de
# por que um sinal de compra pode não ter virado uma reserva - ver
# core.services.gerar_sinais_robo_para_usuario).
# --------------------------------------------------------------------------
COLUNAS_ROBO = [
    "Ativo", "Em carteira", "Tendência", "RSI", "MACD", "Sinal geral", "Votos compra", "Votos venda",
]


def _linha_robo(item: dict) -> list:
    return [
        item["ativo"].ticker,
        "Sim" if item["em_carteira"] else "Não",
        item["tendencia_label"],
        item["rsi"],
        item["macd_label"],
        item["sinal_geral_label"],
        item["votos_compra"],
        item["votos_venda"],
    ]


def gerar_excel_robo(sinais: list[dict]) -> bytes:
    """
    Gera uma planilha .xlsx com a indicação atual do robô consultor (sinal
    geral, tendência, RSI e MACD) para cada ativo acompanhado, pronta para
    download.
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Indicação do robô"

    ws.append(["Indicação do robô consultor"])
    ws.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws.append(["Sinal técnico (RSI, MACD, tendência) - não é recomendação de investimento."])
    ws.append([])
    ws.append(COLUNAS_ROBO)
    for celula in ws[5]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D1526")
        celula.alignment = Alignment(horizontal="center")

    for item in sinais:
        ws.append(_linha_robo(item))

    for indice in range(1, len(COLUNAS_ROBO) + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 18
    ws.freeze_panes = "A6"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_robo(sinais: list[dict], nome_usuario: str) -> bytes:
    """
    Gera um PDF com a indicação atual do robô consultor (sinal geral,
    tendência, RSI e MACD) para cada ativo acompanhado, pronto para download.
    """
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()

    elementos = [
        Paragraph("BolsaTrader - Indicação do Robô Consultor", estilos["Title"]),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')} - "
            "sinal técnico (RSI, MACD, tendência), não é recomendação de investimento.",
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    if not sinais:
        elementos.append(Paragraph("Nenhum ativo para analisar.", estilos["Normal"]))
    else:
        dados = [COLUNAS_ROBO]
        for item in sinais:
            linha = _linha_robo(item)
            dados.append([
                linha[0], linha[1], linha[2],
                f"{linha[3]:.2f}" if linha[3] is not None else "—",
                linha[4], linha[5], linha[6], linha[7],
            ])

        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela)

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Notícias de mercado (fontes configuráveis, raspagem simples de manchetes)
# --------------------------------------------------------------------------
class NoticiaScrapingError(Exception):
    """Erro ao buscar ou ler as manchetes de uma fonte de notícias."""


_NOTICIAS_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)


def buscar_manchetes(url: str, limite: int = 15) -> list[dict]:
    """
    Busca a página em `url` e extrai as manchetes (título + link) mais
    prováveis: primeiro títulos (h1/h2/h3) que contêm um link, depois
    qualquer link com texto longo o bastante pra parecer uma manchete de
    verdade - descarta menus, rodapés e links de uma palavra só.

    Cada site de notícias tem uma estrutura de HTML diferente, então isso é
    uma heurística genérica (não um parser dedicado a cada site) - captura a
    maioria dos títulos, mas pode incluir algum link que não é bem uma
    notícia, ou deixar escapar alguma manchete com marcação incomum.
    """
    try:
        resposta = requests.get(url, headers={"User-Agent": _NOTICIAS_USER_AGENT}, timeout=15)
        resposta.raise_for_status()
    except requests.RequestException as exc:
        raise NoticiaScrapingError(f"Falha ao acessar {url}: {exc}") from exc

    soup = BeautifulSoup(resposta.text, "html.parser")

    encontradas = []
    vistos = set()

    candidatos = soup.select("h1 a[href], h2 a[href], h3 a[href]") + soup.select("a[href]")
    for link in candidatos:
        titulo = link.get_text(strip=True)
        href = link.get("href")
        if not titulo or not href or len(titulo) < 25:
            continue  # texto curto demais pra ser um título de matéria (menu, "Entrar", etc.)

        link_absoluto = urljoin(url, href)
        if link_absoluto in vistos:
            continue
        vistos.add(link_absoluto)

        encontradas.append({"titulo": titulo, "url": link_absoluto})
        if len(encontradas) >= limite:
            break

    return encontradas


def atualizar_noticias_fonte(fonte: FonteNoticia, limite: int = 15) -> list[int]:
    """
    Busca as manchetes que o site está exibindo agora (as "notícias do dia")
    e sincroniza a fonte com elas: grava as que ainda não existem
    (identificadas pelo link - já vistas não duplicam) e exclui as manchetes
    antigas que não estão mais na página (não são mais exibidas pelo site).

    Se o site não retornar nenhuma manchete reconhecível, trata como falha de
    leitura (NoticiaScrapingError) em vez de apagar o histórico existente -
    evita esvaziar a fonte por causa de uma leitura ruim.

    Retorna os ids das notícias novas gravadas (não só a contagem) - usado
    pela tela de Notícias do Mercado para destacar as recém-encontradas.
    """
    manchetes = buscar_manchetes(fonte.url, limite=limite)
    if not manchetes:
        raise NoticiaScrapingError(f"Nenhuma notícia encontrada em {fonte.url}.")

    urls_atuais = {item["url"] for item in manchetes}

    novas_ids = []
    for item in manchetes:
        noticia, criada = Noticia.objects.get_or_create(
            url=item["url"],
            defaults={"fonte": fonte, "titulo": item["titulo"]},
        )
        if criada:
            novas_ids.append(noticia.id)

    fonte.noticias.exclude(url__in=urls_atuais).delete()

    return novas_ids


def _sem_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn").upper()


def termos_destaque_noticias(texto: str) -> list[str]:
    """
    Termos usados pra destacar as manchetes de um ativo na tela de Notícias do
    Mercado. O usuário digita um ou mais termos separados por vírgula (ex:
    "PETR4, Petrobras"); para um ticker de ação (ex: PETR4) entra também a raiz
    de 4 letras (PETR), que costuma aparecer em manchetes sobre PETR3/PETR4.
    """
    termos = []
    for parte in texto.split(","):
        termo = _sem_acentos(parte.strip())
        if not termo:
            continue
        termos.append(termo)
        raiz = re.match(r"^([A-Z]{4})\d{1,2}$", termo)
        if raiz:
            termos.append(raiz.group(1))
    return list(dict.fromkeys(termos))


def noticia_cita_termos(titulo: str, termos: list[str]) -> bool:
    """True se o título contém algum dos termos como palavra inteira (sem diferenciar maiúsculas/acentos) - "PETR" casa com "PETR3"/"PETR4"."""
    titulo_normalizado = _sem_acentos(titulo)
    return any(re.search(rf"(?<![A-Z0-9]){re.escape(t)}(?![A-Z])", titulo_normalizado) for t in termos)


# --------------------------------------------------------------------------
# Conta corrente (caixa das operações de bolsa)
# --------------------------------------------------------------------------
def obter_ou_criar_conta_corrente(usuario) -> ContaCorrente:
    """Cada usuário tem no máximo uma conta corrente - cria na primeira vez que for preciso (saldo inicial R$ 0)."""
    conta, _ = ContaCorrente.objects.get_or_create(usuario=usuario)
    return conta


def saldo_conta_corrente(conta: ContaCorrente) -> Decimal:
    """Saldo atual = saldo inicial + soma dos créditos - soma dos débitos de todos os lançamentos."""
    total_creditos = conta.lancamentos.filter(
        tipo=LancamentoContaCorrente.CREDITO
    ).aggregate(total=Sum("valor"))["total"] or Decimal("0")
    total_debitos = conta.lancamentos.filter(
        tipo=LancamentoContaCorrente.DEBITO
    ).aggregate(total=Sum("valor"))["total"] or Decimal("0")
    return conta.saldo_inicial + total_creditos - total_debitos


def sincronizar_lancamento_compra(operacao: Operacao) -> None:
    """
    Cria/atualiza o lançamento de DÉBITO (compra) da conta corrente do usuário
    com o valor/quantidade/data atuais da Operacao - chamar sempre depois de
    criar, efetivar (reserva -> compra) ou editar uma compra. Reservas
    (Operacao.RESERVAR) não geram lançamento, já que ainda não é dinheiro
    comprometido de fato.
    """
    # lote gerado por evento corporativo (desdobramento etc.) não é dinheiro saindo
    if operacao.tipo != Operacao.COMPRA or operacao.evento_id:
        return

    conta = obter_ou_criar_conta_corrente(operacao.usuario)
    LancamentoContaCorrente.objects.update_or_create(
        operacao=operacao,
        origem=LancamentoContaCorrente.ORIGEM_COMPRA,
        defaults={
            "conta": conta,
            "tipo": LancamentoContaCorrente.DEBITO,
            "valor": operacao.quantidade_comprada * operacao.preco_unitario + (operacao.taxas_compra or Decimal("0")),
            "descricao": f"Compra de {operacao.quantidade}x {operacao.ativo.ticker}",
            "data": operacao.data_operacao,
        },
    )


def recalcular_resumo_vendas(operacao: Operacao) -> None:
    """
    Atualiza os campos de resumo do lote (quantidade vendida, preço médio de
    venda, data e percentual da corretora da última venda) a partir dos
    VendaLote dele, e sincroniza os lançamentos da conta corrente - chamar
    sempre depois de registrar, corrigir ou excluir uma venda.
    """
    vendas = list(operacao.vendas.order_by("data", "pk"))
    if vendas:
        quantidade = sum(venda.quantidade for venda in vendas)
        total = sum((venda.valor_bruto for venda in vendas), Decimal("0"))
        operacao.quantidade_vendida = quantidade
        operacao.preco_venda = (total / quantidade).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        operacao.data_venda = vendas[-1].data
        operacao.percentual_corretora = vendas[-1].percentual_corretora
    else:
        operacao.quantidade_vendida = 0
        operacao.preco_venda = None
        operacao.data_venda = None
        operacao.percentual_corretora = None
    operacao.save(update_fields=["quantidade_vendida", "preco_venda", "data_venda", "percentual_corretora"])
    sincronizar_lancamento_venda(operacao)


def sincronizar_lancamento_venda(operacao: Operacao) -> None:
    """
    Cria/atualiza um lançamento de CRÉDITO na conta corrente para cada venda
    (VendaLote) deste lote de compra, com o valor líquido (já descontadas
    corretora e taxas), que é o que realmente entra na conta, na data da
    própria venda. Remove os lançamentos de vendas que não existem mais (ex:
    venda excluída). Lote sem VendaLote mas com o resumo preenchido (dados
    antigos/gravados direto no lote) ganha um lançamento único, sem venda_lote.
    """
    existentes = LancamentoContaCorrente.objects.filter(
        operacao=operacao, origem=LancamentoContaCorrente.ORIGEM_VENDA
    )
    vendas = operacao.vendas_registradas()
    if not vendas:
        existentes.delete()
        return

    conta = obter_ou_criar_conta_corrente(operacao.usuario)
    vendas_salvas = [venda for venda in vendas if venda.pk]
    if not vendas_salvas:
        existentes.exclude(venda_lote__isnull=True).delete()
        venda = vendas[0]
        LancamentoContaCorrente.objects.update_or_create(
            operacao=operacao,
            origem=LancamentoContaCorrente.ORIGEM_VENDA,
            venda_lote=None,
            defaults={
                "conta": conta,
                "tipo": LancamentoContaCorrente.CREDITO,
                "valor": venda.valor_liquido,
                "descricao": f"Venda de {venda.quantidade}x {operacao.ativo.ticker}",
                "data": venda.data or timezone.localdate(),
            },
        )
        return

    existentes.exclude(venda_lote__in=vendas_salvas).delete()
    for venda in vendas_salvas:
        LancamentoContaCorrente.objects.update_or_create(
            venda_lote=venda,
            defaults={
                "conta": conta,
                "operacao": operacao,
                "origem": LancamentoContaCorrente.ORIGEM_VENDA,
                "tipo": LancamentoContaCorrente.CREDITO,
                "valor": venda.valor_liquido,
                "descricao": f"Venda de {venda.quantidade}x {operacao.ativo.ticker}",
                "data": venda.data,
            },
        )


def sincronizar_lancamento_compra_cripto(operacao: OperacaoCripto) -> None:
    """Como sincronizar_lancamento_compra, mas para OperacaoCripto - mesma conta corrente do usuário, lançamento próprio (operacao_cripto), não misturado com os de ações."""
    if operacao.tipo != OperacaoCripto.COMPRA:
        return

    conta = obter_ou_criar_conta_corrente(operacao.usuario)
    LancamentoContaCorrente.objects.update_or_create(
        operacao_cripto=operacao,
        origem=LancamentoContaCorrente.ORIGEM_COMPRA_CRIPTO,
        defaults={
            "conta": conta,
            "tipo": LancamentoContaCorrente.DEBITO,
            "valor": operacao.valor_total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
            "descricao": f"Compra de {operacao.quantidade}x {operacao.cripto_ativo.coin}",
            "data": operacao.data_operacao,
        },
    )


def sincronizar_lancamento_venda_cripto(operacao: OperacaoCripto) -> None:
    """Como sincronizar_lancamento_venda, mas para OperacaoCripto (sem percentual de corretora - esse campo não existe em OperacaoCripto)."""
    if not operacao.quantidade_vendida or operacao.preco_venda is None:
        LancamentoContaCorrente.objects.filter(
            operacao_cripto=operacao, origem=LancamentoContaCorrente.ORIGEM_VENDA_CRIPTO
        ).delete()
        return

    conta = obter_ou_criar_conta_corrente(operacao.usuario)
    valor = operacao.valor_total_vendido.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    LancamentoContaCorrente.objects.update_or_create(
        operacao_cripto=operacao,
        origem=LancamentoContaCorrente.ORIGEM_VENDA_CRIPTO,
        defaults={
            "conta": conta,
            "tipo": LancamentoContaCorrente.CREDITO,
            "valor": valor,
            "descricao": f"Venda de {operacao.quantidade_vendida}x {operacao.cripto_ativo.coin}",
            "data": operacao.data_venda or timezone.localdate(),
        },
    )


def registrar_transferencia_conta_corrente(
    usuario, tipo: str, valor: Decimal, descricao: str, data: date | None = None
) -> LancamentoContaCorrente:
    """Registra uma transferência manual (crédito ou débito) de/para outra conta (ex: Nubank, corretora)."""
    conta = obter_ou_criar_conta_corrente(usuario)
    return LancamentoContaCorrente.objects.create(
        conta=conta,
        tipo=tipo,
        origem=LancamentoContaCorrente.ORIGEM_TRANSFERENCIA,
        valor=valor,
        descricao=descricao,
        data=data or timezone.localdate(),
    )


def extrato_conta_corrente(
    conta: ContaCorrente, data_inicio: date | None = None, data_fim: date | None = None
) -> list[dict]:
    """
    Lançamentos da conta corrente no período informado (todos, se nenhuma
    data for passada), com o saldo acumulado (saldo_apos) já calculado a
    partir do saldo inicial mais tudo que aconteceu antes do período - assim
    o extrato bate com o saldo real da conta mesmo filtrando só uma fatia do
    tempo. Devolvido do lançamento mais recente para o mais antigo, como as
    outras grids do sistema.
    """
    todos = conta.lancamentos.select_related("operacao__ativo", "operacao_cripto__cripto_ativo").order_by("data", "criado_em")
    saldo = conta.saldo_inicial
    linhas = []
    for lancamento in todos:
        sinal = 1 if lancamento.tipo == LancamentoContaCorrente.CREDITO else -1

        if data_inicio and lancamento.data < data_inicio:
            saldo += sinal * lancamento.valor
            continue
        if data_fim and lancamento.data > data_fim:
            continue

        saldo += sinal * lancamento.valor
        linhas.append({"lancamento": lancamento, "saldo_apos": saldo})

    linhas.reverse()
    return linhas


COLUNAS_EXTRATO_CONTA_CORRENTE = ["Data", "Tipo", "Origem", "Descrição", "Valor (R$)", "Saldo após (R$)"]


def _linha_extrato_conta_corrente(item: dict) -> list:
    lancamento = item["lancamento"]
    sinal = 1 if lancamento.tipo == LancamentoContaCorrente.CREDITO else -1
    return [
        lancamento.data,
        lancamento.get_tipo_display(),
        lancamento.get_origem_display(),
        lancamento.descricao,
        float(sinal * lancamento.valor),
        float(item["saldo_apos"]),
    ]


def gerar_excel_extrato_conta_corrente(linhas: list[dict], saldo_inicial: Decimal, saldo_atual: Decimal) -> bytes:
    """Gera uma planilha .xlsx com o extrato da conta corrente do usuário, pronta para download."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato conta corrente"

    ws.append(["Extrato da Conta Corrente"])
    ws.append([f"Gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([f"Saldo inicial: R$ {saldo_inicial}", f"Saldo atual: R$ {saldo_atual}"])
    ws.append([])
    ws.append(COLUNAS_EXTRATO_CONTA_CORRENTE)
    for celula in ws[5]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0D1526")
        celula.alignment = Alignment(horizontal="center")

    for item in linhas:
        linha = _linha_extrato_conta_corrente(item)
        linha[0] = linha[0].strftime("%d/%m/%Y")
        ws.append(linha)

    for indice in range(1, len(COLUNAS_EXTRATO_CONTA_CORRENTE) + 1):
        ws.column_dimensions[get_column_letter(indice)].width = 22
    ws.freeze_panes = "A6"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_extrato_conta_corrente(
    linhas: list[dict], saldo_inicial: Decimal, saldo_atual: Decimal, periodo_label: str, nome_usuario: str
) -> bytes:
    """Gera um PDF com o extrato da conta corrente do usuário, pronto para download."""
    import io
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    nome_usuario = escape(nome_usuario)
    periodo_label = escape(periodo_label)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()

    elementos = [
        Paragraph("BolsaTrader - Extrato da Conta Corrente", estilos["Title"]),
        Paragraph(
            f"{nome_usuario} - gerado em {timezone.localtime().strftime('%d/%m/%Y %H:%M')}",
            estilos["Normal"],
        ),
        Paragraph(periodo_label, estilos["Normal"]),
        Paragraph(
            f"Saldo inicial: R$ {saldo_inicial:.2f} — Saldo atual: R$ {saldo_atual:.2f}",
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    if not linhas:
        elementos.append(Paragraph("Nenhum lançamento no período selecionado.", estilos["Normal"]))
    else:
        dados = [COLUNAS_EXTRATO_CONTA_CORRENTE]
        for item in linhas:
            linha = _linha_extrato_conta_corrente(item)
            dados.append([
                linha[0].strftime("%d/%m/%Y"), linha[1], linha[2], linha[3],
                f"{linha[4]:.2f}", f"{linha[5]:.2f}",
            ])

        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0D1526")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        elementos.append(tabela)

    doc.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Post-it (bloco de notas pessoal e fixo, nas telas Menu/Operações/Posições)
# --------------------------------------------------------------------------
def obter_post_it(usuario) -> PostIt | None:
    """
    O post-it do usuário, se ele já tiver escrito algo alguma vez - None
    quando ainda não existe. Só busca (nunca cria): é chamado pelo context
    processor em TODA requisição autenticada (ver core.context_processors.
    post_it), então criar um registro vazio aqui gravaria no banco à toa para
    quem nunca usou o post-it. A criação de verdade só acontece ao salvar
    (ver salvar_post_it).
    """
    return PostIt.objects.filter(usuario=usuario).first()


def salvar_post_it(usuario, texto: str | None = None, minimizado: bool | None = None) -> PostIt:
    """
    Cria (na primeira vez) ou atualiza o post-it do usuário - só grava os
    campos realmente informados, então salvar só o estado de minimizado não
    mexe no texto e vice-versa.
    """
    post_it, _ = PostIt.objects.get_or_create(usuario=usuario)
    campos = []
    if texto is not None:
        post_it.texto = texto
        campos.append("texto")
    if minimizado is not None:
        post_it.minimizado = minimizado
        campos.append("minimizado")
    if campos:
        post_it.save(update_fields=campos)
    return post_it


# --------------------------------------------------------------------------
# Criptomoedas - tela própria, à parte das ações da B3 (ver CriptoAtivo,
# OperacaoCripto, CotacaoCripto em core.models). Usa a mesma API brapi.dev
# (endpoint /v2/crypto) e o mesmo orçamento compartilhado de requisições (ver
# _get_brapi) - não é uma conta nem uma chave separada. Cripto negocia 24
# horas por dia, todo santo dia - por isso as funções abaixo não têm (nem
# usam) nenhum portão de "horário de pregão" como mercado_b3_aberto(), que só
# faz sentido pra B3. Sem robô consultor, Scanner Técnico ou comparação com
# Ibovespa/CDI aqui de propósito - esses recursos continuam só sobre ações.
# --------------------------------------------------------------------------
MOEDAS_CRIPTO_SUGERIDAS = ["BTC", "ETH", "BNB", "SOL", "XRP", "ADA", "DOGE"]


def _para_decimal_cripto(valor) -> Decimal | None:
    """Como _para_decimal, mas preservando 8 casas decimais - preço de cripto costuma ter muito mais casas do que preço de ação."""
    if valor is None:
        return None
    try:
        return Decimal(str(valor)).quantize(Decimal("0.00000001"), rounding=ROUND_HALF_UP)
    except (ArithmeticError, ValueError, TypeError):
        return None


def buscar_cotacoes_cripto(coins: list[str], moeda: str = "BRL") -> dict[str, dict]:
    """
    Consulta a cotação atual de uma ou mais criptomoedas na API brapi.dev
    (endpoint /v2/crypto) numa única chamada, não importa quantas moedas
    forem pedidas - devolve um dict {"BTC": {...}, "ETH": {...}} pronto pra
    consulta por código (mesmo espírito de buscar_cotacoes_em_lote, ações).
    """
    if not coins:
        return {}

    url = f"{settings.BRAPI_BASE_URL}/v2/crypto"
    params = {"coin": ",".join(c.upper() for c in coins), "currency": moeda}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN

    try:
        resposta = _get_brapi(url, params=params, timeout=10)
        resposta.raise_for_status()
        dados = resposta.json()
    except requests.RequestException as exc:
        raise BrapiError(_mensagem_falha_api(exc, "cotações de criptomoedas")) from exc

    return {item["coin"].upper(): item for item in (dados.get("coins") or []) if item.get("coin")}


def atualizar_cotacao_cripto_diaria(cripto_ativo: CriptoAtivo, dados_api: dict | None = None) -> CotacaoCripto:
    """Busca (ou recebe) a cotação atual de uma criptomoeda e grava/atualiza o registro do dia - espelha atualizar_cotacao_diaria (ações)."""
    if dados_api is None:
        lote = buscar_cotacoes_cripto([cripto_ativo.coin])
        dados_api = lote.get(cripto_ativo.coin.upper())
        if dados_api is None:
            raise BrapiError(f"Criptomoeda {cripto_ativo.coin} não encontrada na API de cotações.")

    preco = dados_api.get("regularMarketPrice")
    if preco is None:
        raise BrapiError(f"API não retornou preço para {cripto_ativo.coin}.")

    cripto_ativo.nome = dados_api.get("coinName") or cripto_ativo.nome
    cripto_ativo.moeda = dados_api.get("currency") or cripto_ativo.moeda
    cripto_ativo.logo_url = dados_api.get("coinImageUrl") or cripto_ativo.logo_url
    cripto_ativo.preco_atual = _para_decimal_cripto(preco)
    cripto_ativo.variacao_dia_pct = _para_decimal(dados_api.get("regularMarketChangePercent"))
    cripto_ativo.variacao_dia_valor = _para_decimal_cripto(dados_api.get("regularMarketChange"))
    cripto_ativo.maxima_dia = _para_decimal_cripto(dados_api.get("regularMarketDayHigh"))
    cripto_ativo.minima_dia = _para_decimal_cripto(dados_api.get("regularMarketDayLow"))
    cripto_ativo.volume = _para_decimal(dados_api.get("regularMarketVolume"))
    cripto_ativo.valor_mercado = _para_decimal(dados_api.get("marketCap"))

    hora_cotacao = dados_api.get("regularMarketTime")
    if hora_cotacao:
        cripto_ativo.hora_cotacao = parse_datetime(hora_cotacao)

    cripto_ativo.save()

    variacao = dados_api.get("regularMarketChangePercent")
    cotacao, _ = CotacaoCripto.objects.update_or_create(
        cripto_ativo=cripto_ativo,
        data=timezone.localdate(),
        defaults={
            "preco_fechamento": _para_decimal_cripto(preco),
            "variacao_dia_pct": (
                Decimal(str(variacao)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                if variacao is not None else None
            ),
            "maxima": _para_decimal_cripto(dados_api.get("regularMarketDayHigh")),
            "minima": _para_decimal_cripto(dados_api.get("regularMarketDayLow")),
        },
    )
    return cotacao


def atualizar_cotacoes_cripto_ativos(cripto_ativos) -> tuple[int, int]:
    """Atualiza a cotação de uma lista/queryset de criptomoedas com uma única chamada em lote à API - espelha atualizar_cotacoes_ativos (ações). Retorna (atualizados, falhas)."""
    cripto_ativos = list(cripto_ativos)
    try:
        lote = buscar_cotacoes_cripto([c.coin for c in cripto_ativos])
    except BrapiError:
        lote = {}  # orçamento de requisições esgotado (ver _get_brapi) - todos contam como falha

    atualizados, falhas = 0, 0
    for cripto_ativo in cripto_ativos:
        dados_api = lote.get(cripto_ativo.coin.upper())
        try:
            if dados_api is None:
                raise BrapiError(f"Criptomoeda {cripto_ativo.coin} não encontrada na consulta em lote.")
            atualizar_cotacao_cripto_diaria(cripto_ativo, dados_api=dados_api)
            atualizados += 1
        except BrapiError:
            falhas += 1

    return atualizados, falhas


@dataclass
class PosicaoCripto:
    """Espelha Posicao (ações), sem os campos ligados ao robô consultor (que não existe para cripto)."""

    cripto_ativo: CriptoAtivo
    quantidade: Decimal
    preco_medio: Decimal
    valor_investido: Decimal
    preco_atual: Decimal | None = None
    variacao_dia_pct: Decimal | None = None
    valor_atual: Decimal | None = None
    lucro_perda_valor: Decimal | None = None
    lucro_perda_pct: Decimal | None = None
    meta_lucro_pct: Decimal | None = None
    meta_perda_pct: Decimal | None = None
    data_abertura: date | None = None
    quantidade_reservada: Decimal = Decimal("0")

    @property
    def apenas_reservado(self) -> bool:
        return self.quantidade == 0 and self.quantidade_reservada > 0

    @property
    def dias_desde_compra(self) -> int | None:
        if self.data_abertura is None:
            return None
        return (timezone.localdate() - self.data_abertura).days

    @property
    def lucro_perda_por_dia_valor(self) -> Decimal | None:
        """Espelha Posicao.lucro_perda_por_dia_valor (ações)."""
        dias = self.dias_desde_compra
        if self.lucro_perda_valor is None or not dias:
            return None
        return (self.lucro_perda_valor / dias).quantize(Decimal("0.01"))

    @property
    def lucro_perda_por_dia_pct(self) -> Decimal | None:
        """Espelha Posicao.lucro_perda_por_dia_pct (ações)."""
        dias = self.dias_desde_compra
        if self.lucro_perda_pct is None or not dias:
            return None
        return (self.lucro_perda_pct / dias).quantize(Decimal("0.01"))


def calcular_posicoes_cripto(usuario) -> list[PosicaoCripto]:
    """
    Consolida os lotes de compra de criptomoeda do usuário em posições por
    moeda (preço médio do saldo não vendido, quantidade líquida em carteira)
    - mesma lógica de calcular_posicoes (ações), simplificada: sem
    rastreamento de reserva automática do robô consultor, que não existe
    para cripto.
    """
    operacoes = (
        OperacaoCripto.objects.filter(usuario=usuario)
        .select_related("cripto_ativo")
        .order_by("data_operacao", "criado_em")
    )

    agregados = defaultdict(lambda: {
        "quantidade": Decimal("0"),
        "custo_total": Decimal("0"),
        "meta_lucro_pct": None,
        "meta_perda_pct": None,
        "data_abertura": None,
        "quantidade_reservada": Decimal("0"),
        "custo_reservado": Decimal("0"),
    })

    for op in operacoes:
        item = agregados[op.cripto_ativo_id]
        if op.tipo == OperacaoCripto.COMPRA:
            saldo_lote = op.saldo
            if saldo_lote > 0:
                if item["data_abertura"] is None or op.data_operacao < item["data_abertura"]:
                    item["data_abertura"] = op.data_operacao
                item["quantidade"] += saldo_lote
                item["custo_total"] += saldo_lote * op.preco_unitario
        else:  # RESERVAR: intenção de compra futura, não altera a posição real
            item["quantidade_reservada"] += op.quantidade
            item["custo_reservado"] += op.valor_total

        if op.meta_lucro_pct is not None:
            item["meta_lucro_pct"] = op.meta_lucro_pct
        if op.meta_perda_pct is not None:
            item["meta_perda_pct"] = op.meta_perda_pct
        item["cripto_ativo"] = op.cripto_ativo

    posicoes = []
    for item in agregados.values():
        if item["quantidade"] <= 0 and item["quantidade_reservada"] <= 0:
            continue

        cripto_ativo = item["cripto_ativo"]

        if item["quantidade"] > 0:
            preco_medio = (item["custo_total"] / item["quantidade"]).quantize(
                Decimal("0.00000001"), rounding=ROUND_HALF_UP
            )
            valor_investido = (preco_medio * item["quantidade"]).quantize(Decimal("0.01"))
        else:
            preco_medio = (item["custo_reservado"] / item["quantidade_reservada"]).quantize(
                Decimal("0.00000001"), rounding=ROUND_HALF_UP
            )
            valor_investido = Decimal("0.00")

        posicao = PosicaoCripto(
            cripto_ativo=cripto_ativo,
            quantidade=item["quantidade"],
            preco_medio=preco_medio,
            valor_investido=valor_investido,
            meta_lucro_pct=item["meta_lucro_pct"],
            meta_perda_pct=item["meta_perda_pct"],
            data_abertura=item["data_abertura"],
            quantidade_reservada=item["quantidade_reservada"],
        )

        ultima_cotacao = cripto_ativo.ultima_cotacao()
        if ultima_cotacao:
            posicao.preco_atual = ultima_cotacao.preco_fechamento
            posicao.variacao_dia_pct = ultima_cotacao.variacao_dia_pct

        if ultima_cotacao and not posicao.apenas_reservado:
            posicao.valor_atual = (posicao.preco_atual * posicao.quantidade).quantize(Decimal("0.01"))
            posicao.lucro_perda_valor = (posicao.valor_atual - posicao.valor_investido).quantize(Decimal("0.01"))
            if posicao.valor_investido:
                posicao.lucro_perda_pct = (
                    (posicao.lucro_perda_valor / posicao.valor_investido) * 100
                ).quantize(Decimal("0.01"))

        posicoes.append(posicao)

    posicoes.sort(key=lambda p: (
        p.dias_desde_compra is None,
        p.dias_desde_compra if p.dias_desde_compra is not None else 0,
        p.cripto_ativo.coin,
    ))
    return posicoes


def totais_posicoes_cripto(posicoes_cripto: list[PosicaoCripto]) -> dict:
    """
    Totais consolidados das posições de cripto efetivamente COMPRADAS (reservas
    não entram, mesmo critério de calcular_posicoes/dashboard para ações) -
    reaproveitado pela tela Criptomoedas e pelo cartão de resumo do Painel de
    Controle (ver core.views.criptomoedas/dashboard).
    """
    compradas = [p for p in posicoes_cripto if not p.apenas_reservado]
    valor_investido_total = sum((p.valor_investido for p in compradas), Decimal("0"))
    valor_atual_total = sum((p.valor_atual for p in compradas if p.valor_atual is not None), Decimal("0"))
    lucro_perda_total = valor_atual_total - valor_investido_total
    lucro_perda_pct_total = (
        (lucro_perda_total / valor_investido_total) * 100 if valor_investido_total else None
    )
    return {
        "total_moedas": len(compradas),
        "valor_investido_total": valor_investido_total,
        "valor_atual_total": valor_atual_total,
        "lucro_perda_total": lucro_perda_total,
        "lucro_perda_pct_total": lucro_perda_pct_total,
    }


def registrar_atualizacao_carteira_cripto(
    usuario, posicoes_cripto: list[PosicaoCripto] | None = None,
) -> RegistroAtualizacaoCarteiraCripto:
    """
    Grava um retrato (snapshot) dos totais da carteira de criptomoeda do
    usuário - espelha registrar_atualizacao_carteira (ações), chamado toda
    vez que as cotações de cripto são atualizadas (ver core.views.
    cripto_atualizar_cotacoes). Alimenta o gráfico "Variações de hoje" da
    tela Criptomoedas e do Painel de Controle (reaproveita
    construir_grafico_atualizacoes_dia, que só depende dos nomes de campo,
    iguais aos de RegistroAtualizacaoCarteira).
    """
    if posicoes_cripto is None:
        posicoes_cripto = calcular_posicoes_cripto(usuario)

    compradas = [p for p in posicoes_cripto if not p.apenas_reservado]
    valor_investido = sum((p.valor_investido for p in compradas), Decimal("0"))
    valor_atual = sum((p.valor_atual for p in compradas if p.valor_atual is not None), Decimal("0"))
    lucro_perda = valor_atual - valor_investido
    lucro_perda_pct = (
        (lucro_perda / valor_investido * 100).quantize(Decimal("0.01")) if valor_investido else None
    )

    return RegistroAtualizacaoCarteiraCripto.objects.create(
        usuario=usuario,
        total_moedas=len(compradas),
        valor_investido=valor_investido,
        valor_atual=valor_atual,
        lucro_perda=lucro_perda,
        lucro_perda_pct=lucro_perda_pct,
    )


# --------------------------------------------------------------------------
# Contas logadas agora (Painel de Controle, só para administrador/staff) -
# lê as sessões válidas direto de django.contrib.sessions, sem precisar de
# nenhum modelo ou rastreamento próprio. É uma aproximação de "logado agora":
# reflete quem tem uma sessão ainda não expirada (não expulsa por logout nem
# vencida), não necessariamente quem está com a tela aberta neste segundo.
# --------------------------------------------------------------------------
def usuarios_logados_agora() -> list[dict]:
    """
    Uma linha por usuário com pelo menos uma sessão válida (não expirada) no
    momento, com o último login e até quando a sessão mais recente dele
    ainda vale. Decodifica as sessões da tabela django_session - usuários
    com sessão corrompida ou de um cookie secret antigo são ignorados
    silenciosamente (não é motivo pra quebrar o Painel de quem é staff).
    """
    from django.contrib.sessions.models import Session

    agora = timezone.now()
    expira_em_por_usuario_id: dict[int, timezone.datetime] = {}
    for sessao in Session.objects.filter(expire_date__gte=agora):
        try:
            dados = sessao.get_decoded()
        except Exception:
            continue
        usuario_id = dados.get("_auth_user_id")
        if usuario_id is None:
            continue
        usuario_id = int(usuario_id)
        atual = expira_em_por_usuario_id.get(usuario_id)
        if atual is None or sessao.expire_date > atual:
            expira_em_por_usuario_id[usuario_id] = sessao.expire_date

    if not expira_em_por_usuario_id:
        return []

    User = get_user_model()
    usuarios = User.objects.filter(id__in=expira_em_por_usuario_id.keys())
    linhas = [
        {
            "usuario": usuario,
            "expira_em": expira_em_por_usuario_id[usuario.id],
        }
        for usuario in usuarios
    ]
    linhas.sort(key=lambda linha: linha["usuario"].username.lower())
    return linhas


# --------------------------------------------------------------------------
# Limpeza geral dos dados (botão "Limpar dados do sistema" em Minha Conta)
# --------------------------------------------------------------------------
# Apagados de TODOS os usuários, na ordem em que aparecem (filhos antes dos
# pais, por causa dos on_delete=PROTECT de Operacao.ativo e
# OperacaoCripto.cripto_ativo). Ficam de fora só o catálogo Ações da B3
# (AcaoB3), o contador de consumo da API brapi.dev (ConsumoApiBrapi - controle
# do orçamento mensal do plano, não é cadastro) e, no app accounts, os
# usuários e seus perfis (Minha Conta). Model novo no core precisa entrar
# numa das duas listas - core.tests.LimparDadosSistemaTests confere isso.
MODELOS_LIMPOS_NA_LIMPEZA_GERAL = [
    NotaCorretagem,
    LancamentoContaCorrente,
    VendaLote,
    ContaCorrente,
    OperacaoCripto,
    CotacaoCripto,
    RegistroAtualizacaoCarteiraCripto,
    CriptoAtivo,
    Alerta,
    Operacao,
    EventoCorporativo,  # depois dos lotes (que o protegem) e antes dos ativos
    Provento,
    Cotacao,
    RegistroCotacao,
    PrevisaoAnalista,
    Ativo,
    RegistroAtualizacaoCarteira,
    CotacaoIndice,
    Noticia,
    FonteNoticia,
    MensagemWhatsapp,
    PostIt,
    SaldoInicialFiscal,
    Darf,
    AvisoLimiteIsencao,
    FundamentosAtivo,  # cache de dado público - volta a ser buscado
    TaxaSelicMensal,  # cache do Banco Central - volta a ser buscado
    MapeamentoTitulo,  # nome de pregão -> ticker confirmado nas importações de notas
]
# Ficam de propósito: o catálogo Ações da B3; o consumo da API brapi (zerar faria
# o sistema achar que ainda tem cota no mês e estourar o plano); e os parâmetros
# tributários (alíquotas e limites da lei - sem eles o IR cai no padrão e perde
# regras com vigência, como a retenção de dividendos de 2026).
MODELOS_PRESERVADOS_NA_LIMPEZA_GERAL = [AcaoB3, ConsumoApiBrapi, ParametroTributario]
MANTIDOS_NA_LIMPEZA_GERAL = (
    "Ações da B3, usuários, Minha Conta, parâmetros tributários e o controle de consumo da API de cotações"
)


@transaction.atomic
def limpar_dados_sistema() -> int:
    """
    Apaga todos os dados do sistema, exceto o que está em MANTIDOS_NA_LIMPEZA_GERAL
    (e as tabelas internas do Django: permissões, grupos, tipos de conteúdo).
    Também apaga o histórico de ações do Admin e as sessões de login já
    expiradas - as ativas ficam, senão quem está limpando cairia no meio da
    operação. Tudo numa transação só: se algo falhar, nada é apagado. Devolve
    o total de registros apagados.
    """
    from django.contrib.admin.models import LogEntry
    from django.contrib.sessions.models import Session

    total = 0
    for modelo in MODELOS_LIMPOS_NA_LIMPEZA_GERAL:
        apagados, _ = modelo.objects.all().delete()
        total += apagados
    total += LogEntry.objects.all().delete()[0]
    total += Session.objects.filter(expire_date__lt=timezone.now()).delete()[0]
    return total


# Limpeza de UMA conta (combo de usuário no mesmo botão): modelo -> caminho
# até o usuário dono do registro, na ordem de exclusão (filhos antes dos
# pais). Os demais modelos de MODELOS_LIMPOS_NA_LIMPEZA_GERAL (ativos,
# cotações, índices, notícias, WhatsApp...) são compartilhados entre os
# usuários e ficam intactos - ver MODELOS_COMPARTILHADOS_NA_LIMPEZA_USUARIO.
MODELOS_LIMPOS_NA_LIMPEZA_USUARIO = {
    NotaCorretagem: "usuario",
    LancamentoContaCorrente: "conta__usuario",
    VendaLote: "operacao__usuario",
    ContaCorrente: "usuario",
    OperacaoCripto: "usuario",
    RegistroAtualizacaoCarteiraCripto: "usuario",
    Alerta: "usuario",
    Operacao: "usuario",
    EventoCorporativo: "usuario",  # depois dos lotes (que o protegem) e antes dos ativos
    Provento: "usuario",
    RegistroCotacao: "usuario",
    PrevisaoAnalista: "usuario",
    RegistroAtualizacaoCarteira: "usuario",
    PostIt: "usuario",
    SaldoInicialFiscal: "usuario",
    Darf: "usuario",
    AvisoLimiteIsencao: "usuario",
}
MODELOS_COMPARTILHADOS_NA_LIMPEZA_USUARIO = [
    CotacaoCripto,
    CriptoAtivo,
    Cotacao,
    Ativo,
    CotacaoIndice,
    Noticia,
    FonteNoticia,
    MensagemWhatsapp,
    FundamentosAtivo,
    TaxaSelicMensal,
    MapeamentoTitulo,
]


@transaction.atomic
def limpar_dados_usuario(usuario) -> int:
    """
    Apaga os dados de um usuário só (operações, alertas, conta corrente,
    cripto, post-it, previsões, históricos de atualização), mantendo o
    usuário, o perfil (Minha Conta) e os dados compartilhados. Devolve o
    total de registros apagados.
    """
    total = 0
    for modelo, campo_usuario in MODELOS_LIMPOS_NA_LIMPEZA_USUARIO.items():
        apagados, _ = modelo.objects.filter(**{campo_usuario: usuario}).delete()
        total += apagados
    return total


# --------------------------------------------------------------------------
# Previsões dos analistas (preço-alvo de 12 meses) x preço real
#
# Responde "os analistas acertam a tendência de um ano?" com os dados do
# próprio usuário: ele registra o consenso (o que o TradingView mostra na aba
# "Previsão", ou o que a API de cotações devolver) e, vencido o prazo, o
# sistema compara com o preço real.
# --------------------------------------------------------------------------
_RECOMENDACAO_API = {
    "strong_buy": PrevisaoAnalista.COMPRA_FORTE,
    "buy": PrevisaoAnalista.COMPRA,
    "hold": PrevisaoAnalista.NEUTRO,
    "underperform": PrevisaoAnalista.VENDA,
    "sell": PrevisaoAnalista.VENDA,
    "strong_sell": PrevisaoAnalista.VENDA_FORTE,
}


def buscar_consenso_analistas(ticker: str) -> dict:
    """
    Cotação atual e consenso dos analistas (módulo financialData da
    brapi.dev: targetMeanPrice, targetLowPrice, targetHighPrice,
    recommendationKey, numberOfAnalystOpinions) de um ticker - usado para
    pré-preencher o formulário de Previsões dos Analistas.

    O módulo nem sempre está disponível (depende do plano contratado e da
    cobertura do ativo): nesse caso tenta de novo só a cotação, e os campos
    do consenso voltam None para o usuário completar com o que vê no
    TradingView. Levanta BrapiError só quando nem a cotação vem.
    """
    url = f"{settings.BRAPI_BASE_URL}/quote/{ticker.upper()}"
    params = {"modules": "financialData"}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN

    try:
        resposta = _get_brapi(url, params=params, timeout=15)
        resposta.raise_for_status()
        resultados = resposta.json().get("results") or []
        dados = resultados[0] if resultados else {}
    except (requests.RequestException, ValueError):
        dados = {}

    if dados.get("regularMarketPrice") is None:
        dados = buscar_cotacao_atual(ticker)

    financeiro = dados.get("financialData") or {}
    return {
        "ticker": ticker.upper(),
        "preco": _para_decimal(dados.get("regularMarketPrice")),
        "alvo_medio": _para_decimal(financeiro.get("targetMeanPrice")),
        "alvo_minimo": _para_decimal(financeiro.get("targetLowPrice")),
        "alvo_maximo": _para_decimal(financeiro.get("targetHighPrice")),
        "recomendacao": _RECOMENDACAO_API.get((financeiro.get("recommendationKey") or "").lower(), ""),
        "numero_analistas": financeiro.get("numberOfAnalystOpinions"),
    }


def ativos_com_previsao_em_andamento() -> set[int]:
    """IDs dos ativos com previsão de analistas ainda dentro do prazo - entram no ciclo automático de cotações."""
    inicio = timezone.localdate() - timedelta(days=PrevisaoAnalista.PRAZO_DIAS)
    return set(
        PrevisaoAnalista.objects.filter(data_registro__gte=inicio).values_list("ativo_id", flat=True).distinct()
    )


def atualizar_precos_previsoes(usuario) -> dict:
    """
    Cotação de hoje e histórico recente (até 90 dias - limite do plano
    gratuito da brapi.dev, ver _range_brapi) dos ativos com previsão do
    usuário ainda dentro do prazo, para a avaliação não depender de o ativo
    estar em carteira. Ativos de previsões já vencidas não são consultados:
    o preço que conta para elas é o do vencimento.
    """
    inicio = timezone.localdate() - timedelta(days=PrevisaoAnalista.PRAZO_DIAS)
    ativos = list(
        Ativo.objects.filter(
            previsoes_analistas__usuario=usuario, previsoes_analistas__data_registro__gte=inicio,
        ).distinct()
    )
    atualizados, falhas = atualizar_cotacoes_ativos(ativos)
    for ativo in ativos:
        try:
            backfill_historico_cotacoes(ativo, dias=90)
        except BrapiError:
            pass  # a cotação do dia já foi tentada acima - o histórico é complemento
    return {"ativos": len(ativos), "atualizados": atualizados, "falhas": falhas}


def _pct(valor: Decimal, base: Decimal) -> Decimal:
    return ((valor / base - 1) * 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _direcao(preco_final: Decimal, preco_inicial: Decimal) -> str:
    if preco_final > preco_inicial:
        return "ALTA"
    if preco_final < preco_inicial:
        return "BAIXA"
    return "NEUTRA"


def avaliar_previsao_analista(previsao: PrevisaoAnalista, cotacoes: list[Cotacao], hoje: date | None = None) -> dict:
    """
    Compara uma previsão com as cotações diárias do ativo (em qualquer
    ordem) entre o dia seguinte ao registro e o vencimento (12 meses) - ou
    até hoje, enquanto o prazo não acabou (resultado parcial).

    - acertou_direcao: o preço andou para o lado do alvo (None se o alvo era
      igual à cotação do registro ou ainda não há cotação no período);
    - alvo_no_fim: o último preço do período já está no alvo ou além dele;
    - alvo_no_periodo: em algum dia do período a máxima (alta) ou a mínima
      (baixa) chegou ao alvo;
    - erro_alvo_pct: quanto o último preço ficou acima (+) ou abaixo (-) do alvo.
    """
    hoje = hoje or timezone.localdate()
    vencimento = previsao.data_vencimento
    concluida = hoje >= vencimento
    fim = min(vencimento, hoje)
    periodo = sorted(
        (c for c in cotacoes if previsao.data_registro < c.data <= fim), key=lambda c: c.data,
    )
    prevista = previsao.direcao_prevista
    resultado = {
        "previsao": previsao,
        "concluida": concluida,
        "dias_restantes": max((vencimento - hoje).days, 0),
        "direcao_prevista": prevista,
        "data_preco_final": None,
        "preco_final": None,
        "variacao_real_pct": None,
        "direcao_real": None,
        "acertou_direcao": None,
        "alvo_no_fim": None,
        "alvo_no_periodo": None,
        "erro_alvo_pct": None,
        "dados_incompletos": False,
    }
    if not periodo:
        resultado["dados_incompletos"] = True
        return resultado

    ultima = periodo[-1]
    preco_final = ultima.preco_fechamento
    alvo = previsao.alvo_medio
    real = _direcao(preco_final, previsao.preco_registro)
    resultado.update({
        "data_preco_final": ultima.data,
        "preco_final": preco_final,
        "variacao_real_pct": _pct(preco_final, previsao.preco_registro),
        "direcao_real": real,
        "erro_alvo_pct": _pct(preco_final, alvo),
        # vencida, mas a última cotação gravada ficou bem antes do vencimento
        "dados_incompletos": concluida and (vencimento - ultima.data).days > 10,
    })
    if prevista == "ALTA":
        resultado["acertou_direcao"] = real == "ALTA"
        resultado["alvo_no_fim"] = preco_final >= alvo
        resultado["alvo_no_periodo"] = max(c.maxima or c.preco_fechamento for c in periodo) >= alvo
    elif prevista == "BAIXA":
        resultado["acertou_direcao"] = real == "BAIXA"
        resultado["alvo_no_fim"] = preco_final <= alvo
        resultado["alvo_no_periodo"] = min(c.minima or c.preco_fechamento for c in periodo) <= alvo
    return resultado


def _taxa(avaliacoes: list[dict], chave: str) -> dict:
    validas = [a[chave] for a in avaliacoes if a[chave] is not None]
    acertos = sum(1 for v in validas if v)
    return {
        "acertos": acertos,
        "total": len(validas),
        "pct": round(acertos / len(validas) * 100, 1) if validas else None,
    }


def resumir_acertos_analistas(avaliacoes: list[dict]) -> dict:
    """
    Taxas de acerto de um grupo de avaliações (ver avaliar_previsao_analista).
    "base_alta_pct" é a fatia das previsões em que o preço simplesmente
    subiu - a referência de quem "chutasse alta para tudo": se a taxa de
    acerto da direção não passa dela, a previsão não acrescentou informação.
    """
    com_preco = [a for a in avaliacoes if a["preco_final"] is not None]
    erros = [abs(a["erro_alvo_pct"]) for a in com_preco]
    altas = sum(1 for a in com_preco if a["direcao_real"] == "ALTA")
    previstas_alta = sum(1 for a in avaliacoes if a["direcao_prevista"] == "ALTA")
    return {
        "total": len(avaliacoes),
        "com_preco": len(com_preco),
        "direcao": _taxa(com_preco, "acertou_direcao"),
        "alvo_no_fim": _taxa(com_preco, "alvo_no_fim"),
        "alvo_no_periodo": _taxa(com_preco, "alvo_no_periodo"),
        "erro_medio_abs_pct": round(sum(erros) / len(erros), 2) if erros else None,
        "base_alta_pct": round(altas / len(com_preco) * 100, 1) if com_preco else None,
        "previstas_alta_pct": round(previstas_alta / len(avaliacoes) * 100, 1) if avaliacoes else None,
    }


def avaliar_previsoes_analistas(usuario, hoje: date | None = None) -> dict:
    """
    Avalia todas as previsões do usuário e resume a taxa de acerto em dois
    grupos: concluídas (prazo de 12 meses vencido - resultado definitivo) e
    em andamento (resultado parcial, com o preço até hoje).
    """
    hoje = hoje or timezone.localdate()
    previsoes = list(PrevisaoAnalista.objects.filter(usuario=usuario).select_related("ativo"))
    cotacoes_por_ativo = defaultdict(list)
    if previsoes:
        for cotacao in Cotacao.objects.filter(
            ativo_id__in={p.ativo_id for p in previsoes},
            data__gt=min(p.data_registro for p in previsoes),
            data__lte=hoje,
        ):
            cotacoes_por_ativo[cotacao.ativo_id].append(cotacao)

    avaliacoes = [avaliar_previsao_analista(p, cotacoes_por_ativo[p.ativo_id], hoje) for p in previsoes]
    return {
        "avaliacoes": avaliacoes,
        "concluidas": resumir_acertos_analistas([a for a in avaliacoes if a["concluida"]]),
        "em_andamento": resumir_acertos_analistas([a for a in avaliacoes if not a["concluida"]]),
    }


# --------------------------------------------------------------------------
# Qualidade fundamentalista (tela Qualidade das Empresas)
#
# Filtro de qualidade ANTES do gráfico: empresas que dão lucro todo ano, com
# boa rentabilidade, crescendo, sem dívida excessiva e sem preço exagerado
# tendem a decepcionar menos no longo prazo do que sinais técnicos de curto
# prazo (ver o Backtesting do Robô). Usa os módulos que o plano Startup da
# brapi.dev libera (defaultKeyStatistics, balanço, DRE e proventos) - o
# financialData, com ROE/dívida prontos, só existe no plano Pro, então os
# indicadores são calculados a partir dos demonstrativos.
#
# Cada critério vale pontos; critério sem dado (ou que não se aplica, como
# endividamento de banco) sai da conta, e a nota é proporcional ao máximo
# possível. É um filtro de qualidade, não uma previsão de preço.
# --------------------------------------------------------------------------
FUNDAMENTOS_MODULOS = "defaultKeyStatistics,balanceSheetHistory,incomeStatementHistory,summaryProfile"
# balanço muda por trimestre - mais que isso é desperdício de requisição
FUNDAMENTOS_VALIDADE_DIAS = 7
FUNDAMENTOS_MAX_TICKERS_POR_CONSULTA = 10
_SETORES_FINANCEIROS = ("finance", "financial", "bank", "banco", "insurance", "seguro")


def _numero(valor) -> float | None:
    if valor is None or isinstance(valor, bool):
        return None
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def _lucro_do_ano(demonstrativo: dict) -> float | None:
    # bancos às vezes não trazem netIncome, só os campos "para os acionistas"
    for campo in ("netIncome", "netIncomeApplicableToCommonShares", "netIncomeFromContinuingOps"):
        valor = _numero(demonstrativo.get(campo))
        if valor is not None:
            return valor
    return None


def _dividend_yield_12m(dividendos: list, preco: float | None) -> float | None:
    """Soma os proventos por ação com data-ex nos últimos 12 meses (sem repetir o mesmo evento) / preço."""
    if not preco or preco <= 0:
        return None
    limite = timezone.now() - timedelta(days=365)
    eventos = set()
    for provento in dividendos or []:
        valor = _numero(provento.get("rate"))
        data_ex = provento.get("exDate")
        if not valor or not data_ex:
            continue
        try:
            data = datetime.fromisoformat(str(data_ex).replace("Z", "+00:00"))
        except ValueError:
            continue
        if data.tzinfo is None:
            data = data.replace(tzinfo=dt_timezone.utc)
        if data >= limite:
            # a mesma distribuição aparece uma vez por classe de ação (ON/PN)
            eventos.add((data.date(), round(valor, 6)))
    return sum(valor for _, valor in eventos) / preco * 100


def calcular_qualidade_fundamentos(dados: dict, setor_cadastrado: str = "") -> dict:
    """
    Núcleo puro (sem banco nem rede): recebe o "results[0]" da brapi.dev com
    os módulos de FUNDAMENTOS_MODULOS (+ dividendsData) e devolve os
    indicadores, a lista de critérios avaliados e a nota de 0 a 10.
    """
    estatisticas = dados.get("defaultKeyStatistics") or {}
    perfil = dados.get("summaryProfile") or {}
    preco = _numero(dados.get("regularMarketPrice"))

    # o setor da própria brapi tem prioridade: o cadastrado no Ativo vem de
    # outra classificação, que põe até construtoras/imobiliárias em "Finance"
    if perfil.get("sector"):
        setor = perfil["sector"]
        texto_setor = f"{setor} {perfil.get('industry') or ''}".lower()
    else:
        setor = setor_cadastrado or ""
        texto_setor = setor.lower()
    financeira = any(chave in texto_setor for chave in _SETORES_FINANCEIROS)

    p_l = _numero(dados.get("priceEarnings")) or _numero(estatisticas.get("trailingPE"))
    p_vp = _numero(estatisticas.get("priceToBook"))
    lpa = _numero(estatisticas.get("trailingEps"))
    vpa = _numero(estatisticas.get("bookValue"))
    # ROE pelo lucro e patrimônio POR AÇÃO: funciona igual para bancos, cujo
    # balanço/DRE vêm em outro formato
    roe = lpa / vpa * 100 if lpa is not None and vpa and vpa > 0 else None
    patrimonio_negativo = vpa is not None and vpa <= 0
    margem = _numero(estatisticas.get("profitMargins"))
    margem = margem * 100 if margem is not None else None

    anuais = sorted(
        (d for d in dados.get("incomeStatementHistory") or [] if d.get("type", "yearly") == "yearly"),
        key=lambda d: d.get("endDate") or "", reverse=True,
    )
    lucros = [lucro for lucro in (_lucro_do_ano(d) for d in anuais[:5]) if lucro is not None]
    anos_lucro_positivo = sum(1 for lucro in lucros if lucro > 0) if lucros else None
    lucros_4 = [_lucro_do_ano(d) for d in anuais[:4]]
    cagr = None
    if len(lucros_4) == 4 and None not in lucros_4 and lucros_4[0] > 0 and lucros_4[3] > 0:
        cagr = ((lucros_4[0] / lucros_4[3]) ** (1 / 3) - 1) * 100

    balancos = sorted(
        (b for b in dados.get("balanceSheetHistory") or [] if b.get("type", "yearly") == "yearly"),
        key=lambda b: b.get("endDate") or "", reverse=True,
    )
    endividamento = None
    if balancos:
        ativo_total = _numero(balancos[0].get("totalAssets"))
        passivo_total = _numero(balancos[0].get("totalLiab"))
        if ativo_total and passivo_total is not None:
            endividamento = passivo_total / ativo_total * 100

    dividend_yield = _dividend_yield_12m((dados.get("dividendsData") or {}).get("cashDividends"), preco)

    criterios = []

    def avaliar(nome, maximo, pontos, valor_texto, regra, aplicavel=True):
        criterios.append({
            "nome": nome, "maximo": maximo, "pontos": pontos if aplicavel else None,
            "valor": valor_texto, "regra": regra, "aplicavel": aplicavel,
        })

    if anos_lucro_positivo is not None and len(lucros) >= 3:
        pontos = 2 if anos_lucro_positivo == len(lucros) else 1 if anos_lucro_positivo >= len(lucros) - 1 else 0
        avaliar("Lucro consistente", 2, pontos, f"lucro em {anos_lucro_positivo} de {len(lucros)} anos",
                "2 pts: lucro todo ano; 1 pt: no máximo 1 ano de prejuízo")
    else:
        avaliar("Lucro consistente", 2, 0, "sem DRE anual suficiente", "", aplicavel=False)

    if patrimonio_negativo:
        avaliar("Rentabilidade (ROE)", 2, 0, "patrimônio líquido negativo", "2 pts: ROE ≥ 15%; 1 pt: ≥ 10%")
    elif roe is not None:
        avaliar("Rentabilidade (ROE)", 2, 2 if roe >= 15 else 1 if roe >= 10 else 0, f"{roe:.1f}%",
                "2 pts: ROE ≥ 15%; 1 pt: ≥ 10%")
    else:
        avaliar("Rentabilidade (ROE)", 2, 0, "sem dado", "", aplicavel=False)

    if margem is not None:
        avaliar("Margem líquida", 1, 1 if margem >= 10 else 0, f"{margem:.1f}%", "1 pt: margem ≥ 10%")
    else:
        avaliar("Margem líquida", 1, 0, "sem dado", "", aplicavel=False)

    if cagr is not None:
        avaliar("Crescimento do lucro", 2, 2 if cagr >= 10 else 1 if cagr > 0 else 0, f"{cagr:+.1f}% ao ano (3 anos)",
                "2 pts: ≥ 10% ao ano; 1 pt: crescendo")
    elif len(lucros_4) == 4 and None not in lucros_4:
        avaliar("Crescimento do lucro", 2, 0, "prejuízo no início ou no fim do período",
                "2 pts: ≥ 10% ao ano; 1 pt: crescendo")
    else:
        avaliar("Crescimento do lucro", 2, 0, "sem 4 anos de DRE", "", aplicavel=False)

    if financeira:
        avaliar("Endividamento", 1, 0, "não se aplica a banco/seguradora", "", aplicavel=False)
    elif endividamento is not None:
        avaliar("Endividamento", 1, 1 if endividamento <= 60 else 0, f"passivo = {endividamento:.0f}% do ativo",
                "1 pt: passivo ≤ 60% do ativo total")
    else:
        avaliar("Endividamento", 1, 0, "sem balanço", "", aplicavel=False)

    if p_l is not None:
        barato = 0 < p_l <= 20
        avaliar("Preço (P/L)", 1, 1 if barato else 0,
                f"P/L {p_l:.1f}" + (" (prejuízo)" if p_l <= 0 else ""), "1 pt: P/L entre 0 e 20")
    else:
        avaliar("Preço (P/L)", 1, 0, "sem dado (empresa com prejuízo?)", "1 pt: P/L entre 0 e 20")

    if dividend_yield is not None:
        avaliar("Dividendos", 1, 1 if dividend_yield >= 4 else 0, f"{dividend_yield:.1f}% em 12 meses",
                "1 pt: dividend yield ≥ 4%")
    else:
        avaliar("Dividendos", 1, 0, "sem dado", "", aplicavel=False)

    maximo = sum(c["maximo"] for c in criterios if c["aplicavel"])
    pontos = sum(c["pontos"] for c in criterios if c["aplicavel"])
    # com menos da metade dos pontos avaliáveis, a nota diria pouco
    nota = round(pontos / maximo * 10, 1) if maximo >= 5 else None

    return {
        "nome": dados.get("longName") or dados.get("shortName") or "",
        "setor": setor,
        "financeira": financeira,
        "preco": preco,
        "p_l": p_l,
        "p_vp": p_vp,
        "roe_pct": roe,
        "margem_liquida_pct": margem,
        "cagr_lucro_pct": cagr,
        "anos_lucro_positivo": anos_lucro_positivo,
        "anos_avaliados": len(lucros) or None,
        "endividamento_pct": endividamento,
        "dividend_yield_pct": dividend_yield,
        "nota": nota,
        "criterios": criterios,
    }


def buscar_fundamentos(ticker: str) -> dict:
    """Uma requisição à brapi.dev com balanço, DRE, estatísticas, perfil e proventos do ticker."""
    url = f"{settings.BRAPI_BASE_URL}/quote/{ticker.upper()}"
    params = {"modules": FUNDAMENTOS_MODULOS, "fundamental": "true", "dividends": "true"}
    if settings.BRAPI_TOKEN:
        params["token"] = settings.BRAPI_TOKEN
    try:
        resposta = _get_brapi(url, params=params, timeout=30)
    except requests.RequestException as exc:
        raise BrapiError(f"Falha ao consultar os fundamentos de {ticker}. Tente novamente em instantes.") from exc
    if resposta.status_code == 404:
        raise BrapiError(f"{ticker} não encontrado na B3.")
    try:
        corpo = resposta.json()
    except ValueError:
        corpo = {}
    if resposta.status_code != 200:
        mensagem = corpo.get("message") if isinstance(corpo, dict) else None
        raise BrapiError(mensagem or f"Falha ao consultar os fundamentos de {ticker} (HTTP {resposta.status_code}).")
    resultados = corpo.get("results") or []
    if not resultados:
        raise BrapiError(f"{ticker} não encontrado na B3.")
    return resultados[0]


def atualizar_fundamentos(ticker: str, forcar: bool = False) -> FundamentosAtivo:
    """
    Fundamentos e nota de `ticker`, buscando na API só se não houver dado
    salvo há menos de FUNDAMENTOS_VALIDADE_DIAS (ou com `forcar`). Levanta
    BrapiError se precisar buscar e a API falhar.
    """
    ticker = ticker.strip().upper()
    existente = FundamentosAtivo.objects.filter(ticker=ticker).first()
    validade = timezone.now() - timedelta(days=FUNDAMENTOS_VALIDADE_DIAS)
    if existente and not forcar and existente.atualizado_em >= validade:
        return existente

    ativo = Ativo.objects.filter(ticker=ticker).first()
    calculado = calcular_qualidade_fundamentos(buscar_fundamentos(ticker), ativo.setor if ativo else "")
    criterios = calculado.pop("criterios")
    valores = {
        campo: (_para_decimal(valor) if isinstance(valor, float) else valor)
        for campo, valor in calculado.items()
    }
    valores["nome"] = (valores["nome"] or "")[:255]
    valores["setor"] = (valores["setor"] or "")[:150]
    valores["criterios"] = criterios
    fundamentos, _ = FundamentosAtivo.objects.update_or_create(ticker=ticker, defaults=valores)
    return fundamentos


def tickers_da_carteira(usuario) -> list[str]:
    """Tickers que o usuário tem comprados ou reservados agora."""
    return sorted({p.ativo.ticker for p in calcular_posicoes(usuario) if p.quantidade > 0 or p.quantidade_reservada > 0})
