"""
Relatório anual para a declaração de IR (DIRPF): o que vai em cada ficha,
calculado a partir das operações, proventos, eventos e DARFs do ano.

O relatório é uma lista de seções (título, colunas, linhas, nota) - a mesma
estrutura alimenta a tela, o Excel (uma aba por seção) e o PDF.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from django.utils import timezone

from core.classificacao import ACAO, BDR, ETF, FII, UNIT
from core.models import Ativo

from .models import Darf, Provento
from .services import (
    GRUPO_COMUM, GRUPO_DAY_TRADE, GRUPO_FII, GRUPOS, GRUPO_LABELS, apurar_meses, apurar_operacoes,
    prejuizos_iniciais_do_usuario,
)

ZERO = Decimal("0")

AVISO = (
    "Valores calculados pelo BolsaTrader a partir das operações cadastradas - confira com os informes de "
    "rendimentos e as notas de corretagem antes de declarar. Não substitui a orientação de um contador."
)

# códigos da ficha Bens e Direitos conferidos; ETF e BDR mudam de código entre versões do programa
CODIGO_BENS = {ACAO: "03 - 01", UNIT: "03 - 01", FII: "07 - 03"}
NOME_CLASSE = {ACAO: "ações", UNIT: "units", FII: "cotas do FII", ETF: "cotas do ETF", BDR: "BDRs"}


@dataclass
class Secao:
    titulo: str
    colunas: list[str]
    linhas: list[list] = field(default_factory=list)
    nota: str = ""


def _reais(valor) -> str:
    if valor is None or valor == "":
        return ""
    numero = Decimal(valor).quantize(Decimal("0.01"))
    texto = f"{abs(numero):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-R$ " if numero < 0 else "R$ ") + texto


def _secao_bens_e_direitos(usuario, ano: int) -> Secao:
    anterior = apurar_operacoes(usuario, ate=date(ano - 1, 12, 31)).posicoes
    atual = apurar_operacoes(usuario, ate=date(ano, 12, 31)).posicoes
    nomes = dict(Ativo.objects.filter(ticker__in=set(anterior) | set(atual)).values_list("ticker", "nome"))

    secao = Secao(
        titulo=f"Bens e Direitos - situação em 31/12/{ano - 1} e 31/12/{ano}",
        colunas=["Grupo - código", "Ativo", "Discriminação sugerida", f"Qtd. 31/12/{ano - 1}",
                 f"Situação 31/12/{ano - 1}", f"Qtd. 31/12/{ano}", f"Situação 31/12/{ano}"],
        nota="Valor = custo de aquisição pelo preço médio (não o valor de mercado). Para ETF e BDR, confira o "
        "grupo e o código no programa da Receita. Inclua o CNPJ e a corretora na discriminação.",
    )
    for ticker in sorted(set(anterior) | set(atual)):
        antes, depois = anterior.get(ticker), atual.get(ticker)
        if not (antes and antes.quantidade) and not (depois and depois.quantidade):
            continue
        classe = (depois or antes).classe
        quantidade = depois.quantidade if depois else 0
        discriminacao = (
            f"{quantidade} {NOME_CLASSE.get(classe, 'unidades')} {ticker}"
            + (f" - {nomes[ticker]}" if nomes.get(ticker) and nomes[ticker].upper() != ticker else "")
            + (f" - custo médio R$ {depois.preco_medio:.2f}".replace(".", ",") if depois and depois.quantidade else "")
        )
        secao.linhas.append([
            CODIGO_BENS.get(classe, "conferir"), ticker, discriminacao,
            antes.quantidade if antes else 0, _reais(antes.custo_total_centavos if antes else ZERO),
            quantidade, _reais(depois.custo_total_centavos if depois else ZERO),
        ])
    return secao


def _proventos_do_ano(usuario, ano: int, tipo: str):
    return Provento.objects.filter(usuario=usuario, tipo=tipo, data_pagamento__year=ano).select_related("ativo")


def _por_ativo(proventos) -> list[list]:
    totais: dict[str, list[Decimal]] = {}
    for provento in proventos:
        bruto, ir = totais.setdefault(provento.ativo.ticker, [ZERO, ZERO])
        totais[provento.ativo.ticker] = [bruto + provento.valor_bruto, ir + provento.ir_retido]
    return [[ticker, bruto, ir, bruto - ir] for ticker, (bruto, ir) in sorted(totais.items())]


def _secao_isentos(usuario, ano: int, meses) -> Secao:
    secao = Secao(
        titulo="Rendimentos isentos e não tributáveis",
        colunas=["Rendimento", "Ativo", "Valor bruto", "IR retido", "Valor a declarar"],
        nota="Dividendos e rendimentos de FII: informe o CNPJ de cada fonte pagadora. Dividendos com IR retido "
        "(acima de R$ 50 mil/mês da mesma empresa, desde 2026): confira no informe como a retenção deve ser declarada.",
    )
    ganho_isento = sum((m.grupos[GRUPO_COMUM].isento for m in meses if m.ano == ano), ZERO)
    if ganho_isento:
        secao.linhas.append([
            "Ganhos líquidos com ações - vendas até R$ 20 mil no mês", "", "", "", _reais(ganho_isento),
        ])
    for tipo, rotulo in ((Provento.DIVIDENDO, "Lucros e dividendos recebidos"),
                         (Provento.RENDIMENTO_FII, "Rendimentos de fundos imobiliários")):
        for ticker, bruto, ir, liquido in _por_ativo(_proventos_do_ano(usuario, ano, tipo)):
            secao.linhas.append([rotulo, ticker, _reais(bruto), _reais(ir), _reais(bruto if tipo == Provento.DIVIDENDO else liquido)])
    return secao


def _secao_exclusiva(usuario, ano: int) -> Secao:
    secao = Secao(
        titulo="Rendimentos sujeitos à tributação exclusiva/definitiva",
        colunas=["Rendimento", "Ativo", "Valor bruto", "IR retido", "Valor a declarar (líquido)"],
        nota="JCP é declarado pelo valor líquido (já descontados os 15% retidos), com o CNPJ da fonte pagadora.",
    )
    for ticker, bruto, ir, liquido in _por_ativo(_proventos_do_ano(usuario, ano, Provento.JCP)):
        secao.linhas.append(["Juros sobre capital próprio", ticker, _reais(bruto), _reais(ir), _reais(liquido)])
    return secao


def _secao_renda_variavel(meses, ano: int, darfs_pagas: dict[date, Decimal]) -> Secao:
    secao = Secao(
        titulo=f"Renda Variável - operações comuns e day trade ({ano})",
        colunas=["Mês", "Resultado comum", "Resultado day trade", "Prejuízo compensado (comum)",
                 "Prejuízo compensado (DT)", "Base de cálculo comum", "Base de cálculo DT", "IR devido",
                 "IR retido na fonte", "DARF paga"],
        nota="Ganho de ações isento (vendas até R$ 20 mil) não entra aqui - vai em Rendimentos Isentos.",
    )
    por_mes = {m.mes: m for m in meses if m.ano == ano}
    for numero in range(1, 13):
        m = por_mes.get(numero)
        if m is None:
            continue
        comum, dt = m.grupos[GRUPO_COMUM], m.grupos[GRUPO_DAY_TRADE]
        if not (comum.resultado or dt.resultado):
            continue
        secao.linhas.append([
            f"{numero:02d}/{ano}", _reais(comum.resultado_tributavel), _reais(dt.resultado),
            _reais(comum.prejuizo_compensado), _reais(dt.prejuizo_compensado),
            _reais(comum.base_calculo), _reais(dt.base_calculo), _reais(comum.imposto + dt.imposto),
            _reais(m.irrf_mes), _reais(darfs_pagas.get(date(ano, numero, 1), ZERO)),
        ])
    return secao


def _secao_fii(meses, ano: int) -> Secao:
    secao = Secao(
        titulo=f"Renda Variável - fundos imobiliários ({ano})",
        colunas=["Mês", "Resultado", "Prejuízo compensado", "Base de cálculo", "IR devido (20%)"],
    )
    for m in meses:
        fii = m.grupos[GRUPO_FII]
        if m.ano == ano and fii.resultado:
            secao.linhas.append([
                m.competencia, _reais(fii.resultado), _reais(fii.prejuizo_compensado),
                _reais(fii.base_calculo), _reais(fii.imposto),
            ])
    return secao


def _secao_prejuizos(meses, ano: int, iniciais: dict[str, Decimal]) -> Secao:
    do_ano = [m for m in meses if m.ano <= ano]
    finais = {g: (do_ano[-1].grupos[g].prejuizo_a_compensar if do_ano else iniciais.get(g, ZERO)) for g in GRUPOS}
    return Secao(
        titulo=f"Prejuízos a compensar em 31/12/{ano}",
        colunas=["Tipo de operação", "Prejuízo a compensar"],
        linhas=[[GRUPO_LABELS[g], _reais(finais[g])] for g in GRUPOS],
        nota="Esses saldos passam para o ano seguinte - informe-os em dezembro da ficha Renda Variável.",
    )


def _secao_darfs(darfs) -> Secao:
    return Secao(
        titulo="DARFs pagas no ano (código 6015)",
        colunas=["Mês de apuração", "Vencimento", "Data do pagamento", "Principal", "Multa", "Juros", "Total pago"],
        linhas=[
            [f"{d.competencia:%m/%Y}", f"{d.vencimento:%d/%m/%Y}", f"{d.data_pagamento:%d/%m/%Y}",
             _reais(d.valor_principal), _reais(d.multa_paga), _reais(d.juros_pagos), _reais(d.valor_pago)]
            for d in darfs
        ],
    )


def gerar_relatorio_anual(usuario, ano: int) -> list[Secao]:
    """Todas as seções do relatório anual, na ordem das fichas da declaração."""
    meses = apurar_meses(usuario)
    darfs = list(Darf.objects.filter(usuario=usuario, data_pagamento__year=ano).order_by("competencia"))
    pagas_por_competencia: dict[date, Decimal] = {}
    for darf in Darf.objects.filter(usuario=usuario, competencia__year=ano, data_pagamento__isnull=False):
        pagas_por_competencia[darf.competencia] = darf.valor_pago or ZERO
    return [
        _secao_bens_e_direitos(usuario, ano),
        _secao_isentos(usuario, ano, meses),
        _secao_exclusiva(usuario, ano),
        _secao_renda_variavel(meses, ano, pagas_por_competencia),
        _secao_fii(meses, ano),
        _secao_prejuizos(meses, ano, prejuizos_iniciais_do_usuario(usuario)),
        _secao_darfs(darfs),
    ]


def _celula(valor):
    return valor if isinstance(valor, (int, str)) else _reais(valor)


def gerar_excel_relatorio_anual(secoes: list[Secao], ano: int, nome_usuario: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for indice, secao in enumerate(secoes, start=1):
        ws = wb.create_sheet(f"{indice}. {secao.titulo}"[:31].replace("/", "-"))
        ws.append([f"Declaração IR {ano + 1} (ano-calendário {ano}) - {secao.titulo}"])
        ws.append([f"{nome_usuario} - gerado em {timezone.localtime():%d/%m/%Y %H:%M}"])
        ws.append([])
        ws.append(secao.colunas)
        for celula in ws[4]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="0D1526")
            celula.alignment = Alignment(horizontal="center", wrap_text=True)
        for linha in secao.linhas:
            ws.append([_celula(valor) for valor in linha])
        if not secao.linhas:
            ws.append(["Nada a declarar nesta ficha."])
        ws.append([])
        if secao.nota:
            ws.append([secao.nota])
        ws.append([AVISO])
        for coluna in range(1, len(secao.colunas) + 1):
            ws.column_dimensions[get_column_letter(coluna)].width = 60 if secao.colunas[coluna - 1].startswith("Discrimina") else 22
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_relatorio_anual(secoes: list[Secao], ano: int, nome_usuario: str) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.2 * cm, rightMargin=1.2 * cm,
    )
    estilos = getSampleStyleSheet()
    celula = estilos["BodyText"].clone("celula", fontSize=7.5, leading=9)
    elementos = [
        Paragraph(f"BolsaTrader - Relatório para a Declaração de IR {ano + 1} (ano-calendário {ano})", estilos["Title"]),
        Paragraph(f"{escape(nome_usuario)} - gerado em {timezone.localtime():%d/%m/%Y %H:%M}", estilos["Normal"]),
        Paragraph(escape(AVISO), estilos["Italic"]),
        Spacer(1, 0.4 * cm),
    ]
    for secao in secoes:
        elementos.append(Paragraph(escape(secao.titulo), estilos["Heading2"]))
        if secao.linhas:
            dados = [[Paragraph(f"<b>{escape(c)}</b>", celula) for c in secao.colunas]]
            dados += [[Paragraph(escape(str(_celula(v))), celula) for v in linha] for linha in secao.linhas]
            tabela = Table(dados, repeatRows=1)
            tabela.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E1F2")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            elementos.append(tabela)
        else:
            elementos.append(Paragraph("Nada a declarar nesta ficha.", estilos["Normal"]))
        if secao.nota:
            elementos.append(Paragraph(f"<i>{escape(secao.nota)}</i>", celula))
        elementos.append(Spacer(1, 0.4 * cm))
    doc.build(elementos)
    return buffer.getvalue()
