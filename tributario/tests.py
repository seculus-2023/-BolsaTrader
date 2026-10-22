"""
Testes do módulo tributário (preço médio fiscal, day trade, classe dos ativos e apuração mensal).

Executar com:
    python manage.py test tributario
"""
import importlib
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import Mock, patch

import requests

from django.apps import apps as django_apps
from django.contrib.auth.models import User
from django.contrib.messages import get_messages
from django.test import TestCase
from django.urls import reverse

from core.classificacao import ACAO, BDR, ETF, FII, UNIT, inferir_classe_ativo
from core.models import Ativo, LancamentoContaCorrente, Operacao, VendaLote
from core.services import obter_ou_criar_conta_corrente, recalcular_resumo_vendas, sincronizar_lancamento_compra

from accounts.models import PerfilUsuario

from .calendario import feriados_bancarios, pascoa, ultimo_dia_util
from .darf import (
    calcular_acrescimos, desfazer_pagamento, notificar_darfs_telegram, obter_taxas_selic, registrar_pagamento,
    situacao_das_darfs, sincronizar_darfs,
)
from .eventos import EventoInvalido, aplicar_evento, desfazer_evento
from .models import (
    AvisoLimiteIsencao, Darf, EventoCorporativo, ParametroTributario, Provento, SaldoInicialFiscal, TaxaSelicMensal,
)
from .relatorio import gerar_relatorio_anual
from .services import (
    COMUM, DAY_TRADE, GRUPO_COMUM, GRUPO_DAY_TRADE, GRUPO_FII, MovimentoDia, ResultadoVenda, apurar_meses,
    apurar_operacoes, calcular_apuracao, calcular_meses, ir_estimado_por_lote,
)

D1 = date(2026, 3, 2)
D2 = date(2026, 3, 3)
D3 = date(2026, 3, 4)


class ClassificacaoAtivoTests(TestCase):
    def test_inferencia_pelo_ticker(self):
        self.assertEqual(inferir_classe_ativo("PETR4"), ACAO)
        self.assertEqual(inferir_classe_ativo("VALE3"), ACAO)
        self.assertEqual(inferir_classe_ativo("PETR4F"), ACAO)
        self.assertEqual(inferir_classe_ativo("AAPL34"), BDR)
        self.assertEqual(inferir_classe_ativo("HGLG11"), FII)
        self.assertEqual(inferir_classe_ativo("BOVA11"), ETF)
        self.assertEqual(inferir_classe_ativo("TAEE11"), UNIT)
        self.assertEqual(inferir_classe_ativo("XPTO11", "iShares Qualquer Índice"), ETF)

    def test_ativo_novo_ganha_classe_e_classe_manual_e_respeitada(self):
        self.assertEqual(Ativo.objects.create(ticker="HGLG11").classe, FII)
        corrigido = Ativo.objects.create(ticker="ABCD11", classe=UNIT)
        corrigido.save()
        corrigido.refresh_from_db()
        self.assertEqual(corrigido.classe, UNIT)


class CalculoPrecoMedioPuroTests(TestCase):
    """Regras de preço médio e day trade direto em calcular_apuracao, sem banco."""

    def test_venda_comum_da_baixa_pelo_preco_medio_ponderado_com_taxas(self):
        movimentos = {
            ("PETR4", D1): MovimentoDia(quantidade_compra=100, custo_compra=Decimal("1005.00")),  # 100 x 10 + 5 de taxas
            ("PETR4", D2): MovimentoDia(quantidade_compra=100, custo_compra=Decimal("2000.00")),
            ("PETR4", D3): MovimentoDia(quantidade_venda=50, valor_venda=Decimal("1500.00"), custos_venda=Decimal("3.00")),
        }
        apuracao = calcular_apuracao(movimentos, {"PETR4": ACAO})

        posicao = apuracao.posicoes["PETR4"]
        self.assertEqual(posicao.quantidade, 150)
        self.assertEqual(posicao.preco_medio, Decimal("15.0250"))  # (1005 + 2000) / 200
        self.assertEqual(posicao.custo_total_centavos, Decimal("2253.75"))

        [resultado] = apuracao.resultados
        self.assertEqual(resultado.modalidade, COMUM)
        self.assertEqual(resultado.resultado, Decimal("745.75"))  # 1500 - 3 - 50 x 15,025
        self.assertEqual(apuracao.inconsistencias, [])

    def test_compra_e_venda_no_mesmo_dia_sao_day_trade_e_nao_mexem_no_preco_medio(self):
        movimentos = {
            ("PETR4", D1): MovimentoDia(quantidade_compra=100, custo_compra=Decimal("1000.00")),
            ("PETR4", D2): MovimentoDia(
                quantidade_compra=50, custo_compra=Decimal("600.00"),
                quantidade_venda=50, valor_venda=Decimal("650.00"),
            ),
        }
        apuracao = calcular_apuracao(movimentos, {})

        [resultado] = apuracao.resultados
        self.assertEqual(resultado.modalidade, DAY_TRADE)
        self.assertEqual(resultado.resultado, Decimal("50.00"))
        self.assertEqual(apuracao.posicoes["PETR4"].quantidade, 100)
        self.assertEqual(apuracao.posicoes["PETR4"].preco_medio, Decimal("10.0000"))

    def test_sobra_do_dia_separa_day_trade_da_operacao_comum(self):
        movimentos = {
            ("PETR4", D1): MovimentoDia(quantidade_compra=100, custo_compra=Decimal("1000.00")),
            ("PETR4", D2): MovimentoDia(
                quantidade_compra=20, custo_compra=Decimal("240.00"),
                quantidade_venda=50, valor_venda=Decimal("650.00"),
            ),
        }
        apuracao = calcular_apuracao(movimentos, {})

        por_modalidade = {r.modalidade: r for r in apuracao.resultados}
        self.assertEqual(por_modalidade[DAY_TRADE].quantidade, 20)
        self.assertEqual(por_modalidade[DAY_TRADE].resultado, Decimal("20.00"))  # 20 x 13 - 240
        self.assertEqual(por_modalidade[COMUM].quantidade, 30)
        self.assertEqual(por_modalidade[COMUM].resultado, Decimal("90.00"))  # 30 x 13 - 30 x 10
        self.assertEqual(apuracao.posicoes["PETR4"].quantidade, 70)

    def test_vender_tudo_zera_o_custo(self):
        movimentos = {
            ("VALE3", D1): MovimentoDia(quantidade_compra=3, custo_compra=Decimal("100.00")),
            ("VALE3", D2): MovimentoDia(quantidade_venda=3, valor_venda=Decimal("120.00")),
        }
        posicao = calcular_apuracao(movimentos, {}).posicoes["VALE3"]
        self.assertEqual(posicao.quantidade, 0)
        self.assertEqual(posicao.custo_total, Decimal("0"))

    def test_venda_maior_que_a_posicao_gera_aviso(self):
        movimentos = {
            ("VALE3", D1): MovimentoDia(quantidade_compra=10, custo_compra=Decimal("100.00")),
            ("VALE3", D2): MovimentoDia(quantidade_venda=15, valor_venda=Decimal("180.00")),
        }
        apuracao = calcular_apuracao(movimentos, {})
        self.assertEqual(len(apuracao.inconsistencias), 1)
        self.assertIn("VALE3", apuracao.inconsistencias[0])
        self.assertEqual(apuracao.posicoes["VALE3"].quantidade, 0)


class ApurarOperacoesTests(TestCase):
    """apurar_operacoes a partir dos lotes e vendas gravados."""

    def setUp(self):
        self.usuario = User.objects.create_user(username="contribuinte", password="SenhaForte123!")
        self.ativo = Ativo.objects.create(ticker="PETR4")

    def _lote(self, quantidade, preco, data, taxas="0", usuario=None):
        return Operacao.objects.create(
            usuario=usuario or self.usuario, ativo=self.ativo, tipo=Operacao.COMPRA,
            quantidade=quantidade, preco_unitario=Decimal(preco), data_operacao=data, taxas_compra=Decimal(taxas),
        )

    def test_vender_o_lote_mais_caro_nao_muda_o_custo_fiscal(self):
        self._lote(100, "10.00", D1)
        lote_caro = self._lote(100, "20.00", D2)
        VendaLote.objects.create(operacao=lote_caro, quantidade=50, preco=Decimal("30.00"), data=D3)

        apuracao = apurar_operacoes(self.usuario)

        # visão do lote: (30 - 20) x 50; visão fiscal: 30 x 50 - 50 x 15 (preço médio)
        self.assertEqual(lote_caro.lucro_perda_realizado, Decimal("500.00"))
        self.assertEqual(apuracao.resultados[0].resultado, Decimal("750.00"))
        self.assertEqual(apuracao.posicoes["PETR4"].preco_medio, Decimal("15.0000"))

    def test_vender_lote_antigo_no_dia_de_outra_compra_e_day_trade(self):
        lote_antigo = self._lote(100, "10.00", D1)
        self._lote(50, "12.00", D3)
        VendaLote.objects.create(operacao=lote_antigo, quantidade=50, preco=Decimal("13.00"), data=D3)

        [resultado] = apurar_operacoes(self.usuario).resultados
        self.assertEqual(resultado.modalidade, DAY_TRADE)
        self.assertEqual(resultado.resultado, Decimal("50.00"))  # 50 x 13 - 50 x 12

    def test_taxas_entram_no_custo_e_corretora_reduz_o_ganho(self):
        lote = self._lote(10, "10.00", D1, taxas="1.00")
        VendaLote.objects.create(
            operacao=lote, quantidade=10, preco=Decimal("12.00"), data=D2,
            percentual_corretora=Decimal("0.50"), taxas=Decimal("0.40"),
        )
        [resultado] = apurar_operacoes(self.usuario).resultados
        self.assertEqual(resultado.resultado, Decimal("18.00"))  # 120 - 0,60 - 0,40 - 101

    def test_lote_com_venda_so_no_resumo_tambem_entra(self):
        Operacao.objects.create(
            usuario=self.usuario, ativo=self.ativo, tipo=Operacao.COMPRA, quantidade=10,
            preco_unitario=Decimal("10.00"), data_operacao=D1,
            quantidade_vendida=10, preco_venda=Decimal("11.00"), data_venda=D2,
        )
        [resultado] = apurar_operacoes(self.usuario).resultados
        self.assertEqual(resultado.resultado, Decimal("10.00"))

    def test_data_limite_ignora_o_que_vem_depois(self):
        lote = self._lote(10, "10.00", D1)
        VendaLote.objects.create(operacao=lote, quantidade=10, preco=Decimal("12.00"), data=D3)
        apuracao = apurar_operacoes(self.usuario, ate=D2)
        self.assertEqual(apuracao.resultados, [])
        self.assertEqual(apuracao.posicoes["PETR4"].quantidade, 10)

    def test_operacoes_de_outro_usuario_nao_entram(self):
        outro = User.objects.create_user(username="outro_contribuinte", password="SenhaForte123!")
        self._lote(10, "10.00", D1, usuario=outro)
        self.assertEqual(apurar_operacoes(self.usuario).posicoes, {})

    def test_reserva_nao_entra(self):
        Operacao.objects.create(
            usuario=self.usuario, ativo=self.ativo, tipo=Operacao.RESERVAR,
            quantidade=1, preco_unitario=Decimal("10.00"), data_operacao=D1,
        )
        self.assertEqual(apurar_operacoes(self.usuario).posicoes, {})


class TaxasCompraContaCorrenteTests(TestCase):
    def test_debito_da_compra_inclui_as_taxas(self):
        usuario = User.objects.create_user(username="pagador_taxas", password="SenhaForte123!")
        lote = Operacao.objects.create(
            usuario=usuario, ativo=Ativo.objects.create(ticker="ITUB4"), tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("30.00"), data_operacao=D1, taxas_compra=Decimal("1.25"),
        )
        sincronizar_lancamento_compra(lote)
        lancamento = obter_ou_criar_conta_corrente(usuario).lancamentos.get(origem=LancamentoContaCorrente.ORIGEM_COMPRA)
        self.assertEqual(lancamento.valor, Decimal("301.25"))


class MigracaoVendasLoteTests(TestCase):
    """A migração cria um VendaLote para cada lote já vendido e liga o lançamento de venda existente a ele."""

    def test_cria_venda_e_liga_lancamento(self):
        usuario = User.objects.create_user(username="dados_antigos", password="SenhaForte123!")
        lote = Operacao.objects.create(
            usuario=usuario, ativo=Ativo.objects.create(ticker="BBAS3"), tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("20.00"), data_operacao=D1,
            quantidade_vendida=6, preco_venda=Decimal("25.00"), data_venda=D2, percentual_corretora=Decimal("0.50"),
        )
        lancamento = LancamentoContaCorrente.objects.create(
            conta=obter_ou_criar_conta_corrente(usuario), tipo=LancamentoContaCorrente.CREDITO,
            origem=LancamentoContaCorrente.ORIGEM_VENDA, valor=Decimal("149.25"),
            descricao="Venda de 6x BBAS3", data=D2, operacao=lote,
        )

        migracao = importlib.import_module("core.migrations.0028_preenche_classe_e_vendas_lote")
        migracao.preencher(django_apps, None)

        venda = lote.vendas.get()
        self.assertEqual((venda.quantidade, venda.preco, venda.data), (6, Decimal("25.00"), D2))
        self.assertEqual(venda.percentual_corretora, Decimal("0.50"))
        lancamento.refresh_from_db()
        self.assertEqual(lancamento.venda_lote, venda)

        migracao.preencher(django_apps, None)  # rodar de novo não duplica
        self.assertEqual(lote.vendas.count(), 1)


class PosicaoFiscalViewTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user(username="tela_fiscal", password="SenhaForte123!")
        self.client.login(username="tela_fiscal", password="SenhaForte123!")

    def test_exige_login(self):
        self.client.logout()
        resposta = self.client.get(reverse("tributario:posicao_fiscal"))
        self.assertEqual(resposta.status_code, 302)

    def test_mostra_posicao_e_resultado_do_ano(self):
        ativo = Ativo.objects.create(ticker="WEGE3")
        lote = Operacao.objects.create(
            usuario=self.usuario, ativo=ativo, tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("40.00"), data_operacao=D1,
        )
        VendaLote.objects.create(operacao=lote, quantidade=4, preco=Decimal("50.00"), data=D1)

        resposta = self.client.get(reverse("tributario:posicao_fiscal"), {"ano": "2026"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "WEGE3")
        self.assertContains(resposta, "Day trade")
        self.assertContains(resposta, "R$ 240,00")  # custo dos 6 que sobraram

    def test_sem_operacoes_nao_quebra(self):
        resposta = self.client.get(reverse("tributario:posicao_fiscal"), {"ano": "abc"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Nenhum ativo em carteira")

    def test_menu_tem_link(self):
        self.assertContains(self.client.get(reverse("core:menu")), reverse("tributario:imposto_renda"))


def _venda(data, valor_venda, ganho, classe=ACAO, modalidade=COMUM, quantidade=100):
    """Resultado de venda pronto para calcular_meses: ganho = valor de venda - custo de aquisição."""
    return ResultadoVenda(
        data=data, ticker="TEST3", classe=classe, modalidade=modalidade, quantidade=quantidade,
        valor_venda=Decimal(valor_venda), custos_venda=Decimal("0"),
        custo_aquisicao=Decimal(valor_venda) - Decimal(ganho),
    )


MAR = date(2026, 3, 10)
ABR = date(2026, 4, 10)


class ApuracaoMensalTests(TestCase):
    """Regras da apuração mensal direto em calcular_meses (parâmetro padrão cadastrado pela migração)."""

    def test_vendas_ate_20_mil_deixam_o_ganho_com_acoes_isento(self):
        [mes] = calcular_meses([_venda(MAR, "15000", "3000")])
        self.assertTrue(mes.isento)
        self.assertEqual(mes.grupos[GRUPO_COMUM].isento, Decimal("3000"))
        self.assertEqual(mes.imposto_bruto, Decimal("0"))
        self.assertEqual(mes.disponivel_isencao, Decimal("5000.00"))
        self.assertEqual(mes.irrf_mes, Decimal("0"))  # 0,75 de dedo-duro: abaixo de R$ 1, não é retido

    def test_acima_do_limite_tributa_o_ganho_inteiro_e_abate_o_irrf(self):
        [mes] = calcular_meses([_venda(MAR, "25000", "5000")])
        self.assertFalse(mes.isento)
        self.assertEqual(mes.grupos[GRUPO_COMUM].imposto, Decimal("750.00"))  # 15% de 5.000, não só do excedente
        self.assertEqual(mes.irrf_mes, Decimal("1.25"))  # 0,005% de 25.000
        self.assertEqual(mes.valor_darf, Decimal("748.75"))
        self.assertEqual(mes.vencimento, date(2026, 4, 30))

    def test_vendas_em_day_trade_contam_para_o_limite(self):
        [mes] = calcular_meses([
            _venda(MAR, "15000", "2000"),
            _venda(MAR, "6000", "500", modalidade=DAY_TRADE),
        ])
        self.assertFalse(mes.isento)
        self.assertEqual(mes.grupos[GRUPO_COMUM].imposto, Decimal("300.00"))
        self.assertEqual(mes.grupos[GRUPO_DAY_TRADE].imposto, Decimal("100.00"))  # 20%
        self.assertEqual(mes.irrf_mes, Decimal("5.00"))  # 1% do ganho em day trade (o comum, 0,75, não é retido)
        self.assertEqual(mes.valor_darf, Decimal("395.00"))

    def test_prejuizo_de_um_mes_compensa_o_ganho_dos_seguintes(self):
        marco, abril = calcular_meses([_venda(MAR, "30000", "-1000"), _venda(ABR, "30000", "3000")])
        self.assertEqual(marco.grupos[GRUPO_COMUM].prejuizo_a_compensar, Decimal("1000"))
        self.assertEqual(abril.grupos[GRUPO_COMUM].prejuizo_compensado, Decimal("1000"))
        self.assertEqual(abril.grupos[GRUPO_COMUM].imposto, Decimal("300.00"))
        # IR retido em março (sem imposto) abate abril
        self.assertEqual(marco.irrf_saldo_final, Decimal("1.50"))
        self.assertEqual(abril.valor_darf, Decimal("297.00"))

    def test_prejuizo_em_mes_isento_continua_compensavel(self):
        marco, abril = calcular_meses([_venda(MAR, "10000", "-500"), _venda(ABR, "30000", "1000")])
        self.assertTrue(marco.isento)
        self.assertEqual(marco.grupos[GRUPO_COMUM].prejuizo_a_compensar, Decimal("500"))
        self.assertEqual(abril.grupos[GRUPO_COMUM].imposto, Decimal("75.00"))

    def test_prejuizo_em_day_trade_nao_abate_operacao_comum(self):
        [mes] = calcular_meses([
            _venda(MAR, "30000", "2000"),
            _venda(MAR, "5000", "-1000", modalidade=DAY_TRADE),
        ])
        self.assertEqual(mes.grupos[GRUPO_COMUM].imposto, Decimal("300.00"))
        self.assertEqual(mes.grupos[GRUPO_DAY_TRADE].prejuizo_a_compensar, Decimal("1000"))

    def test_fii_e_etf_nao_tem_isencao(self):
        [mes] = calcular_meses([
            _venda(MAR, "5000", "1000", classe=FII),
            _venda(MAR, "3000", "200", classe=ETF),
        ])
        self.assertEqual(mes.vendas_acoes, Decimal("0.00"))
        self.assertEqual(mes.grupos[GRUPO_FII].imposto, Decimal("200.00"))  # 20%
        self.assertEqual(mes.grupos[GRUPO_COMUM].imposto, Decimal("30.00"))  # ETF: 15%, mesmo com mês "isento"

    def test_imposto_abaixo_de_10_reais_acumula_para_o_mes_seguinte(self):
        marco, abril = calcular_meses([_venda(MAR, "25000", "40"), _venda(ABR, "25000", "100")])
        self.assertEqual(marco.valor_darf, Decimal("0"))
        self.assertEqual(marco.imposto_acumulado_proximo, Decimal("4.75"))  # 6,00 - 1,25 de IRRF
        self.assertEqual(abril.valor_darf, Decimal("18.50"))  # 15,00 - 1,25 + 4,75

    def test_saldo_de_irrf_nao_passa_para_outro_ano(self):
        dezembro, janeiro = calcular_meses([
            _venda(date(2025, 12, 10), "30000", "-100"),
            _venda(date(2026, 1, 10), "30000", "1000"),
        ])
        self.assertEqual(dezembro.irrf_saldo_final, Decimal("1.50"))
        self.assertEqual(janeiro.irrf_saldo_anterior, Decimal("0"))
        self.assertEqual(janeiro.vencimento, date(2026, 2, 27))  # 28/02/2026 é sábado

    def test_parametro_novo_vale_a_partir_da_vigencia(self):
        parametros = [
            ParametroTributario(vigencia_inicio=date(2005, 1, 1), vigencia_fim=date(2026, 3, 31)),
            ParametroTributario(vigencia_inicio=date(2026, 4, 1), aliquota_comum_pct=Decimal("17.5")),
        ]
        marco, abril = calcular_meses([_venda(MAR, "30000", "1000"), _venda(ABR, "30000", "1000")], parametros=parametros)
        self.assertEqual(marco.grupos[GRUPO_COMUM].imposto, Decimal("150.00"))
        self.assertEqual(abril.grupos[GRUPO_COMUM].imposto, Decimal("175.00"))

    def test_parametros_cadastrados_pelas_migracoes(self):
        ate_2025, desde_2026 = ParametroTributario.objects.order_by("vigencia_inicio")
        self.assertEqual(ate_2025.vigencia_fim, date(2025, 12, 31))
        self.assertIsNone(ate_2025.limite_dividendos_mes)
        self.assertEqual(desde_2026.vigencia_inicio, date(2026, 1, 1))
        self.assertIsNone(desde_2026.vigencia_fim)
        self.assertEqual(desde_2026.limite_dividendos_mes, Decimal("50000.00"))
        self.assertEqual(desde_2026.limite_isencao_acoes, Decimal("20000.00"))


class SemBancoCentralMixin:
    """O painel calcula juros de DARF vencida - nos testes, a Selic nunca vem da internet."""

    def setUp(self):
        super().setUp()
        patcher = patch("tributario.darf.buscar_selic_mensal", return_value={})
        self.buscar_selic = patcher.start()
        self.addCleanup(patcher.stop)


class ApuracaoMensalComBancoTests(SemBancoCentralMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="apuracao_mensal", password="SenhaForte123!")
        self.client.login(username="apuracao_mensal", password="SenhaForte123!")
        self.ativo = Ativo.objects.create(ticker="PETR4")

    def _lote_vendido(self, quantidade, preco_compra, preco_venda, compra=D1, venda=D2):
        lote = Operacao.objects.create(
            usuario=self.usuario, ativo=self.ativo, tipo=Operacao.COMPRA,
            quantidade=quantidade, preco_unitario=Decimal(preco_compra), data_operacao=compra,
        )
        VendaLote.objects.create(operacao=lote, quantidade=quantidade, preco=Decimal(preco_venda), data=venda)
        recalcular_resumo_vendas(lote)
        return lote

    def test_prejuizo_anterior_informado_entra_na_compensacao(self):
        SaldoInicialFiscal.objects.create(usuario=self.usuario, prejuizo_comum=Decimal("4000"))
        self._lote_vendido(100, "100.00", "300.00")  # vende 30.000, ganha 20.000
        [mes] = apurar_meses(self.usuario)
        self.assertEqual(mes.grupos[GRUPO_COMUM].prejuizo_compensado, Decimal("4000"))
        self.assertEqual(mes.grupos[GRUPO_COMUM].imposto, Decimal("2400.00"))  # 15% de 16.000

    def test_ir_estimado_do_lote_usa_a_aliquota_efetiva_do_mes(self):
        tributado = self._lote_vendido(100, "100.00", "300.00")
        estimativas = ir_estimado_por_lote(self.usuario, [tributado])
        self.assertEqual(estimativas[tributado.id], Decimal("3000.00"))

    def test_ir_estimado_e_zero_em_mes_isento(self):
        isento = self._lote_vendido(10, "100.00", "150.00")  # vende 1.500
        self.assertEqual(ir_estimado_por_lote(self.usuario, [isento])[isento.id], Decimal("0.00"))

    def test_painel_mostra_isencao_e_apuracao(self):
        self._lote_vendido(100, "100.00", "300.00")
        resposta = self.client.get(reverse("tributario:imposto_renda"), {"ano": "2026"})
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "03/2026")
        self.assertContains(resposta, "R$ 2998,50")  # 3.000 - 1,50 de IR retido

    def test_painel_sem_vendas_nao_quebra(self):
        resposta = self.client.get(reverse("tributario:imposto_renda"))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Ainda pode vender")

    def test_salvar_prejuizos_anteriores(self):
        resposta = self.client.post(reverse("tributario:saldo_inicial_salvar"), {
            "prejuizo_comum": "1500.00", "prejuizo_day_trade": "", "prejuizo_fii": "200",
        })
        self.assertRedirects(resposta, reverse("tributario:imposto_renda"))
        saldo = SaldoInicialFiscal.objects.get(usuario=self.usuario)
        self.assertEqual((saldo.prejuizo_comum, saldo.prejuizo_day_trade, saldo.prejuizo_fii),
                         (Decimal("1500.00"), Decimal("0"), Decimal("200.00")))

    def test_rejeita_prejuizo_negativo(self):
        self.client.post(reverse("tributario:saldo_inicial_salvar"), {"prejuizo_comum": "-10"})
        self.assertFalse(SaldoInicialFiscal.objects.filter(usuario=self.usuario, prejuizo_comum=Decimal("-10")).exists())

    def test_minhas_operacoes_mostra_lucro_apos_ir(self):
        self._lote_vendido(100, "100.00", "300.00")
        resposta = self.client.get(reverse("core:operacao_lista"), {"data_de_venda": "", "data_ate_venda": ""})
        self.assertContains(resposta, "Lucro após IR")
        self.assertContains(resposta, "R$ 3000,00 (15,0% do lucro)")
        self.assertContains(resposta, "R$ 17000,00")


class CalendarioTests(TestCase):
    def test_feriados_moveis(self):
        self.assertEqual(pascoa(2026), date(2026, 4, 5))
        feriados = feriados_bancarios(2026)
        for dia in (date(2026, 2, 16), date(2026, 2, 17), date(2026, 4, 3), date(2026, 6, 4), date(2026, 11, 20)):
            self.assertIn(dia, feriados)

    def test_vencimento_pula_fim_de_semana_e_feriado(self):
        self.assertEqual(ultimo_dia_util(2026, 4), date(2026, 4, 30))
        self.assertEqual(ultimo_dia_util(2026, 2), date(2026, 2, 27))  # 28/02 é sábado
        self.assertEqual(ultimo_dia_util(2029, 3), date(2029, 3, 29))  # 30/03/2029 é Sexta-feira Santa


class AcrescimosAtrasoTests(TestCase):
    VENCIMENTO = date(2026, 4, 30)

    def test_pago_ate_o_vencimento_nao_tem_acrescimo(self):
        acrescimos = calcular_acrescimos(Decimal("1000"), self.VENCIMENTO, self.VENCIMENTO, {})
        self.assertEqual((acrescimos.multa, acrescimos.juros), (Decimal("0"), Decimal("0")))

    def test_atraso_no_mes_seguinte_tem_multa_diaria_e_juros_de_1_pct(self):
        acrescimos = calcular_acrescimos(Decimal("1000"), self.VENCIMENTO, date(2026, 5, 10), {})
        self.assertEqual(acrescimos.dias_atraso, 10)
        self.assertEqual(acrescimos.multa, Decimal("33.00"))  # 10 x 0,33%
        self.assertEqual(acrescimos.juros, Decimal("10.00"))  # só o 1% do mês do pagamento
        self.assertTrue(acrescimos.selic_completa)

    def test_multa_limitada_a_20_pct_e_juros_com_selic_dos_meses_intermediarios(self):
        taxas = {date(2026, 5, 1): Decimal("1.1"), date(2026, 6, 1): Decimal("1.0"), date(2026, 7, 1): Decimal("1.2")}
        acrescimos = calcular_acrescimos(Decimal("1000"), self.VENCIMENTO, date(2026, 8, 15), taxas)
        self.assertEqual(acrescimos.multa, Decimal("200.00"))
        self.assertEqual(acrescimos.juros_pct, Decimal("4.3"))  # 1,1 + 1,0 + 1,2 + 1
        self.assertEqual(acrescimos.juros, Decimal("43.00"))

    def test_sem_selic_de_algum_mes_marca_juros_incompletos(self):
        acrescimos = calcular_acrescimos(Decimal("1000"), self.VENCIMENTO, date(2026, 7, 15), {date(2026, 5, 1): Decimal("1")})
        self.assertFalse(acrescimos.selic_completa)


class SelicBancoCentralTests(TestCase):
    @patch("tributario.darf.requests.get")
    def test_busca_grava_no_cache_e_nao_busca_de_novo(self, get):
        get.return_value = Mock(status_code=200, json=lambda: [
            {"data": "01/05/2026", "valor": "1.14"}, {"data": "01/06/2026", "valor": "1.10"},
        ])
        meses = [date(2026, 5, 1), date(2026, 6, 1)]
        self.assertEqual(obter_taxas_selic(meses), {date(2026, 5, 1): Decimal("1.14"), date(2026, 6, 1): Decimal("1.10")})
        self.assertEqual(TaxaSelicMensal.objects.count(), 2)

        obter_taxas_selic(meses)
        self.assertEqual(get.call_count, 1)

    @patch("tributario.darf.requests.get", side_effect=requests.ConnectionError("sem internet"))
    def test_sem_internet_devolve_o_que_tem(self, get):
        TaxaSelicMensal.objects.create(mes=date(2026, 5, 1), taxa_pct=Decimal("1.14"))
        self.assertEqual(obter_taxas_selic([date(2026, 5, 1), date(2026, 6, 1)]), {date(2026, 5, 1): Decimal("1.14")})


class DarfTests(SemBancoCentralMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="contribuinte_darf", password="SenhaForte123!")
        self.client.login(username="contribuinte_darf", password="SenhaForte123!")
        self.lote = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="PETR4"), tipo=Operacao.COMPRA,
            quantidade=100, preco_unitario=Decimal("100.00"), data_operacao=D1,
        )
        # vende 30.000 em março/2026 com 20.000 de ganho: IR 3.000 - 1,50 retido
        self.venda = VendaLote.objects.create(operacao=self.lote, quantidade=100, preco=Decimal("300.00"), data=D2)
        recalcular_resumo_vendas(self.lote)

    def test_cria_a_guia_do_mes_com_imposto(self):
        [darf] = sincronizar_darfs(self.usuario)
        self.assertEqual(darf.competencia, date(2026, 3, 1))
        self.assertEqual(darf.valor_principal, Decimal("2998.50"))
        self.assertEqual(darf.vencimento, date(2026, 4, 30))
        self.assertEqual(darf.codigo_receita, "6015")

    def test_guia_em_aberto_acompanha_correcao_e_some_sem_imposto(self):
        sincronizar_darfs(self.usuario)
        self.venda.preco = Decimal("210.00")  # vende 21.000, ganho 11.000 -> IR 1.650 - 1,05 retido
        self.venda.save()
        [darf] = sincronizar_darfs(self.usuario)
        self.assertEqual(darf.valor_principal, Decimal("1648.95"))

        self.venda.delete()
        recalcular_resumo_vendas(self.lote)
        self.assertEqual(sincronizar_darfs(self.usuario), [])

    def test_guia_paga_fica_congelada_e_mostra_diferenca(self):
        [darf] = sincronizar_darfs(self.usuario)
        registrar_pagamento(darf, date(2026, 4, 28), Decimal("2998.50"))
        self.venda.preco = Decimal("310.00")  # ganho sobe para 21.000
        self.venda.save()

        [situacao] = situacao_das_darfs(self.usuario)
        self.assertEqual(situacao.status, "PAGA")
        self.assertEqual(situacao.darf.valor_principal, Decimal("2998.50"))
        self.assertEqual(situacao.diferenca, Decimal("149.95"))  # 3.150 - 1,55 retido - 2.998,50

    def test_pagamento_lanca_debito_na_conta_corrente_e_desfazer_remove(self):
        [darf] = sincronizar_darfs(self.usuario)
        registrar_pagamento(darf, date(2026, 5, 10), Decimal("3128.44"))
        lancamento = darf.lancamento
        self.assertEqual(lancamento.tipo, LancamentoContaCorrente.DEBITO)
        self.assertEqual(lancamento.origem, LancamentoContaCorrente.ORIGEM_IMPOSTO)
        self.assertEqual((lancamento.valor, lancamento.data), (Decimal("3128.44"), date(2026, 5, 10)))
        self.assertEqual(darf.multa_paga, Decimal("98.95"))  # 10 dias x 0,33% de 2.998,50

        desfazer_pagamento(darf)
        self.assertFalse(LancamentoContaCorrente.objects.filter(origem=LancamentoContaCorrente.ORIGEM_IMPOSTO).exists())
        self.assertIsNone(Darf.objects.get().data_pagamento)

    def test_tela_lista_guia_vencida_com_acrescimos(self):
        with patch("tributario.darf.timezone.localdate", return_value=date(2026, 5, 10)):
            resposta = self.client.get(reverse("tributario:darf_lista"))
        self.assertContains(resposta, "Vencida há 10 dia(s)")
        self.assertContains(resposta, "R$ 98,95")  # multa
        self.assertContains(resposta, "R$ 29,99")  # juros de 1%

    def test_marcar_como_paga_pela_tela(self):
        [darf] = sincronizar_darfs(self.usuario)
        resposta = self.client.post(reverse("tributario:darf_pagar", args=[darf.id]), {
            "data_pagamento": "2026-04-29", "valor_pago": "2998.50",
        })
        self.assertRedirects(resposta, reverse("tributario:darf_lista"))
        darf.refresh_from_db()
        self.assertEqual(darf.data_pagamento, date(2026, 4, 29))
        self.assertContains(self.client.get(reverse("tributario:darf_lista")), "Paga em 29/04/2026")

        self.client.post(reverse("tributario:darf_desfazer_pagamento", args=[darf.id]))
        darf.refresh_from_db()
        self.assertFalse(darf.paga)

    def test_rejeita_data_futura_e_nao_mexe_em_guia_de_outro_usuario(self):
        [darf] = sincronizar_darfs(self.usuario)
        resposta = self.client.post(reverse("tributario:darf_pagar", args=[darf.id]), {
            "data_pagamento": (date.today() + timedelta(days=5)).isoformat(), "valor_pago": "10",
        })
        self.assertEqual(resposta.status_code, 200)
        self.assertFalse(Darf.objects.get().paga)

        outro = User.objects.create_user(username="outro_darf", password="SenhaForte123!")
        self.client.login(username="outro_darf", password="SenhaForte123!")
        self.assertEqual(self.client.get(reverse("tributario:darf_pagar", args=[darf.id])).status_code, 404)

    def test_painel_avisa_guia_vencida(self):
        with patch("tributario.darf.timezone.localdate", return_value=date(2026, 5, 10)):
            resposta = self.client.get(reverse("tributario:imposto_renda"))
        self.assertContains(resposta, "DARF(s) vencida(s)")


@patch("tributario.darf.enviar_notificacao_telegram", return_value=True)
class LembreteDarfTelegramTests(SemBancoCentralMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="lembrete_darf", password="SenhaForte123!")
        PerfilUsuario.objects.create(usuario=self.usuario, telegram_bot_token="123:abc", telegram_chat_id="999")
        lote = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="VALE3"), tipo=Operacao.COMPRA,
            quantidade=100, preco_unitario=Decimal("100.00"), data_operacao=D1,
        )
        VendaLote.objects.create(operacao=lote, quantidade=100, preco=Decimal("300.00"), data=D2)
        recalcular_resumo_vendas(lote)

    def _notificar_em(self, dia):
        with patch("tributario.darf.timezone.localdate", return_value=dia):
            return notificar_darfs_telegram(self.usuario)

    def test_avisa_perto_do_vencimento_uma_vez_por_dia(self, enviar):
        self.assertTrue(self._notificar_em(date(2026, 4, 27)))
        self.assertIn("vence em 3 dia(s)", enviar.call_args[0][1])
        self.assertFalse(self._notificar_em(date(2026, 4, 27)))
        self.assertEqual(enviar.call_count, 1)

    def test_nao_avisa_longe_do_vencimento(self, enviar):
        self.assertFalse(self._notificar_em(date(2026, 4, 10)))
        enviar.assert_not_called()

    def test_avisa_vencida_com_valor_atualizado(self, enviar):
        self.assertTrue(self._notificar_em(date(2026, 5, 10)))
        self.assertIn("VENCIDA há 10 dia(s) - R$ 3127.44", enviar.call_args[0][1])

    def test_nao_avisa_guia_paga(self, enviar):
        [darf] = sincronizar_darfs(self.usuario)
        registrar_pagamento(darf, date(2026, 4, 29), Decimal("2998.50"))
        self.assertFalse(self._notificar_em(date(2026, 5, 10)))


class SimuladorVendaTests(SemBancoCentralMixin, TestCase):
    """Lote de 100 PETR4 a R$ 100 com 50 já vendidas a R$ 300 em março/2026 (R$ 15.000 vendidos, R$ 10.000 de ganho isento)."""

    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="simulador_ir", password="SenhaForte123!")
        self.client.login(username="simulador_ir", password="SenhaForte123!")
        self.lote = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="PETR4"), tipo=Operacao.COMPRA,
            quantidade=100, preco_unitario=Decimal("100.00"), data_operacao=D1,
        )
        self.venda = VendaLote.objects.create(operacao=self.lote, quantidade=50, preco=Decimal("300.00"), data=D2)
        recalcular_resumo_vendas(self.lote)

    def _simular(self, quantidade, preco, data=D3, venda=None, operacao=None):
        params = {"operacao": (operacao or self.lote).id, "quantidade": quantidade, "preco": preco, "data": data.isoformat()}
        if venda:
            params["venda"] = venda.id
        return self.client.get(reverse("tributario:simular_venda"), params).json()

    def test_venda_que_passa_do_limite_avisa_perda_da_isencao(self):
        dados = self._simular(30, "200.00")
        self.assertTrue(dados["ok"])
        self.assertTrue(dados["perde_isencao"])
        self.assertEqual(Decimal(dados["vendas_acoes_antes"]), Decimal("15000.00"))
        self.assertEqual(Decimal(dados["vendas_acoes_depois"]), Decimal("21000.00"))
        self.assertEqual(Decimal(dados["ganho_isento_antes"]), Decimal("10000.00"))
        self.assertEqual(Decimal(dados["ganho_fiscal"]), Decimal("3000.00"))
        # 15% sobre o ganho do mês inteiro (10.000 que eram isentos + 3.000 desta venda)
        self.assertEqual(Decimal(dados["ir_adicional"]), Decimal("1950.00"))
        self.assertEqual(Decimal(dados["lucro_apos_ir"]), Decimal("1050.00"))

    def test_venda_dentro_do_limite_nao_tem_ir(self):
        dados = self._simular(10, "200.00")
        self.assertFalse(dados["perde_isencao"])
        self.assertTrue(dados["isento_depois"])
        self.assertEqual(Decimal(dados["ir_adicional"]), Decimal("0"))
        self.assertEqual(Decimal(dados["disponivel_depois"]), Decimal("3000.00"))

    def test_correcao_compara_com_o_mes_sem_a_venda_original(self):
        dados = self._simular(50, "400.00", data=D2, venda=self.venda)
        self.assertEqual(Decimal(dados["vendas_acoes_antes"]), Decimal("0"))
        self.assertEqual(Decimal(dados["vendas_acoes_depois"]), Decimal("20000.00"))
        self.assertTrue(dados["isento_depois"])
        self.assertEqual(Decimal(dados["ganho_fiscal"]), Decimal("15000.00"))

    def test_fii_nao_tem_isencao(self):
        lote_fii = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="HGLG11"), tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("100.00"), data_operacao=D1,
        )
        dados = self._simular(10, "120.00", operacao=lote_fii)
        self.assertFalse(dados["classe_com_isencao"])
        self.assertEqual(Decimal(dados["ir_adicional"]), Decimal("40.00"))  # 20% de 200

    def test_dados_incompletos_e_lote_de_outro_usuario(self):
        self.assertFalse(self._simular(10, "")["ok"])
        outro = User.objects.create_user(username="outro_simulador", password="SenhaForte123!")
        lote_alheio = Operacao.objects.create(
            usuario=outro, ativo=self.lote.ativo, tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("10.00"), data_operacao=D1,
        )
        resposta = self.client.get(reverse("tributario:simular_venda"), {"operacao": lote_alheio.id, "quantidade": 1, "preco": "10"})
        self.assertEqual(resposta.status_code, 404)

    def test_simulacao_nao_grava_nada(self):
        self._simular(30, "200.00")
        self.assertEqual(VendaLote.objects.count(), 1)

    def test_tela_de_venda_tem_o_simulador(self):
        resposta = self.client.get(reverse("core:operacao_vender", args=[self.lote.id]))
        self.assertContains(resposta, 'id="simulacao-ir"')
        self.assertContains(resposta, reverse("tributario:simular_venda"))


class AvisoLimiteIsencaoTests(SemBancoCentralMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="limite_isencao", password="SenhaForte123!")
        self.client.login(username="limite_isencao", password="SenhaForte123!")
        self.lote = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="ITUB4"), tipo=Operacao.COMPRA,
            quantidade=1000, preco_unitario=Decimal("10.00"), data_operacao=D1,
        )

    def _vender(self, quantidade, preco="30.00"):
        resposta = self.client.post(reverse("core:operacao_vender", args=[self.lote.id]), {
            "quantidade": str(quantidade), "preco": preco, "data": D2.isoformat(),
        })
        return [str(m) for m in get_messages(resposta.wsgi_request)]

    def test_avisa_na_tela_perto_do_limite(self):
        mensagens = self._vender(550)  # 16.500 = 82,5% do limite
        self.assertTrue(any("restam R$ 3500.00" in m for m in mensagens), mensagens)

    def test_avisa_na_tela_quando_passa_do_limite(self):
        mensagens = self._vender(700)  # 21.000
        self.assertTrue(any("acima do limite" in m for m in mensagens), mensagens)

    def test_sem_aviso_com_folga(self):
        mensagens = self._vender(100)  # 3.000
        self.assertFalse(any("limite" in m for m in mensagens), mensagens)

    @patch("tributario.avisos.enviar_notificacao_telegram", return_value=True)
    def test_telegram_avisa_uma_vez_por_mes_e_situacao(self, enviar):
        PerfilUsuario.objects.create(usuario=self.usuario, telegram_bot_token="123:abc", telegram_chat_id="999")
        self._vender(700)
        self._vender(100)
        self.assertEqual(enviar.call_count, 1)
        self.assertTrue(AvisoLimiteIsencao.objects.filter(usuario=self.usuario, tipo=AvisoLimiteIsencao.EXCEDIDO).exists())

    @patch("tributario.avisos.enviar_notificacao_telegram", return_value=False)
    def test_falha_no_telegram_tenta_de_novo_na_proxima_venda(self, enviar):
        PerfilUsuario.objects.create(usuario=self.usuario, telegram_bot_token="123:abc", telegram_chat_id="999")
        self._vender(700)
        self._vender(100)
        self.assertEqual(enviar.call_count, 2)
        self.assertFalse(AvisoLimiteIsencao.objects.exists())


class EventoCorporativoCalculoPuroTests(TestCase):
    def test_desdobramento_muda_quantidade_e_mantem_custo(self):
        movimentos = {
            ("PETR4", D1): MovimentoDia(quantidade_compra=100, custo_compra=Decimal("3000.00")),
            ("PETR4", D3): MovimentoDia(quantidade_venda=50, valor_venda=Decimal("1000.00")),
        }
        evento = EventoCorporativo(tipo=EventoCorporativo.DESDOBRAMENTO, data=D2, proporcao=Decimal("2"))
        apuracao = calcular_apuracao(movimentos, {}, {"PETR4": [evento]})
        self.assertEqual(apuracao.resultados[0].resultado, Decimal("250.00"))  # 1000 - 50 x 15 (PM ajustado)
        self.assertEqual(apuracao.posicoes["PETR4"].quantidade, 150)

    def test_evento_depois_do_ultimo_movimento_ajusta_a_posicao(self):
        movimentos = {("PETR4", D1): MovimentoDia(quantidade_compra=100, custo_compra=Decimal("3000.00"))}
        evento = EventoCorporativo(
            tipo=EventoCorporativo.BONIFICACAO, data=D3, proporcao=Decimal("10"), custo_unitario_bonificacao=Decimal("5"),
        )
        posicao = calcular_apuracao(movimentos, {}, {"PETR4": [evento]}).posicoes["PETR4"]
        self.assertEqual(posicao.quantidade, 110)
        self.assertEqual(posicao.custo_total, Decimal("3050"))


class EventoCorporativoLotesTests(SemBancoCentralMixin, TestCase):
    """Lote de 100 PETR4 a R$ 30 em 02/03/2026, com 20 vendidas a R$ 35 em 03/03."""

    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="eventos_corp", password="SenhaForte123!")
        self.client.login(username="eventos_corp", password="SenhaForte123!")
        self.ativo = Ativo.objects.create(ticker="PETR4")
        self.lote = Operacao.objects.create(
            usuario=self.usuario, ativo=self.ativo, tipo=Operacao.COMPRA,
            quantidade=100, preco_unitario=Decimal("30.00"), data_operacao=D1, meta_lucro_pct=Decimal("10"),
        )
        VendaLote.objects.create(operacao=self.lote, quantidade=20, preco=Decimal("35.00"), data=D2)
        recalcular_resumo_vendas(self.lote)

    def _evento(self, tipo, proporcao, data=date(2026, 3, 10), custo=None):
        evento = EventoCorporativo.objects.create(
            usuario=self.usuario, ativo=self.ativo, tipo=tipo, data=data, proporcao=Decimal(proporcao),
            custo_unitario_bonificacao=custo,
        )
        return evento, aplicar_evento(evento)

    def test_desdobramento_move_o_saldo_para_lote_ajustado(self):
        _, [novo] = self._evento(EventoCorporativo.DESDOBRAMENTO, "2")
        self.lote.refresh_from_db()
        self.assertEqual((novo.quantidade, novo.preco_unitario, novo.data_operacao), (160, Decimal("15.00"), D1))
        self.assertEqual(novo.meta_lucro_pct, Decimal("10.00"))
        self.assertEqual((self.lote.quantidade, self.lote.quantidade_original, self.lote.saldo), (20, 100, 0))

        posicao = apurar_operacoes(self.usuario).posicoes["PETR4"]
        self.assertEqual((posicao.quantidade, posicao.custo_total_centavos), (160, Decimal("2400.00")))

        # venda depois do evento, a partir do lote novo: custo pelo preço médio ajustado (R$ 15)
        VendaLote.objects.create(operacao=novo, quantidade=40, preco=Decimal("20.00"), data=date(2026, 3, 15))
        resultados = {r.data: r.resultado for r in apurar_operacoes(self.usuario).resultados}
        self.assertEqual(resultados[D2], Decimal("100.00"))  # 20 x (35 - 30), antes do evento
        self.assertEqual(resultados[date(2026, 3, 15)], Decimal("200.00"))  # 800 - 40 x 15

    def test_grupamento(self):
        _, [novo] = self._evento(EventoCorporativo.GRUPAMENTO, "10")
        self.assertEqual((novo.quantidade, novo.preco_unitario), (8, Decimal("300.00")))

    def test_bonificacao_cria_lote_ao_custo_informado(self):
        _, [bonus] = self._evento(EventoCorporativo.BONIFICACAO, "10", custo=Decimal("5"))
        self.assertEqual((bonus.quantidade, bonus.preco_unitario, bonus.data_operacao), (8, Decimal("5.00"), date(2026, 3, 10)))
        posicao = apurar_operacoes(self.usuario).posicoes["PETR4"]
        self.assertEqual(posicao.quantidade, 88)
        self.assertEqual(posicao.custo_total_centavos, Decimal("2440.00"))  # 80 x 30 + 8 x 5

    def test_lote_do_evento_nao_debita_conta_corrente(self):
        _, [novo] = self._evento(EventoCorporativo.DESDOBRAMENTO, "2")
        sincronizar_lancamento_compra(novo)
        self.assertFalse(LancamentoContaCorrente.objects.filter(operacao=novo).exists())

    def test_desfazer_devolve_os_lotes(self):
        evento, _ = self._evento(EventoCorporativo.DESDOBRAMENTO, "2")
        desfazer_evento(evento)
        self.lote.refresh_from_db()
        self.assertEqual((self.lote.quantidade, self.lote.quantidade_original), (100, None))
        self.assertEqual(Operacao.objects.filter(usuario=self.usuario).count(), 1)
        self.assertFalse(EventoCorporativo.objects.exists())

    def test_nao_desfaz_com_venda_no_lote_novo(self):
        evento, [novo] = self._evento(EventoCorporativo.DESDOBRAMENTO, "2")
        VendaLote.objects.create(operacao=novo, quantidade=10, preco=Decimal("20.00"), data=date(2026, 3, 15))
        with self.assertRaises(EventoInvalido):
            desfazer_evento(evento)

    def test_nao_aplica_com_venda_registrada_depois_da_data_ex(self):
        VendaLote.objects.create(operacao=self.lote, quantidade=10, preco=Decimal("35.00"), data=date(2026, 3, 12))
        evento = EventoCorporativo.objects.create(
            usuario=self.usuario, ativo=self.ativo, tipo=EventoCorporativo.DESDOBRAMENTO,
            data=date(2026, 3, 10), proporcao=Decimal("2"),
        )
        with self.assertRaises(EventoInvalido):
            aplicar_evento(evento)

    def test_lote_ajustado_nao_pode_ser_editado_nem_excluido(self):
        _, [novo] = self._evento(EventoCorporativo.DESDOBRAMENTO, "2")
        self.assertRedirects(self.client.get(reverse("core:operacao_editar", args=[novo.id])), reverse("core:operacao_lista"))
        self.client.post(reverse("core:operacao_excluir", args=[self.lote.id]))
        self.assertTrue(Operacao.objects.filter(id=self.lote.id).exists())

    def test_tela_registra_e_desfaz_evento(self):
        resposta = self.client.post(reverse("tributario:eventos_corporativos"), {
            "ticker": "PETR4", "tipo": EventoCorporativo.DESDOBRAMENTO, "data": "2026-03-10", "proporcao": "2",
        })
        self.assertRedirects(resposta, reverse("tributario:eventos_corporativos"))
        evento = EventoCorporativo.objects.get()
        self.assertEqual(evento.lotes_gerados.get().quantidade, 160)

        self.client.post(reverse("tributario:evento_desfazer", args=[evento.id]))
        self.assertFalse(EventoCorporativo.objects.exists())

    def test_tela_mostra_erro_e_nao_grava_evento_invalido(self):
        resposta = self.client.post(reverse("tributario:eventos_corporativos"), {
            "ticker": "PETR4", "tipo": EventoCorporativo.BONIFICACAO, "data": "2026-03-10", "proporcao": "10",
        })
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "custo por ação")
        self.assertFalse(EventoCorporativo.objects.exists())


class ProventoTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user(username="proventos_ir", password="SenhaForte123!")
        self.client.login(username="proventos_ir", password="SenhaForte123!")

    def _registrar(self, ticker, tipo, valor, data="2026-05-15", ir="", lancar=True):
        dados = {"ticker": ticker, "tipo": tipo, "data_pagamento": data, "valor_bruto": valor, "ir_retido": ir}
        if lancar:
            dados["lancar_conta_corrente"] = "on"
        return self.client.post(reverse("tributario:proventos"), dados)

    def test_jcp_sugere_15_pct_e_credita_o_liquido(self):
        self._registrar("ITUB4", Provento.JCP, "1000.00")
        provento = Provento.objects.get()
        self.assertEqual(provento.ir_retido, Decimal("150.00"))
        self.assertEqual(provento.lancamento.valor, Decimal("850.00"))
        self.assertEqual(provento.lancamento.origem, LancamentoContaCorrente.ORIGEM_PROVENTO)

    def test_dividendo_acima_de_50_mil_da_mesma_empresa_no_mes_retem_10_pct(self):
        self._registrar("PETR3", Provento.DIVIDENDO, "30000.00")
        self._registrar("PETR4", Provento.DIVIDENDO, "25000.00")
        primeiro, segundo = Provento.objects.order_by("id")
        self.assertEqual(primeiro.ir_retido, Decimal("0"))
        self.assertEqual(segundo.ir_retido, Decimal("2500.00"))

    def test_dividendo_em_2025_nao_tem_retencao(self):
        self._registrar("PETR4", Provento.DIVIDENDO, "60000.00", data="2025-05-15")
        self.assertEqual(Provento.objects.get().ir_retido, Decimal("0"))

    def test_fii_e_isento_e_ir_informado_e_respeitado(self):
        self._registrar("HGLG11", Provento.RENDIMENTO_FII, "80.00")
        self._registrar("VALE3", Provento.JCP, "100.00", ir="14.00")
        self.assertEqual(Provento.objects.get(tipo=Provento.RENDIMENTO_FII).ir_retido, Decimal("0"))
        self.assertEqual(Provento.objects.get(tipo=Provento.JCP).ir_retido, Decimal("14.00"))

    def test_rejeita_ir_maior_que_o_bruto(self):
        resposta = self._registrar("VALE3", Provento.JCP, "100.00", ir="150.00")
        self.assertEqual(resposta.status_code, 200)
        self.assertFalse(Provento.objects.exists())

    def test_editar_sem_lancar_remove_credito_e_excluir_remove_tudo(self):
        self._registrar("ITUB4", Provento.JCP, "1000.00")
        provento = Provento.objects.get()
        self.client.post(reverse("tributario:provento_editar", args=[provento.id]), {
            "ticker": "ITUB4", "tipo": Provento.JCP, "data_pagamento": "2026-05-15", "valor_bruto": "1000.00", "ir_retido": "150.00",
        })
        provento.refresh_from_db()
        self.assertIsNone(provento.lancamento)
        self.assertFalse(LancamentoContaCorrente.objects.exists())

        self._registrar("ITUB4", Provento.DIVIDENDO, "200.00")
        dividendo = Provento.objects.get(tipo=Provento.DIVIDENDO)
        self.client.post(reverse("tributario:provento_excluir", args=[dividendo.id]))
        self.assertFalse(LancamentoContaCorrente.objects.exists())
        self.assertEqual(Provento.objects.count(), 1)

    def test_formulario_novo_deixa_ir_em_branco(self):
        resposta = self.client.get(reverse("tributario:proventos"))
        campo = resposta.context["form"]["ir_retido"]
        self.assertIsNone(campo.value())

    def test_lista_do_ano(self):
        self._registrar("ITUB4", Provento.JCP, "1000.00")
        resposta = self.client.get(reverse("tributario:proventos"), {"ano": "2026"})
        self.assertContains(resposta, "ITUB4")
        self.assertContains(resposta, "R$ 850,00")


class RelatorioAnualTests(SemBancoCentralMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.usuario = User.objects.create_user(username="declaracao", password="SenhaForte123!")
        self.client.login(username="declaracao", password="SenhaForte123!")
        petr = Ativo.objects.create(ticker="PETR4", nome="Petrobras")
        lote = Operacao.objects.create(
            usuario=self.usuario, ativo=petr, tipo=Operacao.COMPRA,
            quantidade=100, preco_unitario=Decimal("10.00"), data_operacao=date(2025, 6, 2),
        )
        VendaLote.objects.create(operacao=lote, quantidade=50, preco=Decimal("30.00"), data=date(2026, 3, 3))
        recalcular_resumo_vendas(lote)
        Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="HGLG11"), tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("100.00"), data_operacao=date(2026, 2, 2),
        )
        Provento.objects.create(usuario=self.usuario, ativo=petr, tipo=Provento.JCP, data_pagamento=date(2026, 4, 1),
                                valor_bruto=Decimal("1000.00"), ir_retido=Decimal("150.00"))
        Provento.objects.create(usuario=self.usuario, ativo=petr, tipo=Provento.DIVIDENDO, data_pagamento=date(2026, 4, 1),
                                valor_bruto=Decimal("500.00"))

    def _secao(self, titulo_inicio):
        return next(s for s in gerar_relatorio_anual(self.usuario, 2026) if s.titulo.startswith(titulo_inicio))

    def test_bens_e_direitos_com_custo_nos_dois_anos(self):
        linhas = {linha[1]: linha for linha in self._secao("Bens e Direitos").linhas}
        self.assertEqual(linhas["PETR4"][0], "03 - 01")
        self.assertEqual(linhas["PETR4"][3:], [100, "R$ 1.000,00", 50, "R$ 500,00"])
        self.assertIn("Petrobras", linhas["PETR4"][2])
        self.assertEqual(linhas["HGLG11"][0], "07 - 03")
        self.assertEqual(linhas["HGLG11"][3:], [0, "R$ 0,00", 10, "R$ 1.000,00"])

    def test_isentos_e_tributacao_exclusiva(self):
        isentos = self._secao("Rendimentos isentos").linhas
        self.assertIn(["Ganhos líquidos com ações - vendas até R$ 20 mil no mês", "", "", "", "R$ 1.000,00"], isentos)
        self.assertIn(["Lucros e dividendos recebidos", "PETR4", "R$ 500,00", "R$ 0,00", "R$ 500,00"], isentos)
        exclusiva = self._secao("Rendimentos sujeitos").linhas
        self.assertEqual(exclusiva, [["Juros sobre capital próprio", "PETR4", "R$ 1.000,00", "R$ 150,00", "R$ 850,00"]])

    def test_tela_e_downloads(self):
        resposta = self.client.get(reverse("tributario:relatorio_anual"), {"ano": "2026"})
        self.assertContains(resposta, "Bens e Direitos")
        self.assertContains(resposta, "R$ 850,00")

        excel = self.client.get(reverse("tributario:relatorio_anual_excel"), {"ano": "2026"})
        self.assertIn("attachment", excel["Content-Disposition"])
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(len(load_workbook(BytesIO(excel.content)).sheetnames), 7)

        pdf = self.client.get(reverse("tributario:relatorio_anual_pdf"), {"ano": "2026"})
        self.assertTrue(pdf.content.startswith(b"%PDF"))


# --------------------------------------------------------------------------
# Importação das notas de corretagem

from .models import MapeamentoTitulo, NotaCorretagem  # noqa: E402
from .notas import (  # noqa: E402
    ExclusaoBloqueada, ImportacaoInvalida, Negocio, NotaInvalida, NotaLida, chave_titulo, excluir_notas,
    importar_notas, ler_notas_pdf, simular_importacao, sugerir_ticker,
)
from .services import ORIGEM_MANUAIS, ORIGEM_NOTAS  # noqa: E402

DIA_NOTA = date(2026, 10, 5)


def _pdf_nota_sinacor(numero="145779266", data_pregao="05/10/2026") -> bytes:
    """PDF mínimo no layout da nota SINACOR (posições das colunas iguais às da Clear)."""
    import pymupdf

    documento = pymupdf.open()
    pagina = documento.new_page(width=595, height=842)

    def escrever(x, y, texto, tamanho=7):
        pagina.insert_text((x, y), texto, fontsize=tamanho)

    escrever(113, 46, "NOTA DE NEGOCIAÇÃO", 9)
    escrever(430, 56, "Nr. nota")
    escrever(432, 63, numero)
    escrever(517, 56, "Data pregão")
    escrever(518, 63, data_pregao)
    escrever(129, 80, "CLEAR CTVM S/A")
    cabecalho = [(35, "Q"), (43, "Negociação"), (94, "C/V"), (109, "Tipo"), (128, "mercado"), (169, "Prazo"),
                 (194, "Especificação"), (249, "do"), (261, "título"), (294, "Obs."), (314, "(*)"),
                 (327, "Quantidade"), (378, "Preço"), (440, "Valor"), (544, "D/C")]
    for x, texto in cabecalho:
        escrever(x, 245, texto)
    linhas = [
        ("C", "FRACIONARIO", ["MAGAZ", "LUIZA", "ON", "NM"], "@", "10", "9,63", "96,30", "D"),
        ("C", "FRACIONARIO", ["PETROBRAS", "PN", "N2"], "D", "2", "53,52", "107,04", "D"),
        ("V", "FRACIONARIO", ["PETROBRAS", "PN", "N2"], "D#", "2", "53,72", "107,44", "C"),
        ("C", "OPCAO DE COMPRA", ["PETRK450", "PN"], "", "100", "0,50", "50,00", "D"),
    ]
    y = 256
    for cv, mercado, especificacao, obs, qtd, preco, valor, dc in linhas:
        escrever(43, y, "1-BOVESPA")
        escrever(98, y, cv)
        x = 109
        for palavra in mercado.split():
            escrever(x, y, palavra)
            x += 20
        x = 194
        for palavra in especificacao:
            escrever(x, y + 1, palavra, 5)
            x += 25
        if obs:
            escrever(304, y, obs)
        escrever(368, y, qtd)
        escrever(418, y, preco)
        escrever(518, y, valor)
        escrever(549, y, dc)
        y += 10
    resumo = [
        (470, "Taxa de liquidação 0,20 D"), (478, "Taxa de Registro 0,00"), (536, "Total Bovespa / Soma 0,05 D"),
        (593, "I.R.R.F. s/ operações, base R$107,44 0,00"), (621, "Total Custos / Despesas 0,00"),
        (630, "Líquido para 07/10/2026 146,15 D"),
    ]
    for y, texto in resumo:
        escrever(300, y, texto)
    conteudo = documento.tobytes()
    documento.close()
    return conteudo


def _nota(numero, dia, negocios, custos="0"):
    return NotaLida(
        corretora="CLEAR CTVM S/A", numero=numero, data=dia, custos=Decimal(custos),
        negocios=[
            Negocio(cv=cv, mercado="FRACIONARIO", especificacao=esp, obs="", quantidade=qtd, preco=Decimal(preco),
                    valor=(Decimal(preco) * qtd).quantize(Decimal("0.01")))
            for cv, esp, qtd, preco in negocios
        ],
    )


TICKERS = {"MAGAZ LUIZA ON": "MGLU3", "PETROBRAS PN": "PETR4", "HYPERA ON": "HYPE3"}


class LeituraNotaPdfTests(TestCase):
    def test_le_negocios_taxas_e_ignora_opcoes(self):
        resultado = ler_notas_pdf(_pdf_nota_sinacor(), "nota.pdf")
        self.assertEqual(len(resultado.notas), 1)
        nota = resultado.notas[0]
        self.assertEqual((nota.numero, nota.data, nota.corretora), ("145779266", DIA_NOTA, "CLEAR CTVM S/A"))
        self.assertEqual(nota.custos, Decimal("0.25"))
        self.assertEqual(nota.liquido, Decimal("-146.15"))
        self.assertEqual(len(nota.negocios), 4)
        self.assertEqual(len(nota.importaveis), 3)
        mglu = nota.negocios[0]
        self.assertEqual(
            (mglu.cv, mglu.especificacao, mglu.obs, mglu.quantidade, mglu.preco),
            ("C", "MAGAZ LUIZA ON NM", "@", 10, Decimal("9.63")),
        )
        self.assertTrue(nota.negocios[2].day_trade)
        self.assertTrue(any("não é importado" in aviso for aviso in nota.avisos))

    def test_pdf_invalido(self):
        with self.assertRaises(NotaInvalida):
            ler_notas_pdf(b"nada a ver", "x.pdf")

    def test_chave_e_sugestao_de_ticker(self):
        self.assertEqual(chave_titulo("COPASA          ON EJ NM"), "COPASA ON")
        self.assertEqual(chave_titulo("TAESA UNT N2"), "TAESA UNT")
        self.assertEqual(sugerir_ticker("TAESA UNT N2"), "TAEE11")
        self.assertEqual(sugerir_ticker("NU HOLDINGS DRN"), "ROXO34")
        from core.models import AcaoB3
        AcaoB3.objects.create(ticker="XPTO3", nome="EMPRESA XPTO PARTICIPAÇÕES S.A.")
        self.assertEqual(sugerir_ticker("XPTO PART ON NM"), "XPTO3")
        MapeamentoTitulo.objects.create(especificacao="XPTO PART ON", ticker="XPTO11")
        self.assertEqual(sugerir_ticker("XPTO PART ON NM"), "XPTO11")


class ImportacaoNotasTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user("importa", password="SenhaForte123!")

    def test_compras_com_taxas_rateadas_e_day_trade(self):
        nota = _nota("1", DIA_NOTA, [
            ("C", "MAGAZ LUIZA ON NM", 10, "10.00"),
            ("C", "PETROBRAS PN N2", 2, "50.00"),
            ("V", "PETROBRAS PN N2", 1, "51.00"),
            ("V", "PETROBRAS PN N2", 1, "51.00"),  # mesmo preço: vira uma venda só
        ], custos="0.31")
        resumo = importar_notas(self.usuario, [nota], TICKERS)

        registro = NotaCorretagem.objects.get(usuario=self.usuario)
        self.assertEqual(registro.operacoes.count(), 2)
        self.assertEqual(registro.vendas.count(), 1)
        mglu = Operacao.objects.get(usuario=self.usuario, ativo__ticker="MGLU3")
        petr = Operacao.objects.get(usuario=self.usuario, ativo__ticker="PETR4")
        venda = VendaLote.objects.get(operacao=petr)
        self.assertEqual(mglu.taxas_compra + petr.taxas_compra + venda.taxas, Decimal("0.31"))
        self.assertEqual(venda.quantidade, 2)
        self.assertEqual(petr.saldo, 0)
        self.assertEqual(len(resumo.linhas), 3)
        self.assertEqual(LancamentoContaCorrente.objects.filter(conta__usuario=self.usuario).count(), 3)
        self.assertEqual(MapeamentoTitulo.objects.get(especificacao="PETROBRAS PN").ticker, "PETR4")
        self.assertEqual([r.modalidade for r in apurar_operacoes(self.usuario).resultados], [DAY_TRADE])

    def test_venda_baixa_primeiro_no_mesmo_dia_e_depois_no_mais_antigo(self):
        importar_notas(self.usuario, [
            _nota("1", D1, [("C", "HYPERA ON NM", 10, "20.00")]),
            _nota("2", D2, [("C", "HYPERA ON NM", 5, "21.00")]),
            _nota("3", D3, [("C", "HYPERA ON NM", 3, "22.00"), ("V", "HYPERA ON NM", 6, "23.00")]),
        ], TICKERS)
        vendas = {v.operacao.data_operacao: v.quantidade for v in VendaLote.objects.select_related("operacao")}
        self.assertEqual(vendas, {D3: 3, D1: 3})

    def test_venda_sem_saldo_nao_grava_nada(self):
        with self.assertRaises(ImportacaoInvalida):
            importar_notas(self.usuario, [
                _nota("1", D1, [("C", "HYPERA ON NM", 1, "20.00")]),
                _nota("2", D2, [("V", "HYPERA ON NM", 5, "21.00")]),
            ], TICKERS)
        self.assertFalse(NotaCorretagem.objects.exists())
        self.assertFalse(Operacao.objects.exists())

    def test_nota_repetida_e_pulada(self):
        nota = _nota("1", D1, [("C", "HYPERA ON NM", 1, "20.00")])
        importar_notas(self.usuario, [nota], TICKERS)
        resumo = importar_notas(self.usuario, [nota], TICKERS)
        self.assertEqual(resumo.notas_repetidas, [f"1 ({D1:%d/%m/%Y})"])
        self.assertEqual(Operacao.objects.count(), 1)

    def test_simulacao_nao_grava(self):
        resumo, meses, erro = simular_importacao(self.usuario, [
            _nota("1", D1, [("C", "HYPERA ON NM", 10, "20.00")]),
            _nota("2", D2, [("V", "HYPERA ON NM", 10, "22.00")]),
        ], TICKERS)
        self.assertEqual(erro, "")
        self.assertEqual(len(resumo.linhas), 2)
        self.assertEqual(meses[0].grupos[GRUPO_COMUM].resultado, Decimal("20.00"))
        self.assertFalse(Operacao.objects.exists())
        self.assertFalse(NotaCorretagem.objects.exists())

    def test_excluir_notas_bloqueia_venda_de_outra_nota_e_desfaz_baixa(self):
        importar_notas(self.usuario, [
            _nota("1", D1, [("C", "HYPERA ON NM", 10, "20.00")]),
            _nota("2", D2, [("V", "HYPERA ON NM", 4, "22.00")]),
        ], TICKERS)
        with self.assertRaises(ExclusaoBloqueada):
            excluir_notas(self.usuario, NotaCorretagem.objects.filter(numero="1"))

        self.assertEqual(excluir_notas(self.usuario, NotaCorretagem.objects.filter(numero="2")), 1)
        lote = Operacao.objects.get()
        self.assertEqual((lote.quantidade_vendida, lote.saldo), (0, 10))
        self.assertFalse(VendaLote.objects.exists())
        self.assertEqual(excluir_notas(self.usuario, NotaCorretagem.objects.all()), 1)
        self.assertFalse(Operacao.objects.exists())
        self.assertFalse(LancamentoContaCorrente.objects.exists())

    def test_compra_ja_lancada_a_mao_e_vinculada_sem_duplicar(self):
        manual = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="HYPE3"), tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("20.05"), data_operacao=D1, observacao="Lançada à mão.",
        )
        resumo = importar_notas(self.usuario, [_nota("1", D1, [("C", "HYPERA ON NM", 10, "20.00")], custos="0.10")], TICKERS)

        self.assertEqual(Operacao.objects.filter(usuario=self.usuario).count(), 1)
        manual.refresh_from_db()
        self.assertEqual(manual.preco_unitario, Decimal("20.00"))
        self.assertEqual(manual.taxas_compra, Decimal("0.10"))
        self.assertIn("Vinculada à nota 1", manual.observacao)
        self.assertEqual(list(NotaCorretagem.objects.get().operacoes.all()), [manual])
        self.assertEqual(resumo.linhas[0].destino, "lote existente")
        self.assertEqual(LancamentoContaCorrente.objects.get(operacao=manual).valor, Decimal("200.10"))

    def test_compra_a_mao_com_quantidade_diferente_nao_e_vinculada(self):
        Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="HYPE3"), tipo=Operacao.COMPRA,
            quantidade=5, preco_unitario=Decimal("20"), data_operacao=D1,
        )
        resumo = importar_notas(self.usuario, [_nota("1", D1, [("C", "HYPERA ON NM", 10, "20.00")])], TICKERS)
        self.assertEqual(Operacao.objects.filter(usuario=self.usuario).count(), 2)
        self.assertEqual(resumo.linhas[0].destino, "lote novo")

    def test_apuracao_por_origem(self):
        manual = Operacao.objects.create(
            usuario=self.usuario, ativo=Ativo.objects.create(ticker="VALE3"), tipo=Operacao.COMPRA,
            quantidade=10, preco_unitario=Decimal("50"), data_operacao=D1,
        )
        VendaLote.objects.create(operacao=manual, quantidade=10, preco=Decimal("55"), data=D2)
        importar_notas(self.usuario, [
            _nota("1", D1, [("C", "HYPERA ON NM", 10, "20.00")]),
            _nota("2", D2, [("V", "HYPERA ON NM", 10, "21.00")]),
        ], TICKERS)

        def tickers(origem):
            return sorted(r.ticker for r in apurar_operacoes(self.usuario, origem=origem).resultados)

        self.assertEqual(tickers("todas"), ["HYPE3", "VALE3"])
        self.assertEqual(tickers(ORIGEM_NOTAS), ["HYPE3"])
        self.assertEqual(tickers(ORIGEM_MANUAIS), ["VALE3"])


class TelasNotasTests(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user("tela_notas", password="SenhaForte123!")
        self.client.login(username="tela_notas", password="SenhaForte123!")

    def test_upload_previa_confirmacao_e_exclusao_do_mes(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        arquivo = SimpleUploadedFile("nota.pdf", _pdf_nota_sinacor(), content_type="application/pdf")
        resposta = self.client.post(reverse("tributario:notas_importar"), {"arquivos": [arquivo]})
        self.assertRedirects(resposta, reverse("tributario:notas_previa"))

        previa = self.client.get(reverse("tributario:notas_previa"))
        self.assertContains(previa, "MGLU3")
        self.assertContains(previa, "Confirmar importação")
        self.assertFalse(Operacao.objects.exists())

        dados = {f"ticker_{t['indice']}": t["ticker"] for t in previa.context["titulos"]}
        resposta = self.client.post(reverse("tributario:notas_previa"), {**dados, "acao": "confirmar"})
        self.assertRedirects(
            resposta, reverse("tributario:imposto_renda") + "?origem=notas", fetch_redirect_response=False
        )
        self.assertEqual(NotaCorretagem.objects.filter(usuario=self.usuario).count(), 1)
        self.assertEqual(Operacao.objects.filter(usuario=self.usuario).count(), 2)

        tela = self.client.get(reverse("tributario:imposto_renda"), {"origem": "notas", "ano": "2026"})
        self.assertContains(tela, "Excluir notas do mês (1)")
        self.assertContains(tela, "Só notas importadas")
        self.assertContains(self.client.get(reverse("tributario:notas_importar")), "145779266")

        resposta = self.client.post(reverse("tributario:notas_mes_excluir", args=[2026, 10]))
        self.assertRedirects(resposta, reverse("tributario:imposto_renda"), fetch_redirect_response=False)
        self.assertFalse(NotaCorretagem.objects.exists())
        self.assertFalse(Operacao.objects.exists())

    def test_ticker_em_branco_pede_preenchimento(self):
        session = self.client.session
        session["notas_corretagem_lidas"] = [
            _nota("9", D1, [("C", "EMPRESA DESCONHECIDA ON NM", 1, "1.00")]).como_dict()
        ]
        session.save()
        resposta = self.client.get(reverse("tributario:notas_previa"))
        self.assertContains(resposta, "Informe o ticker de")
        self.assertNotContains(resposta, "Confirmar importação")

    def test_nao_exclui_nota_de_outro_usuario(self):
        outro = User.objects.create_user("outro_notas", password="SenhaForte123!")
        importar_notas(outro, [_nota("1", D1, [("C", "HYPERA ON NM", 1, "20.00")])], TICKERS)
        nota = NotaCorretagem.objects.get()
        self.assertEqual(self.client.post(reverse("tributario:nota_excluir", args=[nota.id])).status_code, 404)
        self.client.post(reverse("tributario:notas_mes_excluir", args=[D1.year, D1.month]))
        self.assertTrue(NotaCorretagem.objects.exists())
